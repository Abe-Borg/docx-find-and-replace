"""
main.py
Word Document Batch Find & Replace - GUI Application

A tkinter-based GUI for performing batch find-and-replace operations
across multiple Word documents (.docx) with a preview-before-commit workflow.

Threading model: scanning and applying run on daemon worker threads; every
widget update is marshalled back to the main thread with `root.after`. The
results tree is inert while a worker is running, because the worker reads the
same Match objects the tree would mutate. A worker never reads a tk variable:
everything it needs is snapshotted on the main thread before it starts.
"""

import os
import re
import subprocess
import sys
import threading
import tkinter as tk
from datetime import datetime
from tkinter import ttk, filedialog, messagebox
from typing import Dict, List, Optional

import settings as app_settings
from document_processor import (Backup, FileResult, Match, Rule, apply_changes,
                                find_backups, restore_backups, scan_documents_detailed,
                                write_change_log, _normalise_rules)
from version import APP_NAME, __version__


# Prefix shown before each match's location in the results tree.
LOCATION_PREFIXES = {
    'table': '[Table] ',
    'header': '[Header] ',
    'footer': '[Footer] ',
    'footnote': '[Footnote] ',
    'endnote': '[Endnote] ',
    'comment': '[Comment] ',
}

CHECKED = "☑"      # ☑
UNCHECKED = "☐"    # ☐
PARTIAL = "☒"      # ☒
WARNING = "⚠"      # ⚠
ARROW = "→"

CHANGE_LOG_PREFIX = "DocxFindReplace-changelog-"

_GEOMETRY_RE = re.compile(r"^\d+x\d+([+-]\d+[+-]\d+)?$")


def _open_document(path: str) -> None:
    """Open a document in whatever the platform associates with .docx."""
    if hasattr(os, "startfile"):
        os.startfile(path)                      # noqa: S606 - Windows only
        return
    opener = "open" if sys.platform == "darwin" else "xdg-open"
    subprocess.Popen([opener, path])


def _plural(count: int, singular: str, plural: Optional[str] = None) -> str:
    plural = plural if plural is not None else singular + "s"
    return f"{count} {singular if count == 1 else plural}"


def _rule_text(find_text: str, replace_text: str) -> str:
    return f"{find_text} {ARROW} {replace_text}"


class FindReplaceApp:
    """Main application window."""

    def __init__(self, root: tk.Tk, settings_path: Optional[str] = None):
        self.root = root
        self.root.title(f"{APP_NAME} {__version__}")
        self.root.geometry("980x760")
        self.root.minsize(760, 560)
        self.settings_path = settings_path or app_settings.default_settings_path()

        # State
        self.scan_results: List[FileResult] = []
        self.all_matches: List[Match] = []
        self.match_tree_map: Dict[str, Match] = {}   # tree item id -> Match object
        self.file_tree_map: Dict[str, FileResult] = {}  # tree item id -> FileResult
        self.rules: List[Rule] = []                  # the set, in display order
        self.rule_iids: List[str] = []               # rules_tree rows, parallel to rules

        # 'scan', 'apply', or None. Guards the tree and the close button.
        self.processing_kind: Optional[str] = None
        self.scan_cancel: Optional[threading.Event] = None
        self.close_when_idle = False

        # The settings the current results were produced with. Apply uses these
        # rather than reading the widgets live, so a setting changed after the
        # preview cannot be applied to offsets found under the old one.
        self.scan_folder = ""
        self.scan_rules: List[Rule] = []
        self.scan_case_sensitive = True
        self.scan_whole_word = False
        self.scan_recursive = False

        # Configure styles
        style = ttk.Style()
        style.configure("TButton", padding=4)
        style.configure("Header.TLabel", font=("Segoe UI", 10, "bold"))
        style.configure("Status.TLabel", font=("Segoe UI", 9))

        self._build_ui()
        self._load_settings()

        # Changing what gets searched invalidates the results that are on screen.
        for var in (self.folder_var, self.case_var, self.whole_word_var,
                    self.recursive_var):
            var.trace_add("write", self._on_search_settings_changed)

        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

    @property
    def is_processing(self) -> bool:
        return self.processing_kind is not None

    # ------------------------------------------------------------------
    # Layout
    # ------------------------------------------------------------------

    def _build_ui(self):
        """Build the complete GUI layout."""
        main_frame = ttk.Frame(self.root, padding=10)
        main_frame.pack(fill=tk.BOTH, expand=True)

        # --- Input Section ---
        input_frame = ttk.LabelFrame(main_frame, text="Search Settings", padding=8)
        input_frame.pack(fill=tk.X, pady=(0, 8))

        # Folder row
        folder_frame = ttk.Frame(input_frame)
        folder_frame.pack(fill=tk.X, pady=2)
        ttk.Label(folder_frame, text="Folder:", width=8).pack(side=tk.LEFT)
        self.folder_var = tk.StringVar()
        self.folder_entry = ttk.Entry(folder_frame, textvariable=self.folder_var)
        self.folder_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 4))
        self.browse_btn = ttk.Button(folder_frame, text="Browse...", command=self._browse_folder)
        self.browse_btn.pack(side=tk.LEFT)
        self.recursive_var = tk.BooleanVar(value=False)
        self.recursive_check = ttk.Checkbutton(folder_frame, text="Include subfolders",
                                               variable=self.recursive_var)
        self.recursive_check.pack(side=tk.LEFT, padx=(12, 0))

        # Rule editor row
        editor_frame = ttk.Frame(input_frame)
        editor_frame.pack(fill=tk.X, pady=(6, 2))
        ttk.Label(editor_frame, text="Find:", width=8).pack(side=tk.LEFT)
        self.find_var = tk.StringVar()
        self.find_entry = ttk.Entry(editor_frame, textvariable=self.find_var)
        self.find_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Label(editor_frame, text="Replace:", width=9).pack(side=tk.LEFT, padx=(8, 0))
        self.replace_var = tk.StringVar()
        self.replace_entry = ttk.Entry(editor_frame, textvariable=self.replace_var)
        self.replace_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.add_rule_btn = ttk.Button(editor_frame, text="Add / Update",
                                       command=self._add_rule)
        self.add_rule_btn.pack(side=tk.LEFT, padx=(8, 0))
        self.remove_rule_btn = ttk.Button(editor_frame, text="Remove",
                                          command=self._remove_rules)
        self.remove_rule_btn.pack(side=tk.LEFT, padx=(4, 0))

        # Rule set list
        rules_frame = ttk.Frame(input_frame)
        rules_frame.pack(fill=tk.X, pady=2)
        ttk.Label(rules_frame, text="Rules:", width=8).pack(side=tk.LEFT, anchor="n")
        list_frame = ttk.Frame(rules_frame)
        list_frame.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.rules_tree = ttk.Treeview(list_frame, columns=("find", "replace"),
                                       show="headings", height=4, selectmode="extended")
        self.rules_tree.heading("find", text="Find")
        self.rules_tree.heading("replace", text="Replace with")
        self.rules_tree.column("find", width=300, minwidth=120)
        self.rules_tree.column("replace", width=300, minwidth=120)
        rules_scroll = ttk.Scrollbar(list_frame, orient=tk.VERTICAL,
                                     command=self.rules_tree.yview)
        self.rules_tree.configure(yscrollcommand=rules_scroll.set)
        self.rules_tree.grid(row=0, column=0, sticky="nsew")
        rules_scroll.grid(row=0, column=1, sticky="ns")
        list_frame.columnconfigure(0, weight=1)
        self.rules_tree.bind("<<TreeviewSelect>>", self._on_rule_selected)

        set_buttons = ttk.Frame(rules_frame)
        set_buttons.pack(side=tk.LEFT, fill=tk.Y, padx=(8, 0))
        self.load_set_btn = ttk.Button(set_buttons, text="Load set...",
                                       command=self._load_rule_set)
        self.load_set_btn.pack(fill=tk.X)
        self.save_set_btn = ttk.Button(set_buttons, text="Save set...",
                                       command=self._save_rule_set)
        self.save_set_btn.pack(fill=tk.X, pady=(4, 0))

        # Options and preview row
        options_frame = ttk.Frame(input_frame)
        options_frame.pack(fill=tk.X, pady=(6, 2))

        self.case_var = tk.BooleanVar(value=True)
        self.case_check = ttk.Checkbutton(options_frame, text="Case sensitive",
                                          variable=self.case_var)
        self.case_check.pack(side=tk.LEFT)

        self.whole_word_var = tk.BooleanVar(value=False)
        self.whole_word_check = ttk.Checkbutton(options_frame, text="Whole word only",
                                                variable=self.whole_word_var)
        self.whole_word_check.pack(side=tk.LEFT, padx=(16, 0))

        self.preview_btn = ttk.Button(options_frame, text="Preview Changes",
                                      command=self._start_preview)
        self.preview_btn.pack(side=tk.RIGHT)

        # --- Results Section ---
        results_frame = ttk.LabelFrame(main_frame, text="Results", padding=8)
        results_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 8))

        # Toolbar for select all / deselect all
        toolbar = ttk.Frame(results_frame)
        toolbar.pack(fill=tk.X, pady=(0, 4))
        self.select_all_btn = ttk.Button(toolbar, text="Select All", command=self._select_all)
        self.select_all_btn.pack(side=tk.LEFT, padx=(0, 4))
        self.deselect_all_btn = ttk.Button(toolbar, text="Deselect All", command=self._deselect_all)
        self.deselect_all_btn.pack(side=tk.LEFT)
        self.restore_btn = ttk.Button(toolbar, text="Restore backups...",
                                      command=self._open_restore_dialog)
        self.restore_btn.pack(side=tk.LEFT, padx=(16, 0))
        self.match_count_label = ttk.Label(toolbar, text="", style="Status.TLabel")
        self.match_count_label.pack(side=tk.RIGHT)

        # Treeview with checkboxes
        tree_frame = ttk.Frame(results_frame)
        tree_frame.pack(fill=tk.BOTH, expand=True)

        self.tree = ttk.Treeview(tree_frame, columns=("rule", "section"),
                                 show="tree headings", selectmode="browse")
        self.tree.heading("#0", text="Match")
        self.tree.heading("rule", text="Rule")
        self.tree.heading("section", text="Section")
        self.tree.column("#0", width=460, minwidth=240)
        self.tree.column("rule", width=200, minwidth=100)
        self.tree.column("section", width=220, minwidth=100)
        self.tree.tag_configure("conflict", foreground="#a05a00")
        self.tree.tag_configure("error", foreground="#a00000")

        # Scrollbars
        v_scroll = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=self.tree.yview)
        h_scroll = ttk.Scrollbar(tree_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        v_scroll.grid(row=0, column=1, sticky="ns")
        h_scroll.grid(row=1, column=0, sticky="ew")
        tree_frame.columnconfigure(0, weight=1)
        tree_frame.rowconfigure(0, weight=1)

        # Click handlers: single click toggles, double click opens the file.
        self.tree.bind("<Button-1>", self._on_tree_click)
        self.tree.bind("<Double-1>", self._on_tree_double_click)

        # --- Apply options ---
        apply_frame = ttk.LabelFrame(main_frame, text="Apply Options", padding=8)
        apply_frame.pack(fill=tk.X, pady=(0, 8))

        self.backup_var = tk.BooleanVar(value=True)
        self.backup_check = ttk.Checkbutton(apply_frame, text="Create backups (timestamped .bak)",
                                            variable=self.backup_var)
        self.backup_check.pack(side=tk.LEFT)

        self.changelog_var = tk.BooleanVar(value=True)
        self.changelog_check = ttk.Checkbutton(apply_frame, text="Write change log (CSV)",
                                               variable=self.changelog_var)
        self.changelog_check.pack(side=tk.LEFT, padx=(16, 0))

        self.tracked_var = tk.BooleanVar(value=False)
        self.tracked_check = ttk.Checkbutton(apply_frame, text="Write as tracked changes",
                                             variable=self.tracked_var)
        self.tracked_check.pack(side=tk.LEFT, padx=(16, 0))

        ttk.Label(apply_frame, text="Author:").pack(side=tk.LEFT, padx=(16, 4))
        self.author_var = tk.StringVar()
        self.author_entry = ttk.Entry(apply_frame, textvariable=self.author_var, width=18)
        self.author_entry.pack(side=tk.LEFT)

        # --- Bottom Section ---
        bottom_frame = ttk.Frame(main_frame)
        bottom_frame.pack(fill=tk.X)

        self.apply_btn = ttk.Button(bottom_frame, text="Apply Selected Changes",
                                    command=self._start_apply, state=tk.DISABLED)
        self.apply_btn.pack(side=tk.RIGHT)

        self.progress_var = tk.StringVar(value="Ready")
        self.status_label = ttk.Label(bottom_frame, textvariable=self.progress_var,
                                      style="Status.TLabel")
        self.status_label.pack(side=tk.LEFT, fill=tk.X, expand=True)

    # ------------------------------------------------------------------
    # Settings persistence
    # ------------------------------------------------------------------

    def _load_settings(self):
        """Restore the last session's folder, rules and options."""
        try:
            saved = app_settings.load_settings(self.settings_path)
        except Exception:
            return
        self.folder_var.set(saved["folder"])
        self.case_var.set(saved["case_sensitive"])
        self.whole_word_var.set(saved["whole_word"])
        self.recursive_var.set(saved["recursive"])
        self.backup_var.set(saved["create_backups"])
        self.changelog_var.set(saved["write_change_log"])
        self.tracked_var.set(saved["tracked_changes"])
        self.author_var.set(saved["author"])
        self._set_rules(saved["rules"], invalidate=False)
        if saved["geometry"] and _GEOMETRY_RE.match(saved["geometry"]):
            try:
                self.root.geometry(saved["geometry"])
            except Exception:
                pass

    def _current_settings(self) -> dict:
        try:
            geometry = self.root.geometry()
        except Exception:
            geometry = ""
        return {
            "folder": self.folder_var.get() or "",
            "rules": list(self.rules),
            "case_sensitive": bool(self.case_var.get()),
            "whole_word": bool(self.whole_word_var.get()),
            "recursive": bool(self.recursive_var.get()),
            "create_backups": bool(self.backup_var.get()),
            "write_change_log": bool(self.changelog_var.get()),
            "tracked_changes": bool(self.tracked_var.get()),
            "author": self.author_var.get() or "",
            "geometry": geometry if isinstance(geometry, str) else "",
        }

    def _save_settings(self):
        """Remember the current settings. Never lets a failure stop a close."""
        try:
            app_settings.save_settings(self._current_settings(), self.settings_path)
        except Exception:
            pass

    # ------------------------------------------------------------------
    # Input handling
    # ------------------------------------------------------------------

    def _browse_folder(self):
        """Open folder selection dialog."""
        folder = filedialog.askdirectory(title="Select folder containing .docx files")
        if folder:
            self.folder_var.set(folder)

    def _pending_rule(self) -> Optional[Rule]:
        """The rule typed in the editor but not yet added, if any."""
        find_text = self.find_var.get() or ""
        if not find_text:
            return None
        return Rule(find_text, self.replace_var.get() or "")

    def _validate_inputs(self) -> bool:
        """Check that a folder and at least one usable rule are provided."""
        folder = (self.folder_var.get() or "").strip()

        if not folder:
            messagebox.showwarning("Missing Input", "Please select a folder.")
            return False
        if not os.path.isdir(folder):
            messagebox.showwarning("Invalid Folder", "The selected folder does not exist.")
            return False

        # A rule still sitting in the editor is what the user meant to search
        # for; add it rather than making them click twice.
        if self._pending_rule() is not None:
            self._add_rule()

        if not self.rules:
            messagebox.showwarning("Missing Input", "Please enter text to find.")
            return False
        try:
            _normalise_rules(self.rules, bool(self.case_var.get()))
        except ValueError as e:
            messagebox.showwarning("Rules Overlap", str(e))
            return False
        return True

    def _on_search_settings_changed(self, *_args):
        """
        Discard on-screen results when the search settings change.

        The results describe offsets found with one folder, rule set and
        matching options. Applying them under different settings would
        validate old offsets against new rules, so the preview is invalidated
        instead.
        """
        if self.is_processing or not self.all_matches:
            return
        self._invalidate_results()

    def _invalidate_results(self):
        self._clear_results()
        self.progress_var.set("Search settings changed - run Preview Changes again")

    def _clear_results(self):
        """Drop the current results and the tree showing them."""
        self.tree.delete(*self.tree.get_children())
        self.match_tree_map.clear()
        self.file_tree_map.clear()
        self.all_matches.clear()
        self.scan_results = []
        self.apply_btn.configure(state=tk.DISABLED)
        self.match_count_label.configure(text="")

    # ------------------------------------------------------------------
    # Rule set editing
    # ------------------------------------------------------------------

    def _set_rules(self, rules: List[Rule], invalidate: bool = True):
        """Replace the rule set and redraw its list."""
        self.rules = list(rules)
        self.rules_tree.delete(*self.rules_tree.get_children())
        self.rule_iids = [
            self.rules_tree.insert("", tk.END, values=(r.find_text, r.replace_text))
            for r in self.rules
        ]
        if invalidate and self.all_matches and not self.is_processing:
            self._invalidate_results()

    def _selected_rule_index(self) -> Optional[int]:
        """Index of the one rule selected in the list, if exactly one is."""
        selected = self.rules_tree.selection()
        if len(selected) == 1 and selected[0] in self.rule_iids:
            return self.rule_iids.index(selected[0])
        return None

    def _add_rule(self):
        """
        Add the editor's rule to the set, or update an existing rule.

        The rule selected in the list is the one being edited, so it is
        replaced whichever field changed - editing its find text must not
        leave the old rule behind to take part in the next preview. With no
        selection, a rule with the same find text is updated in place and
        anything else is appended.
        """
        if self.is_processing:
            return
        rule = self._pending_rule()
        if rule is None:
            messagebox.showwarning("Missing Input", "Please enter text to find.")
            return
        rules = list(self.rules)
        target = self._selected_rule_index()
        if target is None:
            for index, existing in enumerate(rules):
                if existing.find_text == rule.find_text:
                    target = index
                    break
        else:
            clash = next((i for i, existing in enumerate(rules)
                          if i != target and existing.find_text == rule.find_text), None)
            if clash is not None:
                messagebox.showwarning(
                    "Duplicate Rule",
                    f'"{rule.find_text}" is already rule {clash + 1}. Remove one of '
                    f'them or choose a different find text.')
                return

        if target is None:
            rules.append(rule)
        elif rules[target] == rule:
            self._clear_editor()
            return
        else:
            rules[target] = rule
        self._set_rules(rules)
        self._clear_editor()

    def _clear_editor(self):
        self.find_var.set("")
        self.replace_var.set("")
        selected = self.rules_tree.selection()
        if selected:
            self.rules_tree.selection_remove(*selected)

    def _remove_rules(self):
        """Remove the rules selected in the list."""
        if self.is_processing:
            return
        selected = set(self.rules_tree.selection())
        if not selected:
            messagebox.showinfo("Nothing Selected", "Select a rule in the list to remove it.")
            return
        rules = [r for r, iid in zip(self.rules, self.rule_iids) if iid not in selected]
        self._set_rules(rules)
        self._clear_editor()

    def _on_rule_selected(self, _event=None):
        """Load the selected rule into the editor so it can be changed."""
        selected = self.rules_tree.selection()
        if len(selected) != 1 or selected[0] not in self.rule_iids:
            return
        rule = self.rules[self.rule_iids.index(selected[0])]
        self.find_var.set(rule.find_text)
        self.replace_var.set(rule.replace_text)

    def _load_rule_set(self):
        """Replace the rule set with one read from a CSV file."""
        if self.is_processing:
            return
        path = filedialog.askopenfilename(
            title="Load rule set",
            filetypes=[("Rule sets (CSV)", "*.csv"), ("All files", "*.*")])
        if not path:
            return
        try:
            rules = app_settings.load_rules_csv(path)
        except ValueError as e:
            messagebox.showwarning("Rule Set Not Loaded", f"{os.path.basename(path)}:\n{e}")
            return
        except OSError as e:
            messagebox.showerror("Rule Set Not Loaded", f"Could not read the file:\n{e}")
            return
        self._set_rules(rules)
        self._clear_editor()
        self.progress_var.set(
            f"Loaded {_plural(len(rules), 'rule')} from {os.path.basename(path)}")

    def _save_rule_set(self):
        """Write the rule set (including a pending editor row) to a CSV file."""
        if self.is_processing:
            return
        if self._pending_rule() is not None:
            self._add_rule()
        if not self.rules:
            messagebox.showinfo("Nothing to Save", "There are no rules to save.")
            return
        path = filedialog.asksaveasfilename(
            title="Save rule set", defaultextension=".csv",
            filetypes=[("Rule sets (CSV)", "*.csv"), ("All files", "*.*")])
        if not path:
            return
        try:
            app_settings.save_rules_csv(path, self.rules)
        except OSError as e:
            messagebox.showerror("Rule Set Not Saved", f"Could not write the file:\n{e}")
            return
        self.progress_var.set(
            f"Saved {_plural(len(self.rules), 'rule')} to {os.path.basename(path)}")

    # ------------------------------------------------------------------
    # Preview
    # ------------------------------------------------------------------

    def _start_preview(self):
        """Start the preview scan in a background thread."""
        if not self._validate_inputs():
            return

        self.processing_kind = 'scan'
        self.scan_cancel = threading.Event()

        # Record what this scan is being run with; apply uses these, not the
        # live widgets.
        self.scan_folder = (self.folder_var.get() or "").strip()
        self.scan_rules = list(self.rules)
        self.scan_case_sensitive = bool(self.case_var.get())
        self.scan_whole_word = bool(self.whole_word_var.get())
        self.scan_recursive = bool(self.recursive_var.get())

        self._clear_results()
        self.progress_var.set("Scanning documents...")
        self.preview_btn.configure(text="Cancel Scan", command=self._cancel_scan)
        self._set_controls_enabled(False)

        thread = threading.Thread(target=self._run_preview, daemon=True)
        thread.start()

    def _cancel_scan(self):
        """Ask the running scan to stop after the file it is on."""
        if self.scan_cancel is not None:
            self.scan_cancel.set()
            self.progress_var.set("Cancelling scan...")
            self.preview_btn.configure(state=tk.DISABLED)

    def _run_preview(self):
        """Background thread: scan documents."""
        folder = self.scan_folder
        rules = self.scan_rules
        case_sensitive = self.scan_case_sensitive
        whole_word = self.scan_whole_word
        recursive = self.scan_recursive
        cancel = self.scan_cancel

        def progress(file_name, idx, total):
            self.root.after(0, self.progress_var.set,
                            f"Scanning {file_name} ({idx + 1}/{total})...")

        try:
            results, summary = scan_documents_detailed(
                folder, rules, case_sensitive,
                whole_word=whole_word, recursive=recursive,
                progress_callback=progress, cancel_event=cancel,
            )
            self.root.after(0, self._display_results, results, summary)
        except Exception as e:
            self.root.after(0, self._preview_error, str(e))

    def _preview_error(self, error_msg: str):
        """Handle preview errors on main thread."""
        self.progress_var.set("Error during scan")
        messagebox.showerror("Scan Error", f"An error occurred:\n{error_msg}")
        self._finish_processing()

    def _display_results(self, results: List[FileResult], summary: dict):
        """Populate the treeview with scan results (main thread)."""
        self.scan_results = results
        self.all_matches.clear()

        total_matches = 0
        total_files = 0
        total_conflicts = 0

        for fr in results:
            # File-level node
            if fr.error:
                label = f"{WARNING} {fr.file_name} - ERROR: {fr.error}"
                file_node = self.tree.insert("", tk.END, text=label, values=("", ""),
                                             tags=("error",))
            else:
                file_node = self.tree.insert("", tk.END, text="", values=("", ""))
                self.file_tree_map[file_node] = fr
                total_files += 1

                for match in fr.matches:
                    self.all_matches.append(match)
                    total_matches += 1
                    if match.conflict:
                        total_conflicts += 1
                    match_node = self.tree.insert(
                        file_node, tk.END, text=self._match_label(match),
                        values=(_rule_text(match.effective_find_text, match.replace_text),
                                match.section),
                        tags=("conflict",) if match.conflict else (),
                    )
                    self.match_tree_map[match_node] = match
                self._refresh_file_label(file_node)

            # Expand file nodes
            self.tree.item(file_node, open=True)

        scanned = summary.get('files_scanned', 0)
        cancelled = summary.get('cancelled', False)

        if total_matches > 0:
            self.apply_btn.configure(state=tk.NORMAL)
            status = (f"Found {_plural(total_matches, 'match', 'matches')} "
                      f"in {_plural(total_files, 'file')}")
            if total_conflicts:
                status += (f" - {total_conflicts} overlapping between rules, "
                           f"not selectable ({WARNING})")
            if cancelled:
                status += f" - scan cancelled after {_plural(scanned, 'file')}"
            self.progress_var.set(status)
        elif cancelled:
            self.progress_var.set(f"Scan cancelled after {_plural(scanned, 'file')}")
        elif scanned == 0:
            self.progress_var.set("No .docx files found in the selected folder")
        else:
            self.progress_var.set(
                f"No matches found in {scanned} .docx file{'s' if scanned != 1 else ''}"
            )

        # Last: this may destroy the window if a close was requested mid-scan.
        self._finish_processing()

    # ------------------------------------------------------------------
    # Results tree
    # ------------------------------------------------------------------

    def _match_label(self, match: Match) -> str:
        """The tree row for one match: checkbox, location, context."""
        if match.conflict:
            check = WARNING
        else:
            check = CHECKED if match.is_selected else UNCHECKED
        prefix = LOCATION_PREFIXES.get(match.location_type, "")
        return f"{check} {prefix}{match.location_detail}: {match.display_context}"

    def _on_tree_click(self, event):
        """Handle clicks to toggle checkboxes."""
        # The worker thread reads the same Match objects, so the tree is inert
        # while a scan or apply is running.
        if self.is_processing:
            return "break"

        # Clicking the expand/collapse arrow must expand, not toggle every
        # checkbox in the file.
        if self.tree.identify_element(event.x, event.y) == "Treeitem.indicator":
            return

        item = self.tree.identify_row(event.y)
        if not item:
            return

        if item in self.match_tree_map:
            match = self.match_tree_map[item]
            if match.conflict:
                # Overlapping rules: the user resolves this by changing the
                # rules, not by ticking a box.
                return
            match.is_selected = not match.is_selected
            self._refresh_item_label(item, match)
            parent = self.tree.parent(item)
            if parent in self.file_tree_map:
                self._refresh_file_label(parent)

        elif item in self.file_tree_map:
            fr = self.file_tree_map[item]
            selectable = [m for m in fr.matches if not m.conflict]
            # Toggle: if all selected, deselect all; otherwise select all
            new_state = not all(m.is_selected for m in selectable)

            for m in selectable:
                m.is_selected = new_state

            for child in self.tree.get_children(item):
                if child in self.match_tree_map:
                    self._refresh_item_label(child, self.match_tree_map[child])

            self._refresh_file_label(item)

        self._update_match_count()

    def _on_tree_double_click(self, event):
        """Open the document behind the row in Word (or the platform default)."""
        if self.is_processing:
            return "break"
        item = self.tree.identify_row(event.y)
        if not item:
            return
        if item in self.match_tree_map:
            path = self.match_tree_map[item].file_path
        elif item in self.file_tree_map:
            path = self.file_tree_map[item].file_path
        else:
            return
        try:
            _open_document(path)
        except Exception as e:
            messagebox.showerror("Could Not Open", f"{os.path.basename(path)}:\n{e}")
        return "break"

    def _refresh_item_label(self, item: str, match: Match):
        """Update a match item's label to reflect its checkbox state."""
        self.tree.item(item, text=self._match_label(match))

    def _refresh_file_label(self, item: str):
        """Update a file item's label to reflect its children's state."""
        fr = self.file_tree_map[item]
        selected = fr.selected_count
        selectable = fr.selectable_count
        total = fr.match_count

        if selectable and selected == selectable:
            check = CHECKED
        elif selected == 0:
            check = UNCHECKED
        else:
            check = PARTIAL

        label = f"{check} {fr.file_name} ({_plural(total, 'match', 'matches')}, {selected} selected"
        if fr.conflict_count:
            label += f", {fr.conflict_count} overlapping"
        label += ")"
        self.tree.item(item, text=label)

    def _update_match_count(self):
        """Update the match count label."""
        selectable = sum(1 for m in self.all_matches if not m.conflict)
        selected = sum(1 for m in self.all_matches if m.is_selected and not m.conflict)
        conflicts = len(self.all_matches) - selectable
        text = f"{selected} of {selectable} selected"
        if conflicts:
            text += f" ({conflicts} overlapping, not selectable)"
        self.match_count_label.configure(text=text if self.all_matches else "")

        if self.is_processing:
            return
        self.apply_btn.configure(state=tk.NORMAL if selected > 0 else tk.DISABLED)

    def _select_all(self):
        """Select all matches (overlapping ones stay unselectable)."""
        if self.is_processing:
            return
        for m in self.all_matches:
            if not m.conflict:
                m.is_selected = True
        self._refresh_all_labels()
        self._update_match_count()

    def _deselect_all(self):
        """Deselect all matches."""
        if self.is_processing:
            return
        for m in self.all_matches:
            m.is_selected = False
        self._refresh_all_labels()
        self._update_match_count()

    def _refresh_all_labels(self):
        """Refresh all tree labels."""
        for item, match in self.match_tree_map.items():
            self._refresh_item_label(item, match)
        for item in self.file_tree_map:
            self._refresh_file_label(item)

    # ------------------------------------------------------------------
    # Apply
    # ------------------------------------------------------------------

    def _set_controls_enabled(self, enabled: bool):
        """Enable or disable everything that must not be touched mid-run."""
        state = tk.NORMAL if enabled else tk.DISABLED
        for widget in (self.apply_btn, self.browse_btn,
                       self.select_all_btn, self.deselect_all_btn, self.restore_btn,
                       self.folder_entry, self.find_entry, self.replace_entry,
                       self.add_rule_btn, self.remove_rule_btn,
                       self.load_set_btn, self.save_set_btn,
                       self.recursive_check, self.case_check, self.whole_word_check,
                       self.backup_check, self.changelog_check, self.tracked_check,
                       self.author_entry):
            widget.configure(state=state)

    def _finish_processing(self):
        """Return the UI to its idle state after a scan or apply."""
        self.processing_kind = None
        self.scan_cancel = None
        self.preview_btn.configure(text="Preview Changes",
                                   command=self._start_preview, state=tk.NORMAL)
        self._set_controls_enabled(True)
        self._update_match_count()

        # The user asked to quit while a scan was running; the scan has now
        # stopped, so honour it.
        if self.close_when_idle:
            self._save_settings()
            self.root.destroy()

    def _apply_options(self) -> dict:
        """Snapshot everything the apply worker needs; it never reads a widget."""
        author = (self.author_var.get() or "").strip()
        return {
            'create_backups': bool(self.backup_var.get()),
            'write_change_log': bool(self.changelog_var.get()),
            'tracked_changes': bool(self.tracked_var.get()),
            'author': author or None,
            # The settings the offsets were found under, not the live widgets.
            'case_sensitive': self.scan_case_sensitive,
            'folder': self.scan_folder,
        }

    def _start_apply(self):
        """Start applying changes in a background thread."""
        selected = [m for m in self.all_matches if m.is_selected and not m.conflict]
        file_count = len(set(m.file_path for m in selected))

        if not selected:
            messagebox.showinfo("Nothing to Apply", "No changes are selected.")
            return

        options = self._apply_options()

        # The rules actually scanned, which is what these offsets belong to.
        rules_lines = [f'  "{r.find_text}" {ARROW} "{r.replace_text}"'
                       for r in self.scan_rules[:8]]
        if len(self.scan_rules) > 8:
            rules_lines.append(f"  ... and {len(self.scan_rules) - 8} more")

        if options['tracked_changes']:
            mode = f"Mode: write as tracked changes by {options['author'] or 'the current user'}"
        else:
            mode = "Mode: edit documents in place"

        lines = [
            f"Apply {_plural(len(selected), 'replacement')} "
            f"across {_plural(file_count, 'file')}?",
            "",
            f"Rules ({len(self.scan_rules)}):",
            *rules_lines,
            "",
            f"Case sensitive: {'yes' if self.scan_case_sensitive else 'no'}   "
            f"Whole word: {'yes' if self.scan_whole_word else 'no'}",
            mode,
            "Backups will be created." if options['create_backups']
            else "NO backups will be created!",
        ]
        if options['write_change_log']:
            lines.append(f"A change log (CSV) will be written to {options['folder']}")
        pending = self._pending_rule()
        if pending is not None:
            lines += ["", f"{WARNING} The Find box holds \"{pending.find_text}\", which "
                          f"was not part of this preview. Add it and run Preview "
                          f"Changes again to include it."]

        if not messagebox.askyesno("Confirm Changes", "\n".join(lines)):
            return

        self.processing_kind = 'apply'
        self.preview_btn.configure(state=tk.DISABLED)
        self._set_controls_enabled(False)

        thread = threading.Thread(target=self._run_apply, args=(options,), daemon=True)
        thread.start()

    def _run_apply(self, options: dict):
        """Background thread: apply changes, then write the change log."""
        def progress(file_name, idx, total):
            self.root.after(0, self.progress_var.set,
                            f"Processing {file_name} ({idx + 1}/{total})...")

        try:
            result = apply_changes(
                self.all_matches,
                create_backups=options['create_backups'],
                case_sensitive=options['case_sensitive'],
                tracked_changes=options['tracked_changes'],
                author=options['author'],
                progress_callback=progress,
            )
            if options['write_change_log']:
                stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
                log_path = os.path.join(options['folder'], f"{CHANGE_LOG_PREFIX}{stamp}.csv")
                try:
                    write_change_log(log_path, self.all_matches, result.get('backup_map'))
                    result['change_log'] = log_path
                except Exception as e:
                    result.setdefault('errors', []).append(
                        f"Could not write the change log: {e}")
            self.root.after(0, self._apply_complete, result)
        except Exception as e:
            self.root.after(0, self._apply_error, str(e))

    def _apply_complete(self, result: dict):
        """Handle apply completion on main thread."""
        self.processing_kind = None
        skipped = result.get('total_skipped', 0)

        msg = (f"Completed!\n\n"
               f"Replacements made: {result['total_replaced']}\n"
               f"Files modified: {result['files_modified']}")

        if result.get('backups'):
            msg += f"\nBackups created: {len(result['backups'])}"

        if result.get('change_log'):
            msg += f"\nChange log: {result['change_log']}"

        if skipped:
            # A selected match that could not be applied must never disappear
            # quietly - the preview promised it would be changed.
            msg += (f"\n\n{WARNING} {skipped} selected "
                    f"{'change was' if skipped == 1 else 'changes were'} NOT applied.")
            reasons = result.get('skip_reasons') or {}
            if reasons:
                for reason, count in sorted(reasons.items(), key=lambda kv: -kv[1]):
                    msg += f"\n  - {count}: {reason}"
            else:
                msg += ("\nThe document may have been edited since the preview. "
                        "Re-run Preview Changes to see the current state.")

        if result.get('errors'):
            msg += f"\n\nErrors ({len(result['errors'])}):\n"
            for err in result['errors']:
                msg += f"  - {err}\n"

        status = (f"Done: {result['total_replaced']} replacements "
                  f"in {result['files_modified']} files")
        if skipped:
            status += f" ({skipped} skipped)"

        # Results describe the pre-edit documents, so they are stale now.
        self._clear_results()
        self.progress_var.set(status)
        self._finish_processing()

        messagebox.showinfo("Changes Applied", msg)

    def _apply_error(self, error_msg: str):
        """Handle apply errors on main thread."""
        # Keep the results on screen so the run can be retried without
        # re-scanning.
        self.progress_var.set("Error during apply")
        self._finish_processing()
        messagebox.showerror("Apply Error", f"An error occurred:\n{error_msg}")

    # ------------------------------------------------------------------
    # Restore from backup
    # ------------------------------------------------------------------

    def _open_restore_dialog(self):
        """List the backups under the folder and offer to put them back."""
        if self.is_processing:
            return
        folder = (self.folder_var.get() or "").strip()
        if not folder or not os.path.isdir(folder):
            messagebox.showwarning("Invalid Folder", "Please select a folder first.")
            return
        try:
            backups = find_backups(folder, recursive=bool(self.recursive_var.get()))
        except OSError as e:
            messagebox.showerror("Could Not List Backups", str(e))
            return
        if not backups:
            messagebox.showinfo("No Backups",
                                "No timestamped backups were found in this folder.")
            return
        RestoreDialog(self, folder, backups)

    def _after_restore(self, result: dict):
        """Report a restore and drop results that no longer describe the files."""
        self._clear_results()
        restored = result.get('restored', 0)
        self.progress_var.set(
            f"Restored {_plural(restored, 'document')} - run Preview Changes again")
        if result.get('errors'):
            messagebox.showerror(
                "Restore Errors",
                f"{_plural(restored, 'document')} restored.\n\n"
                + "\n".join(f"- {e}" for e in result['errors']))

    # ------------------------------------------------------------------
    # Shutdown
    # ------------------------------------------------------------------

    def _on_close(self):
        """Refuse to close while a worker thread is mid-run."""
        if self.processing_kind == 'apply':
            messagebox.showwarning(
                "Apply in Progress",
                "Documents are being written right now.\n\n"
                "Closing would interrupt the run and leave part of the batch "
                "unmodified. Please wait for it to finish."
            )
            return

        if self.processing_kind == 'scan':
            if not messagebox.askyesno("Scan in Progress",
                                       "A scan is still running. Cancel it and quit?"):
                return
            self.close_when_idle = True
            self._cancel_scan()
            return

        self._save_settings()
        self.root.destroy()


class RestoreDialog:
    """
    A window listing every backup under the folder, newest first per
    document, with the newest one of each pre-selected.

    Restoring copies the backup back over the document and leaves the backup
    file in place. Only one backup per document can be chosen.
    """

    def __init__(self, app: FindReplaceApp, folder: str, backups: List[Backup]):
        self.app = app
        self.folder = folder
        self.backups = backups
        self.checked: Dict[str, bool] = {}
        self.backup_by_iid: Dict[str, Backup] = {}

        self.top = tk.Toplevel(app.root)
        self.top.title("Restore backups")
        self.top.geometry("760x420")
        self.top.minsize(520, 300)
        self.top.transient(app.root)

        frame = ttk.Frame(self.top, padding=10)
        frame.pack(fill=tk.BOTH, expand=True)
        ttk.Label(frame, text=(
            "Tick the backups to put back. The newest backup of each document is "
            "ticked; a document can be restored from one backup at a time. Backup "
            "files are kept."), wraplength=700).pack(fill=tk.X, pady=(0, 6))

        tree_frame = ttk.Frame(frame)
        tree_frame.pack(fill=tk.BOTH, expand=True)
        self.tree = ttk.Treeview(tree_frame, columns=("when",), show="tree headings",
                                 selectmode="none")
        self.tree.heading("#0", text="Document")
        self.tree.heading("when", text="Backup taken")
        self.tree.column("#0", width=480, minwidth=200)
        self.tree.column("when", width=180, minwidth=120)
        scroll = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        scroll.grid(row=0, column=1, sticky="ns")
        tree_frame.columnconfigure(0, weight=1)
        tree_frame.rowconfigure(0, weight=1)
        self.tree.bind("<Button-1>", self._on_click)

        newest_seen = set()
        for backup in backups:
            key = os.path.normcase(backup.document_path)
            checked = key not in newest_seen
            newest_seen.add(key)
            try:
                name = os.path.relpath(backup.document_path, folder)
            except ValueError:
                name = backup.document_path
            when = backup.timestamp.strftime("%Y-%m-%d %H:%M:%S")
            if backup.sequence:
                when += f" (#{backup.sequence})"
            iid = self.tree.insert("", tk.END, text="", values=(when,))
            self.backup_by_iid[iid] = backup
            self.checked[iid] = checked
            self._refresh(iid, name)

        buttons = ttk.Frame(frame)
        buttons.pack(fill=tk.X, pady=(8, 0))
        self.restore_btn = ttk.Button(buttons, text="Restore Selected", command=self._restore)
        self.restore_btn.pack(side=tk.RIGHT)
        ttk.Button(buttons, text="Close", command=self.top.destroy).pack(
            side=tk.RIGHT, padx=(0, 6))

    def _name_of(self, iid: str) -> str:
        backup = self.backup_by_iid[iid]
        try:
            return os.path.relpath(backup.document_path, self.folder)
        except ValueError:
            return backup.document_path

    def _refresh(self, iid: str, name: Optional[str] = None):
        name = name if name is not None else self._name_of(iid)
        check = CHECKED if self.checked[iid] else UNCHECKED
        self.tree.item(iid, text=f"{check} {name}")

    def _on_click(self, event):
        iid = self.tree.identify_row(event.y)
        if not iid or iid not in self.backup_by_iid:
            return
        self.toggle(iid)
        return "break"

    def toggle(self, iid: str):
        """Tick a backup; ticking one unticks any other backup of that document."""
        new_state = not self.checked[iid]
        self.checked[iid] = new_state
        self._refresh(iid)
        if new_state:
            doc = os.path.normcase(self.backup_by_iid[iid].document_path)
            for other, backup in self.backup_by_iid.items():
                if other != iid and os.path.normcase(backup.document_path) == doc \
                        and self.checked[other]:
                    self.checked[other] = False
                    self._refresh(other)

    def selected_backups(self) -> List[Backup]:
        return [self.backup_by_iid[iid] for iid in self.backup_by_iid if self.checked[iid]]

    def _restore(self):
        chosen = self.selected_backups()
        if not chosen:
            messagebox.showinfo("Nothing Selected", "Tick at least one backup to restore.")
            return
        if not messagebox.askyesno(
                "Confirm Restore",
                f"Put back {_plural(len(chosen), 'document')} from "
                f"{'its' if len(chosen) == 1 else 'their'} backup?\n\n"
                "Any edits made to these documents since the backup was taken "
                "will be lost. The backup files are kept."):
            return
        try:
            result = restore_backups(chosen)
        except ValueError as e:
            messagebox.showerror("Restore Refused", str(e))
            return
        self.top.destroy()
        self.app._after_restore(result)


def selftest(report_path: Optional[str] = None) -> int:
    """
    Prove a packaged build works end to end, without opening a window.

    Run as `DocxFindReplace.exe --selftest`. A frozen build fails in ways the
    source never does - python-docx's XML templates left out of the bundle,
    lxml's dynamically resolved imports not collected - and those only surface
    once the document code actually runs. The CI build runs this against the
    real executable so a broken bundle cannot be shipped.

    Covers the in-place engine, the tracked-changes engine, a rule-set CSV
    round trip and the settings file, because each is a separate path the
    frozen build has to carry.

    Writes a report to `report_path` when given, because a windowed executable
    has no console to print to.
    """
    import tempfile
    import traceback

    lines = [f"{APP_NAME} {__version__}", f"frozen: {getattr(sys, 'frozen', False)}"]
    ok = True

    try:
        from docx import Document
        from docx.oxml.ns import qn
        import document_processor as dp

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "selftest.docx")
            # Document() with no argument loads python-docx's default template,
            # which is exactly the bundled data the frozen build tends to miss.
            doc = Document()
            doc.add_paragraph("Reference the 2022 CBC for compliance.")
            doc.save(path)

            matches = [m for r in dp.scan_documents(tmp, "2022 CBC", True)
                       for m in r.matches]
            if len(matches) != 1:
                raise AssertionError(f"expected 1 match, got {len(matches)}")

            result = dp.apply_changes(matches, "2025 CBC", create_backups=False)
            if result["total_replaced"] != 1:
                raise AssertionError(f"expected 1 replacement, got {result}")

            reloaded = dp._collect_paragraphs(Document(path))[0][0]
            text = dp.paragraph_text(reloaded)
            if text != "Reference the 2025 CBC for compliance.":
                raise AssertionError(f"unexpected text: {text!r}")
            lines.append(f"document engine OK -> {text!r}")

            # Tracked-changes engine on a second document.
            tracked_dir = os.path.join(tmp, "tracked")
            os.mkdir(tracked_dir)
            tracked_path = os.path.join(tracked_dir, "tracked.docx")
            doc = Document()
            doc.add_paragraph("Per NFPA 13, 2022 edition.")
            doc.save(tracked_path)
            rules = [dp.Rule("NFPA 13, 2022 edition", "NFPA 13, 2025 edition")]
            matches = [m for r in dp.scan_documents(tracked_dir, rules) for m in r.matches]
            result = dp.apply_changes(matches, create_backups=False,
                                      tracked_changes=True, author="selftest")
            if result["total_replaced"] != 1:
                raise AssertionError(f"expected 1 tracked replacement, got {result}")
            body = Document(tracked_path).element.body
            dels = body.findall(".//" + qn("w:del"))
            ins = body.findall(".//" + qn("w:ins"))
            if len(dels) != 1 or len(ins) != 1:
                raise AssertionError(f"expected 1 w:del and 1 w:ins, got {len(dels)}/{len(ins)}")
            text = dp.paragraph_text(dp._collect_paragraphs(Document(tracked_path))[0][0])
            if text != "Per NFPA 13, 2025 edition.":
                raise AssertionError(f"unexpected tracked text: {text!r}")
            lines.append(f"tracked changes OK -> {text!r}")

            # Rule-set CSV round trip and the settings file.
            csv_path = os.path.join(tmp, "rules.csv")
            app_settings.save_rules_csv(csv_path, rules)
            if app_settings.load_rules_csv(csv_path) != rules:
                raise AssertionError("rule set did not round-trip")
            settings_path = os.path.join(tmp, "settings.json")
            app_settings.save_settings({"rules": rules, "author": "selftest"}, settings_path)
            loaded = app_settings.load_settings(settings_path)
            if loaded["rules"] != rules or loaded["author"] != "selftest":
                raise AssertionError("settings did not round-trip")
            lines.append("rule set and settings OK")

        lines.append(f"tkinter OK -> Tk {tk.TkVersion}")
        lines.append("SELFTEST PASSED")
    except Exception:
        ok = False
        lines.append("SELFTEST FAILED")
        lines.append(traceback.format_exc())

    report = "\n".join(lines)
    print(report)

    if report_path:
        try:
            with open(report_path, "w", encoding="utf-8") as fh:
                fh.write(report + "\n")
        except OSError:
            pass

    return 0 if ok else 1


def main():
    if "--selftest" in sys.argv:
        after = sys.argv[sys.argv.index("--selftest") + 1:]
        sys.exit(selftest(after[0] if after else None))

    # DPI awareness has to be set before the first Tk window exists, otherwise
    # Tk has already sampled the old DPI and the UI stays blurry on high-DPI
    # displays.
    scaling = None
    try:
        from ctypes import windll
        windll.shcore.SetProcessDpiAwareness(1)
        scaling = windll.user32.GetDpiForSystem() / 72.0
    except Exception:
        pass

    root = tk.Tk()

    if scaling:
        try:
            root.tk.call("tk", "scaling", scaling)
        except Exception:
            pass

    FindReplaceApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
