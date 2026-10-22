"""
Shared fixtures for the document_processor test suite.

Each builder constructs a real .docx on disk exercising one structure that the
scan/replace code has to get right: hyperlinks, merged cells, linked headers,
inline images, nested tables, text boxes and tracked insertions, plus the
elements that are not text but sit at a place in it (pictures, note and
comment reference marks, symbols, field characters, bookmarks), and complex
fields nested in each other or running across paragraphs.
"""

import base64
import os
import sys
from typing import NamedTuple

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# Smallest valid PNG, used wherever a test needs a real inline image.
_PNG_1PX = base64.b64decode(
    b"iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


# ---------------------------------------------------------------- helpers

def add_run(paragraph, text):
    """Append a plain <w:r><w:t> run, bypassing python-docx's text setter."""
    xml = (
        '<w:r xmlns:w="%s"><w:t xml:space="preserve">%s</w:t></w:r>'
        % (NS_W, _escape(text))
    )
    paragraph._p.append(parse_xml(xml))


def add_runs(paragraph, *texts):
    """Append several runs, so a phrase can be split across run boundaries."""
    for text in texts:
        add_run(paragraph, text)


def add_hyperlink(paragraph, text, url="https://example.com/spec"):
    """Append a real external hyperlink containing `text`."""
    r_id = paragraph.part.relate_to(
        url, "%s/hyperlink" % NS_R, is_external=True
    )
    xml = (
        '<w:hyperlink xmlns:w="%s" xmlns:r="%s" r:id="%s">'
        '<w:r><w:t xml:space="preserve">%s</w:t></w:r></w:hyperlink>'
        % (NS_W, NS_R, r_id, _escape(text))
    )
    paragraph._p.append(parse_xml(xml))


def add_tracked_insertion(paragraph, text, author="Reviewer"):
    """Append a <w:ins> tracked insertion containing `text`."""
    xml = (
        '<w:ins xmlns:w="%s" w:id="900" w:author="%s" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:t xml:space="preserve">%s</w:t></w:r></w:ins>'
        % (NS_W, author, _escape(text))
    )
    paragraph._p.append(parse_xml(xml))


def add_deletion(paragraph, text, author="Reviewer"):
    """Append a <w:del> tracked deletion (uses w:delText, must stay untouched)."""
    xml = (
        '<w:del xmlns:w="%s" w:id="901" w:author="%s" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:delText xml:space="preserve">%s</w:delText></w:r></w:del>'
        % (NS_W, author, _escape(text))
    )
    paragraph._p.append(parse_xml(xml))


def add_formatted_run(paragraph, text, bold=False, color=None):
    """Append a run carrying explicit formatting (w:rPr with w:b / w:color)."""
    props = ""
    if bold:
        props += "<w:b/>"
    if color:
        props += '<w:color w:val="%s"/>' % color
    xml = (
        '<w:r xmlns:w="%s"><w:rPr>%s</w:rPr>'
        '<w:t xml:space="preserve">%s</w:t></w:r>'
        % (NS_W, props, _escape(text))
    )
    paragraph._p.append(parse_xml(xml))


def add_run_with_rpr_change(paragraph, text, change_id=902):
    """Append a bold run whose formatting is itself a tracked change."""
    xml = (
        '<w:r xmlns:w="%s"><w:rPr><w:b/>'
        '<w:rPrChange w:id="%d" w:author="Reviewer" w:date="2026-01-01T00:00:00Z">'
        '<w:rPr/></w:rPrChange></w:rPr>'
        '<w:t xml:space="preserve">%s</w:t></w:r>'
        % (NS_W, change_id, _escape(text))
    )
    paragraph._p.append(parse_xml(xml))


def add_tab_run(paragraph):
    """Append a run holding only a <w:tab/>."""
    paragraph._p.append(parse_xml('<w:r xmlns:w="%s"><w:tab/></w:r>' % NS_W))


def add_picture_run(paragraph, png_path, before="", after=""):
    """
    Append one run holding `before` text, a real inline picture, then `after`
    text: ``<w:r><w:t>before</w:t><w:drawing/><w:t>after</w:t></w:r>``.
    """
    run = paragraph.add_run()
    if before:
        run.add_text(before)
    run.add_picture(png_path, width=Inches(0.2))
    if after:
        run.add_text(after)
    return run


def add_run_with_element(paragraph, before, element, after):
    """
    Append one run holding `before` text, a single empty w: element, then
    `after` text. `element` is the tag and its attributes without the prefix,
    e.g. ``'sym w:font="Symbol" w:char="F0B0"'`` or ``'lastRenderedPageBreak'``.
    """
    inner = ""
    if before:
        inner += '<w:t xml:space="preserve">%s</w:t>' % _escape(before)
    inner += "<w:%s/>" % element
    if after:
        inner += '<w:t xml:space="preserve">%s</w:t>' % _escape(after)
    paragraph._p.append(parse_xml('<w:r xmlns:w="%s">%s</w:r>' % (NS_W, inner)))


def add_note_reference(paragraph, kind, note_id):
    """
    Append the run Word writes for a note or comment reference mark. `kind` is
    "footnote", "endnote" or "comment"; `note_id` should name an entry that
    `add_notes_part` created (its first entry is id 2).
    """
    style = {"footnote": "FootnoteReference", "endnote": "EndnoteReference",
             "comment": "CommentReference"}[kind]
    xml = ('<w:r xmlns:w="%s"><w:rPr><w:rStyle w:val="%s"/></w:rPr>'
           '<w:%sReference w:id="%d"/></w:r>' % (NS_W, style, kind, note_id))
    paragraph._p.append(parse_xml(xml))


def add_bookmark_start(paragraph, bookmark_id, name):
    """Append a <w:bookmarkStart> among the paragraph's runs."""
    paragraph._p.append(parse_xml(
        '<w:bookmarkStart xmlns:w="%s" w:id="%d" w:name="%s"/>'
        % (NS_W, bookmark_id, _escape(name))))


def add_bookmark_end(paragraph, bookmark_id):
    """Append the <w:bookmarkEnd> closing `bookmark_id`."""
    paragraph._p.append(parse_xml(
        '<w:bookmarkEnd xmlns:w="%s" w:id="%d"/>' % (NS_W, bookmark_id)))


def add_paragraph_with_outline_level(doc, text, level):
    """A Normal-styled paragraph carrying its own w:outlineLvl."""
    p = doc.add_paragraph(text)
    ppr = p._p.get_or_add_pPr()
    ppr.append(parse_xml('<w:outlineLvl xmlns:w="%s" w:val="%d"/>' % (NS_W, level)))
    return p


def add_custom_paragraph_style(doc, style_id, name, based_on=None, outline_level=None):
    """Append a paragraph style to the styles part. Test code may use
    doc.styles; only the scan path must not."""
    parts = ['<w:style xmlns:w="%s" w:type="paragraph" w:styleId="%s">' % (NS_W, style_id),
             '<w:name w:val="%s"/>' % _escape(name)]
    if based_on:
        parts.append('<w:basedOn w:val="%s"/>' % based_on)
    if outline_level is not None:
        parts.append('<w:pPr><w:outlineLvl w:val="%d"/></w:pPr>' % outline_level)
    parts.append('</w:style>')
    doc.styles.element.append(parse_xml("".join(parts)))


def add_styled_paragraph(doc, text, style_id):
    """A paragraph whose w:pStyle names `style_id` directly (no style lookup)."""
    p = doc.add_paragraph(text)
    ppr = p._p.get_or_add_pPr()
    ppr.insert(0, parse_xml('<w:pStyle xmlns:w="%s" w:val="%s"/>' % (NS_W, style_id)))
    return p


def drop_styles_part(doc):
    """Detach the styles part so the saved package has none."""
    for r_id, rel in list(doc.part.rels.items()):
        if rel.reltype.endswith("/styles"):
            doc.part.drop_rel(r_id)


def has_styles_part(path):
    from zipfile import ZipFile
    with ZipFile(path) as z:
        return "word/styles.xml" in z.namelist()


class NestedField(NamedTuple):
    instruction: object
    result: object = None


def nested_field(instruction, result=None):
    """
    A whole field to place inside another field's code or result, as one
    piece of the `instruction` or `cached_result` of `add_complex_field` (or of
    `add_field_begin`, `add_field_code`, `add_field_separate`). Its own
    instruction and result take pieces the same way, so fields nest to any
    depth. `result=None` gives a field with no separate, like an index entry.
    """
    return NestedField(instruction, result)


def _add_field_pieces(paragraph, pieces, code):
    """
    Append `pieces` - a string or `nested_field`, or a list of them - one run
    each. A string is field code (``w:instrText``) when `code` is true and
    result text (``w:t``) otherwise; an empty one adds nothing.
    """
    if isinstance(pieces, (str, NestedField)):
        pieces = [pieces]
    for piece in pieces:
        if isinstance(piece, NestedField):
            add_complex_field(paragraph, piece.instruction, piece.result)
        elif not piece:
            continue
        elif code:
            paragraph._p.append(parse_xml(
                '<w:r xmlns:w="%s"><w:instrText xml:space="preserve">%s</w:instrText></w:r>'
                % (NS_W, _escape(piece))))
        else:
            add_run(paragraph, piece)


def _add_fld_char(paragraph, kind):
    paragraph._p.append(parse_xml(
        '<w:r xmlns:w="%s"><w:fldChar w:fldCharType="%s"/></w:r>' % (NS_W, kind)))


def add_field_begin(paragraph, instruction):
    """
    Open a complex field: fldChar begin, then its code. The field is finished
    with `add_field_separate` (if it has a result) and `add_field_end`, in this
    paragraph or a later one, since Word lets a field run across paragraphs.
    Any other builder called in between puts its content inside the field.
    """
    _add_fld_char(paragraph, "begin")
    _add_field_pieces(paragraph, instruction, code=True)


def add_field_code(paragraph, instruction):
    """More code for a field opened in an earlier paragraph."""
    _add_field_pieces(paragraph, instruction, code=True)


def add_field_separate(paragraph, result):
    """The separate between an open field's code and its result, then result."""
    _add_fld_char(paragraph, "separate")
    _add_field_pieces(paragraph, result, code=False)


def add_field_result(paragraph, result):
    """More result for a field whose separate came in an earlier paragraph."""
    _add_field_pieces(paragraph, result, code=False)


def add_field_end(paragraph):
    """Close the innermost open field."""
    _add_fld_char(paragraph, "end")


def add_complex_field(paragraph, instruction, cached_result=None):
    """
    Append a complex field: fldChar begin / instrText / separate / cached
    result / fldChar end, each in a run of its own.

    This is how Word stores a TOC entry or a cross-reference. The cached result
    lives in an ordinary <w:t>, which is why it is scanned like any other text.
    `instruction` and `cached_result` are each a string or a list of strings
    and `nested_field` pieces: ``[" IF ", nested_field(" REF x ", "a"),
    ' = "a" "yes" "no" ']`` puts a REF, with its own cached result "a", inside
    the IF's code. With no `cached_result` the field has no separate.
    """
    add_field_begin(paragraph, instruction)
    if cached_result is not None:
        add_field_separate(paragraph, cached_result)
    add_field_end(paragraph)


def add_simple_field(paragraph, instruction, cached_result):
    """Append a <w:fldSimple> whose cached result sits in a child run."""
    xml = (
        '<w:fldSimple xmlns:w="%s" w:instr="%s">'
        '<w:r><w:t xml:space="preserve">%s</w:t></w:r></w:fldSimple>'
        % (NS_W, _escape(instruction), _escape(cached_result))
    )
    paragraph._p.append(parse_xml(xml))


def add_vml_textbox(paragraph, text):
    """Append a legacy VML text box containing a paragraph of `text`."""
    xml = (
        '<w:r xmlns:w="%s" xmlns:v="urn:schemas-microsoft-com:vml">'
        '<w:pict><v:shape id="tb1" style="width:100pt;height:50pt">'
        '<v:textbox><w:txbxContent>'
        '<w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r>'
        % (NS_W, _escape(text))
    )
    paragraph._p.append(parse_xml(xml))


def add_alternate_content_textbox(paragraph, choice_text, fallback_text):
    """
    Append a text box in mc:AlternateContent form.

    Word writes both a modern mc:Choice copy and a legacy mc:Fallback copy of
    the same text box. Only the Choice copy should be surfaced, otherwise one
    visible text box would report two matches.
    """
    xml = (
        '<mc:AlternateContent xmlns:mc="%s" xmlns:w="%s" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        '<mc:Choice Requires="wps">'
        '<w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p>'
        '</mc:Choice>'
        '<mc:Fallback>'
        '<w:p><w:r><w:pict><v:shape id="tb2" style="width:1pt;height:1pt">'
        '<v:textbox><w:txbxContent>'
        '<w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>'
        '</mc:Fallback>'
        '</mc:AlternateContent>'
        % (NS_MC, NS_W, _escape(choice_text), _escape(fallback_text))
    )
    paragraph._p.addnext(parse_xml(xml))


def add_notes_part(doc, kind, texts):
    """
    Attach a footnotes / endnotes / comments part containing one note per text.

    `kind` is "footnote", "endnote" or "comment". A separator entry is included
    for footnotes, mirroring what Word writes, so the traversal has to skip it.
    """
    from docx.opc.constants import CONTENT_TYPE as CT
    from docx.opc.packuri import PackURI
    from docx.opc.part import XmlPart

    spec = {
        "footnote": (CT.WML_FOOTNOTES, "footnotes", "/word/footnotes.xml"),
        "endnote": (CT.WML_ENDNOTES, "endnotes", "/word/endnotes.xml"),
        "comment": (CT.WML_COMMENTS, "comments", "/word/comments.xml"),
    }[kind]
    content_type, root_tag, partname = spec
    reltype = "%s/%s" % (NS_R, root_tag)

    entries = []
    if kind == "footnote":
        entries.append(
            '<w:footnote w:type="separator" w:id="-1">'
            '<w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        )
    for i, text in enumerate(texts, start=2):
        entries.append(
            '<w:%s w:id="%d"><w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p></w:%s>'
            % (kind, i, _escape(text), kind)
        )

    xml = '<w:%s xmlns:w="%s">%s</w:%s>' % (root_tag, NS_W, "".join(entries), root_tag)
    part = XmlPart(PackURI(partname), content_type, parse_xml(xml), doc.part.package)
    doc.part.relate_to(part, reltype)
    return part


def _escape(text):
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def count_drawings(path):
    """Number of <w:drawing> elements anywhere in the document body."""
    doc = Document(path)
    return len(doc.element.body.findall(".//" + qn("w:drawing")))


def all_text(path):
    """Every visible paragraph's text, using the processor's own traversal."""
    import document_processor as dp
    doc = Document(path)
    return [dp.paragraph_text(loc.element, loc.open_fields)
            for loc in dp._locate_paragraphs(doc)]


# How `layout` shows an element that is not text, by local name.
_LAYOUT_MARKS = {
    "drawing": "[drawing]", "pict": "[pict]", "object": "[object]",
    "footnoteReference": "[footnote]", "endnoteReference": "[endnote]",
    "commentReference": "[comment]", "sym": "[sym]", "fldSimple": "[fldSimple]",
    "softHyphen": "[softHyphen]", "del": "[del]", "moveFrom": "[del]",
    "bookmarkStart": "[bm]", "bookmarkEnd": "[/bm]",
}


def layout(p_el):
    """
    A w:p as its text with every non-text element shown where it sits, e.g.
    ``"See Figure 20[drawing]22 CBC ref"`` or
    ``"Article {begin}{separate}2.3{end} above"``. Field codes (w:instrText) are
    left out, as Word does. The walk is its own, independent of the processor's
    traversal, so a test can state where a picture or a field character is
    without trusting the code under test to say so.
    """
    out = []

    def visit(el):
        for child in el:
            if not isinstance(child.tag, str):
                continue
            local = child.tag.rsplit("}", 1)[-1]
            if child.tag == qn("w:t"):
                out.append(child.text or "")
            elif child.tag == qn("w:tab"):
                out.append("\t")
            elif child.tag in (qn("w:br"), qn("w:cr")):
                out.append("\n")
            elif child.tag == qn("w:fldChar"):
                out.append("{%s}" % child.get(qn("w:fldCharType")))
            elif child.tag in (qn("w:instrText"), qn("w:pPr"), qn("w:rPr")):
                continue
            elif child.tag.startswith("{%s}" % NS_W) and local in _LAYOUT_MARKS:
                out.append(_LAYOUT_MARKS[local])
            else:
                visit(child)

    visit(p_el)
    return "".join(out)


def body_layouts(path):
    """`layout` of every top-level body paragraph, in order."""
    return [layout(p) for p in Document(path).element.body.findall(qn("w:p"))]


def all_layouts(path):
    """`layout` of every w:p anywhere in the body, text-box paragraphs and
    paragraphs the processor does not collect included, in document order."""
    return [layout(p) for p in Document(path).element.body.iter(qn("w:p"))]


def revisions(path):
    """Every w:ins / w:del in the body, in document order, as
    (local tag, id, author, text)."""
    doc = Document(path)
    out = []
    for el in doc.element.body.iter(qn("w:ins"), qn("w:del")):
        texts = [t.text or "" for t in el.iter(qn("w:t"), qn("w:delText"))]
        out.append((el.tag.split("}")[1], el.get(qn("w:id")),
                    el.get(qn("w:author")), "".join(texts)))
    return out


def all_revision_ids(path):
    """Every w:ins / w:del id across every XML part of the package."""
    doc = Document(path)
    ids = []
    for part in doc.part.package.iter_parts():
        root = getattr(part, "element", None)
        if root is None:
            continue
        for el in root.iter(qn("w:ins"), qn("w:del"), qn("w:rPrChange")):
            ids.append(el.get(qn("w:id")))
    return ids


def _resolved_texts(path, keep, render=None):
    """Paragraph texts after simulating Word's Accept All (keep='ins') or
    Reject All (keep='del') on the whole package. `render` turns each w:p
    into a string; the processor's own `paragraph_text` by default, `layout`
    to see where pictures and field characters ended up."""
    import document_processor as dp
    doc = Document(path)
    drop, unwrap = (qn("w:del"), qn("w:ins")) if keep == "ins" else (qn("w:ins"), qn("w:del"))
    for part in doc.part.package.iter_parts():
        root = getattr(part, "element", None)
        if root is None:
            continue
        for el in list(root.iter(drop)):
            el.getparent().remove(el)
        for el in list(root.iter(unwrap)):
            parent = el.getparent()
            for child in list(el):
                if keep == "del":
                    for dt in child.iter(qn("w:delText")):
                        t = parse_xml('<w:t xmlns:w="%s" xml:space="preserve"/>' % NS_W)
                        t.text = dt.text
                        dt.getparent().replace(dt, t)
                el.addprevious(child)
            parent.remove(el)
    located = dp._locate_paragraphs(doc)
    if render is None:
        return [dp.paragraph_text(loc.element, loc.open_fields) for loc in located]
    return [render(loc.element) for loc in located]


def accept_all(path, render=None):
    return _resolved_texts(path, "ins", render)


def reject_all(path, render=None):
    return _resolved_texts(path, "del", render)


# ---------------------------------------------------------------- fixtures

@pytest.fixture
def png_file(tmp_path):
    path = tmp_path / "dot.png"
    path.write_bytes(_PNG_1PX)
    return str(path)


@pytest.fixture
def docdir(tmp_path):
    """An empty folder for .docx fixtures, matching the tool's folder-scan API."""
    d = tmp_path / "docs"
    d.mkdir()
    return d


@pytest.fixture
def hyperlink_doc(docdir):
    """'See <hyperlink>2022 CBC link</hyperlink> and done.'"""
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "See ")
    add_hyperlink(p, "2022 CBC link")
    add_run(p, " and done.")
    path = docdir / "hyperlink.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def merged_and_linked_doc(docdir):
    """A merged cell across 3 grid columns plus a header shared by 3 sections."""
    doc = Document()
    doc.add_paragraph("Body 2022 CBC baseline")

    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(0, 0).paragraphs[0].text = "Merged: 2022 CBC note"
    table.cell(1, 0).text = "Plain 2022 CBC cell"

    doc.sections[0].header.paragraphs[0].text = "HEADER 2022 CBC"
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_section(WD_SECTION.NEW_PAGE)

    path = docdir / "merged_linked.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def image_in_run_doc(docdir, png_file):
    """
    'See Figure 20' | [picture]'22' | ' CBC ref'. "2022 CBC" spans three runs
    and the picture sits inside it, between "20" and "22", so that match must
    be skipped: applying it would move the picture.
    """
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "See Figure 20")
    add_picture_run(p, png_file, after="22")
    add_run(p, " CBC ref")
    path = docdir / "image.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def image_beside_match_doc(docdir, png_file):
    """
    'See Figure ' | [picture]'20' | '22 CBC ref'. "2022 CBC" spans two runs and
    the first of them carries a picture, just outside the match: the edit is
    applied and the picture has to survive it, in place.
    """
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "See Figure ")
    add_picture_run(p, png_file, after="20")
    add_run(p, "22 CBC ref")
    path = docdir / "image_beside.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def figure_reference_doc(docdir, png_file):
    """
    ``<w:r><w:t>See Figure</w:t><w:drawing/></w:r><w:r><w:t> 3 for riser
    detail</w:t></w:r>``: the text reads "See Figure 3 for riser detail" with
    a picture between "Figure" and " 3".
    """
    doc = Document()
    p = doc.add_paragraph()
    add_picture_run(p, png_file, before="See Figure")
    add_run(p, " 3 for riser detail")
    path = docdir / "figure_ref.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def article_field_doc(docdir):
    """
    'Refer to Article ' + a REF field whose cached result is '2.3' + ' above.'
    The text reads "Refer to Article 2.3 above."
    """
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Refer to Article ")
    add_complex_field(p, r' REF _Ref1 \h ', "2.3")
    add_run(p, " above.")
    path = docdir / "article_field.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def nested_if_field_doc(docdir):
    """
    'Provide ' + { IF { REF SystemType } = "Wet" "..." "..." } + ' throughout.'

    The REF's cached result "Wet" is an ordinary w:t, but it sits in the IF's
    code, between the IF's begin and its separate. Word shows only the IF's
    result, so the text reads "Provide Wet-pipe system per 2022 CBC
    throughout." and not "Provide WetWet-pipe ...".
    """
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Provide ")
    add_complex_field(
        p, [" IF ", nested_field(r" REF SystemType \h ", "Wet"),
            ' = "Wet" "Wet-pipe system per 2022 CBC" "Dry-pipe system" '],
        "Wet-pipe system per 2022 CBC")
    add_run(p, " throughout.")
    path = docdir / "nested_if.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def toc_doc(docdir):
    """
    A table of contents over three paragraphs, as Word writes one: the TOC
    field begins in the first entry and ends in the last, and each entry
    carries a PAGEREF field nested in the TOC's result. A body paragraph
    follows. Entries 2 and 3 begin inside the TOC's result.
    """
    doc = Document()
    entries = [("PART 1 GENERAL", "_Toc1", "1"), ("PART 2 PRODUCTS", "_Toc2", "3"),
               ("PART 3 EXECUTION", "_Toc3", "5")]
    for i, (title, bookmark, page) in enumerate(entries):
        p = doc.add_paragraph()
        if i == 0:
            add_field_begin(p, r' TOC \o "1-3" \h \z \u ')
            add_field_separate(p, title)
        else:
            add_run(p, title)
        add_tab_run(p)
        add_complex_field(p, r" PAGEREF %s \h " % bookmark, page)
        if i == len(entries) - 1:
            add_field_end(p)
    doc.add_paragraph("Body text after the table of contents.")
    path = docdir / "toc.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def field_code_across_paragraphs_doc(docdir):
    """
    An IF field whose text holds paragraph marks, so its code and its result
    each run across paragraphs:

    1. 'Sprinklers: ' + begin, code ' IF { REF SystemType } = "Wet" "Provide
       wet-pipe system.'                          -> ends inside the code
    2. code 'Coordinate with { REF Standard }.'   -> wholly inside the code
    3. code 'Test per { REF Standard }." "Provide dry-pipe system." ',
       separate, result 'Provide wet-pipe system.' -> begins inside the code
    4. result 'Coordinate with { REF Standard }.' -> wholly inside the result
    5. result 'Test per { REF Standard }.', end, ' Field ends here.'

    Every REF has a cached result ("Wet", "NFPA 13"). The ones in paragraphs
    1-3 sit in the IF's code and are not on the page.
    """
    standard = nested_field(r" REF Standard \h ", "NFPA 13")
    doc = Document()
    p1 = doc.add_paragraph()
    add_run(p1, "Sprinklers: ")
    add_field_begin(p1, [" IF ", nested_field(r" REF SystemType \h ", "Wet"),
                         ' = "Wet" "Provide wet-pipe system.'])
    add_field_code(doc.add_paragraph(), ["Coordinate with ", standard, "."])
    p3 = doc.add_paragraph()
    add_field_code(p3, ["Test per ", standard, '." "Provide dry-pipe system." '])
    add_field_separate(p3, "Provide wet-pipe system.")
    add_field_result(doc.add_paragraph(), ["Coordinate with ", standard, "."])
    p5 = doc.add_paragraph()
    add_field_result(p5, ["Test per ", standard, "."])
    add_field_end(p5)
    add_run(p5, " Field ends here.")
    path = docdir / "field_code_across_paragraphs.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def nested_table_doc(docdir):
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Nested 2022 CBC"
    path = docdir / "nested.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def textbox_doc(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Anchor paragraph.")
    add_vml_textbox(p, "Textbox 2022 CBC")
    path = docdir / "textbox.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def alt_content_doc(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Anchor.")
    add_alternate_content_textbox(p, "Choice 2022 CBC", "Choice 2022 CBC")
    path = docdir / "altcontent.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def field_doc(docdir):
    """A complex field (cross-reference) and a w:fldSimple, both citing 2022 CBC."""
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Refer to ")
    add_complex_field(p, r' REF _Ref1 \\h ', "2022 CBC")
    add_run(p, " and ")
    add_simple_field(p, r' REF _Ref2 \\h ', "2022 CBC simple")
    add_run(p, " end.")
    path = docdir / "fields.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def tracked_changes_doc(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Kept ")
    add_tracked_insertion(p, "2022 CBC inserted")
    add_deletion(p, " 2022 CBC deleted")
    path = docdir / "tracked.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def notes_doc(docdir):
    """Body plus a footnote, an endnote and a comment, all citing 2022 CBC."""
    doc = Document()
    doc.add_paragraph("Body cites 2022 CBC here.")
    add_notes_part(doc, "footnote", ["Footnote cites 2022 CBC too."])
    add_notes_part(doc, "endnote", ["Endnote cites 2022 CBC as well."])
    add_notes_part(doc, "comment", ["Reviewer asks about 2022 CBC."])
    path = docdir / "notes.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def split_run_doc(docdir):
    """'Reference the 2022 CBC for compliance.' split across run boundaries."""
    doc = Document()
    p = doc.add_paragraph()
    add_runs(p, "Reference the 20", "22 ", "CBC for compliance.")
    path = docdir / "split.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def repeated_doc(docdir):
    doc = Document()
    doc.add_paragraph("aaaa and 2022 2022 2022")
    path = docdir / "repeated.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def formatted_split_doc(docdir):
    """'20' bold, '22 CBC' bold+red, ' ref' plain - a phrase spanning formats."""
    doc = Document()
    p = doc.add_paragraph()
    add_formatted_run(p, "20", bold=True)
    add_formatted_run(p, "22 CBC", bold=True, color="FF0000")
    add_run(p, " ref")
    path = docdir / "formatted.docx"
    doc.save(str(path))
    return str(path)


@pytest.fixture
def nested_folder_docs(docdir):
    """a.docx at the top, sub/b.docx, sub/deeper/c.docx, plus a Word lock file."""
    (docdir / "sub" / "deeper").mkdir(parents=True)
    paths = []
    for rel in ("a.docx", os.path.join("sub", "b.docx"),
                os.path.join("sub", "deeper", "c.docx")):
        doc = Document()
        doc.add_paragraph("Nested cites 2022 CBC.")
        path = docdir / rel
        doc.save(str(path))
        paths.append(str(path))
    (docdir / "sub" / "~$lock.docx").write_bytes(b"lock")
    return paths


@pytest.fixture
def heading_doc(docdir):
    """Heading 1 / body / Heading 2 / body / table / Heading 1 / body."""
    doc = Document()
    doc.add_heading("GENERAL", level=1)
    doc.add_paragraph("Intro cites 2022 CBC.")
    doc.add_heading("REFERENCES", level=2)
    doc.add_paragraph("Body cites 2022 CBC.")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell cites 2022 CBC."
    doc.add_heading("PRODUCTS", level=1)
    doc.add_paragraph("Later cites 2022 CBC.")
    path = docdir / "headings.docx"
    doc.save(str(path))
    return str(path)
