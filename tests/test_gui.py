"""
GUI behaviour, exercised headlessly against the tkinter stub in `tkstub.py`.

These cover the decisions the window makes rather than how it looks: which
status message is shown, whether a click toggles a checkbox, whether Apply
survives an error, and whether stale results can still be applied.
"""

import tkstub

tkstub.install()          # must precede `import main`

import os                 # noqa: E402
import json               # noqa: E402
from datetime import datetime  # noqa: E402

import pytest             # noqa: E402
import main               # noqa: E402
import document_processor as dp  # noqa: E402
import settings as app_settings  # noqa: E402


def make_match(location_type='body', detail='Paragraph 1', selected=True,
               file_path='C:\\specs\\a.docx', conflict=False, offset=0,
               find_text='2022 CBC', replace_text='2025 CBC', section=''):
    return dp.Match(
        file_path=file_path,
        location_type=location_type,
        location_detail=detail,
        paragraph_key='/word/document.xml#0',
        char_offset=offset,
        context_before='Per the ',
        match_text='2022 CBC',
        context_after=' requirements.',
        is_selected=selected and not conflict,
        find_text=find_text,
        replace_text=replace_text,
        section=section,
        conflict=conflict,
    )


def click(app, iid, element="text"):
    """Simulate a single click landing on `iid`."""
    app.tree.next_identify_element = element
    app.tree.next_identify_row = iid
    return app._on_tree_click(type("E", (), {"x": 80, "y": 10})())


def no_thread(monkeypatch):
    """Make Thread(...).start() a no-op so _start_* stops before the worker."""
    monkeypatch.setattr(main.threading, "Thread", lambda **kwargs: type(
        "T", (), {"start": lambda self: None})())


def empty_result(**overrides):
    result = {'total_replaced': 0, 'total_skipped': 0, 'files_modified': 0,
              'backups': [], 'backup_map': {}, 'skip_reasons': {}, 'errors': []}
    result.update(overrides)
    return result


@pytest.fixture
def app(tmp_path):
    """A window whose settings file lives under tmp_path, never in %APPDATA%."""
    tkstub.dialogs.reset()
    root = main.tk.Tk()
    return main.FindReplaceApp(root, settings_path=str(tmp_path / "settings.json"))


@pytest.fixture
def previewed(app):
    """An app showing one file with three matches, as after a preview."""
    fr = dp.FileResult(file_path='C:\\specs\\a.docx', file_name='a.docx',
                       matches=[make_match(), make_match(), make_match()])
    app.scan_folder = 'C:\\specs'
    app.scan_rules = [dp.Rule('2022 CBC', '2025 CBC')]
    app.scan_case_sensitive = True
    app._set_rules(app.scan_rules, invalidate=False)
    app._display_results([fr], {'files_scanned': 1, 'files_with_matches': 1,
                                'cancelled': False})
    return app


# ----------------------------------------------------- labels

@pytest.mark.parametrize("location_type,expected", [
    ('body', ''),
    ('table', '[Table] '),
    ('header', '[Header] '),
    ('footer', '[Footer] '),
    ('footnote', '[Footnote] '),
    ('endnote', '[Endnote] '),
    ('comment', '[Comment] '),
])
def test_every_location_type_has_a_label(app, location_type, expected):
    """A location the tree cannot name would show up as a bare, ambiguous row."""
    label = app._match_label(make_match(location_type=location_type))
    assert label.startswith(f"{main.CHECKED} {expected}")


def test_label_reflects_checkbox_state(app):
    assert app._match_label(make_match(selected=True)).startswith(main.CHECKED)
    assert app._match_label(make_match(selected=False)).startswith(main.UNCHECKED)


# ----------------------------------------------------- status messages

def test_no_docx_files_message(app):
    app._display_results([], {'files_scanned': 0, 'cancelled': False})
    assert app.progress_var.get() == "No .docx files found in the selected folder"


def test_no_matches_in_a_folder_full_of_files(app):
    """The old message claimed the folder was empty. It was not."""
    app._display_results([], {'files_scanned': 200, 'cancelled': False})
    assert "No matches found in 200 .docx files" == app.progress_var.get()


def test_cancelled_scan_says_so(app):
    app._display_results([], {'files_scanned': 12, 'cancelled': True})
    assert "cancelled after 12 files" in app.progress_var.get()


def test_cancelled_scan_with_partial_results(previewed):
    fr = dp.FileResult(file_path='b.docx', file_name='b.docx', matches=[make_match()])
    previewed._display_results([fr], {'files_scanned': 3, 'cancelled': True})
    status = previewed.progress_var.get()
    assert "Found 1 match" in status
    assert "cancelled after 3 files" in status


# ----------------------------------------------------- tree interaction

def test_clicking_the_expander_does_not_toggle_the_file(previewed):
    """
    identify_row alone ignores x, so a click on the expand arrow used to flip
    every checkbox in the file.
    """
    file_node = previewed.tree.get_children()[0]
    before = [m.is_selected for m in previewed.all_matches]

    previewed.tree.next_identify_element = "Treeitem.indicator"
    previewed.tree.next_identify_row = file_node
    previewed._on_tree_click(type("E", (), {"x": 12, "y": 10})())

    assert [m.is_selected for m in previewed.all_matches] == before


def test_clicking_a_file_row_toggles_all_its_matches(previewed):
    file_node = previewed.tree.get_children()[0]
    previewed.tree.next_identify_element = "text"
    previewed.tree.next_identify_row = file_node
    previewed._on_tree_click(type("E", (), {"x": 80, "y": 10})())

    assert all(not m.is_selected for m in previewed.all_matches)


def test_clicking_a_match_toggles_only_that_match(previewed):
    file_node = previewed.tree.get_children()[0]
    first_child = previewed.tree.get_children(file_node)[0]
    previewed.tree.next_identify_element = "text"
    previewed.tree.next_identify_row = first_child
    previewed._on_tree_click(type("E", (), {"x": 80, "y": 30})())

    assert [m.is_selected for m in previewed.all_matches] == [False, True, True]


def test_tree_is_inert_while_applying(previewed):
    """The worker thread reads the same Match objects the tree would mutate."""
    file_node = previewed.tree.get_children()[0]
    previewed.processing_kind = 'apply'
    previewed.tree.next_identify_element = "text"
    previewed.tree.next_identify_row = file_node

    result = previewed._on_tree_click(type("E", (), {"x": 80, "y": 10})())

    assert result == "break"
    assert all(m.is_selected for m in previewed.all_matches)


def test_select_all_is_inert_while_applying(previewed):
    previewed._deselect_all()
    previewed.processing_kind = 'apply'
    previewed._select_all()
    assert all(not m.is_selected for m in previewed.all_matches)


# ----------------------------------------------------- stale settings

def test_changing_the_rule_set_invalidates_the_preview(previewed):
    """
    Offsets belong to the rules they were found with. Leaving them applicable
    under a new rule would validate old offsets against new rules.
    """
    previewed.find_var.set("2022 CFC")
    previewed.replace_var.set("2025 CFC")
    assert len(previewed.all_matches) == 3          # typing alone is not a rule

    previewed._add_rule()

    assert previewed.all_matches == []
    assert previewed.tree.get_children() == ()
    assert previewed.apply_btn.state == "disabled"
    assert "run Preview Changes again" in previewed.progress_var.get()


def test_removing_a_rule_invalidates_the_preview(previewed):
    previewed.rules_tree.selection_set(previewed.rule_iids[0])
    previewed._remove_rules()
    assert previewed.rules == []
    assert previewed.all_matches == []


def test_whole_word_recursive_and_folder_invalidate_the_preview(app):
    for var in (app.whole_word_var, app.recursive_var, app.folder_var):
        fr = dp.FileResult(file_path='a.docx', file_name='a.docx', matches=[make_match()])
        app._display_results([fr], {'files_scanned': 1, 'cancelled': False})
        assert app.all_matches
        var.set(True if var is not app.folder_var else "D:\\other")
        assert app.all_matches == []


def test_apply_options_do_not_invalidate_the_preview(previewed):
    for var in (previewed.backup_var, previewed.changelog_var, previewed.tracked_var):
        var.set(not var.get())
    previewed.author_var.set("Abe")
    assert len(previewed.all_matches) == 3


def test_toggling_case_sensitivity_invalidates_the_preview(previewed):
    previewed.case_var.set(False)
    assert previewed.all_matches == []
    assert "run Preview Changes again" in previewed.progress_var.get()


def test_typing_in_the_editor_keeps_the_preview(previewed):
    """The editor is not a rule until Add is clicked, so it does not invalidate."""
    previewed.find_var.set("2022 CFC")
    previewed.replace_var.set("2025 CFC")
    assert len(previewed.all_matches) == 3
    assert previewed.apply_btn.state == "normal"


def test_confirm_dialog_shows_the_rules_that_were_scanned(previewed, monkeypatch):
    """
    The dialog used to print the live entry box, so it could name a search that
    was never run.
    """
    no_thread(monkeypatch)
    previewed._set_rules([dp.Rule("something else entirely", "x")], invalidate=False)

    previewed._start_apply()

    _title, message = tkstub.dialogs.questions[-1]
    assert f'"2022 CBC" {main.ARROW} "2025 CBC"' in message
    assert "something else entirely" not in message
    assert "Case sensitive: yes" in message
    assert "edit documents in place" in message
    assert "Backups will be created." in message


def test_confirm_dialog_names_tracked_mode_and_a_pending_editor_row(previewed, monkeypatch):
    no_thread(monkeypatch)
    previewed.tracked_var.set(True)
    previewed.author_var.set("Jane")
    previewed.backup_var.set(False)
    previewed.find_var.set("not added yet")

    previewed._start_apply()

    _title, message = tkstub.dialogs.questions[-1]
    assert "tracked changes by Jane" in message
    assert "NO backups" in message
    assert 'holds "not added yet"' in message


def test_apply_uses_the_case_setting_from_the_scan(previewed, monkeypatch):
    captured = {}

    def fake_apply(matches, replace_text=None, **kwargs):
        captured.update(kwargs)
        return empty_result()

    monkeypatch.setattr(main, "apply_changes", fake_apply)
    previewed.scan_case_sensitive = True
    previewed.case_var._value = False        # changed after the scan
    previewed.changelog_var.set(False)

    previewed._run_apply(previewed._apply_options())

    assert captured['case_sensitive'] is True


def test_apply_worker_gets_a_snapshot_not_the_live_widgets(previewed, monkeypatch):
    captured = {}

    def fake_apply(matches, replace_text=None, **kwargs):
        captured.update(kwargs)
        return empty_result()

    monkeypatch.setattr(main, "apply_changes", fake_apply)
    previewed.tracked_var.set(True)
    previewed.author_var.set("  Abe  ")
    previewed.backup_var.set(False)
    previewed.changelog_var.set(False)
    options = previewed._apply_options()
    previewed.tracked_var.set(False)         # changed after the snapshot
    previewed.author_var.set("someone else")

    previewed._run_apply(options)

    assert captured['tracked_changes'] is True
    assert captured['author'] == "Abe"
    assert captured['create_backups'] is False


def test_change_log_is_written_to_the_scan_folder(previewed, monkeypatch, tmp_path):
    monkeypatch.setattr(main, "apply_changes", lambda *a, **k: empty_result(
        total_replaced=3, files_modified=1, backup_map={'C:\\specs\\a.docx': 'b'}))
    previewed.scan_folder = str(tmp_path)
    previewed.changelog_var.set(True)

    previewed._run_apply(previewed._apply_options())

    logs = [f for f in os.listdir(str(tmp_path)) if f.startswith(main.CHANGE_LOG_PREFIX)]
    assert len(logs) == 1
    _title, message = tkstub.dialogs.info[-1]
    assert logs[0] in message
    with open(os.path.join(str(tmp_path), logs[0]), encoding="utf-8-sig") as fh:
        assert fh.read().count("\n") == 4       # header + 3 matches


def test_change_log_failure_is_reported_not_fatal(previewed, monkeypatch, tmp_path):
    monkeypatch.setattr(main, "apply_changes", lambda *a, **k: empty_result(total_replaced=3))
    previewed.scan_folder = str(tmp_path / "missing")
    previewed.changelog_var.set(True)

    previewed._run_apply(previewed._apply_options())

    _title, message = tkstub.dialogs.info[-1]
    assert "Could not write the change log" in message
    assert previewed.all_matches == []


def test_change_log_can_be_switched_off(previewed, monkeypatch, tmp_path):
    monkeypatch.setattr(main, "apply_changes", lambda *a, **k: empty_result())
    previewed.scan_folder = str(tmp_path)
    previewed.changelog_var.set(False)
    previewed._run_apply(previewed._apply_options())
    assert os.listdir(str(tmp_path)) == []


# ----------------------------------------------------- apply reporting

def test_skipped_changes_are_reported(previewed):
    previewed._apply_complete({'total_replaced': 6, 'total_skipped': 3,
                               'files_modified': 2, 'backups': ['a', 'b'],
                               'errors': []})
    _title, message = tkstub.dialogs.info[-1]
    assert "3 selected changes were NOT applied" in message
    assert "Backups created: 2" in message
    assert "(3 skipped)" in previewed.progress_var.get()


def test_skip_reasons_are_listed(previewed):
    previewed._apply_complete(empty_result(
        total_replaced=1, total_skipped=3, files_modified=1,
        skip_reasons={dp._REASON_NESTED_INS: 2, dp._REASON_MOVED: 1}))
    _title, message = tkstub.dialogs.info[-1]
    assert f"2: {dp._REASON_NESTED_INS}" in message
    assert f"1: {dp._REASON_MOVED}" in message


def test_clean_run_shows_no_warning(previewed):
    previewed._apply_complete({'total_replaced': 4, 'total_skipped': 0,
                               'files_modified': 1, 'backups': ['a'],
                               'errors': []})
    _title, message = tkstub.dialogs.info[-1]
    assert "NOT applied" not in message


def test_errors_are_listed(previewed):
    previewed._apply_complete({'total_replaced': 0, 'total_skipped': 3,
                               'files_modified': 0, 'backups': [],
                               'errors': ["a.docx: was modified on disk since the preview."]})
    _title, message = tkstub.dialogs.info[-1]
    assert "was modified on disk" in message


def test_results_are_cleared_after_a_successful_apply(previewed):
    previewed._apply_complete({'total_replaced': 3, 'total_skipped': 0,
                               'files_modified': 1, 'backups': [], 'errors': []})
    assert previewed.all_matches == []
    assert previewed.apply_btn.state == "disabled"


def test_apply_error_leaves_apply_available_for_retry(previewed):
    """Losing the button meant re-scanning the whole folder to try again."""
    previewed.processing_kind = 'apply'
    previewed._apply_error("disk full")

    assert previewed.apply_btn.state == "normal"
    assert len(previewed.all_matches) == 3
    assert tkstub.dialogs.errors


# ----------------------------------------------------- shutdown

def test_closing_during_apply_is_refused(previewed):
    previewed.processing_kind = 'apply'
    previewed._on_close()

    assert previewed.root.destroyed is False
    assert tkstub.dialogs.warnings


def test_closing_during_scan_cancels_then_quits(app):
    import threading
    app.processing_kind = 'scan'
    app.scan_cancel = threading.Event()
    tkstub.dialogs.answer = True

    app._on_close()

    assert app.scan_cancel.is_set()
    assert app.close_when_idle is True
    assert app.root.destroyed is False      # not until the scan actually stops

    app._finish_processing()
    assert app.root.destroyed is True


def test_declining_the_scan_close_prompt_keeps_running(app):
    import threading
    app.processing_kind = 'scan'
    app.scan_cancel = threading.Event()
    tkstub.dialogs.answer = False

    app._on_close()

    assert not app.scan_cancel.is_set()
    assert app.close_when_idle is False
    assert app.root.destroyed is False


def test_closing_while_idle_quits(app):
    app._on_close()
    assert app.root.destroyed is True


def test_results_survive_a_scan_that_completes_after_a_close_request(app):
    """
    _finish_processing may destroy the window, so it has to run last - otherwise
    the tree is populated after the root is gone.
    """
    import threading
    app.processing_kind = 'scan'
    app.scan_cancel = threading.Event()
    app.close_when_idle = True

    fr = dp.FileResult(file_path='a.docx', file_name='a.docx', matches=[make_match()])
    app._display_results([fr], {'files_scanned': 1, 'cancelled': True})

    assert app.root.destroyed is True
    assert len(app.all_matches) == 1        # populated before the window went away
    assert "cancelled after 1 file" in app.progress_var.get()


# ----------------------------------------------------- packaged-build selftest

def test_selftest_passes_and_reports(tmp_path):
    """
    `--selftest` is what CI runs against the built .exe to prove the frozen
    bundle actually works. If the check itself is broken, a packaging failure
    would look like a passing build.
    """
    report = tmp_path / "selftest.txt"
    assert main.selftest(str(report)) == 0

    text = report.read_text(encoding="utf-8")
    assert "SELFTEST PASSED" in text
    assert "document engine OK" in text
    assert main.__version__ in text


def test_selftest_needs_no_report_path(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert main.selftest() == 0


def test_selftest_reports_failure_rather_than_raising(monkeypatch, tmp_path):
    """A broken bundle must produce a non-zero exit, not an unhandled crash."""
    import document_processor as dp_mod

    def boom(*args, **kwargs):
        raise RuntimeError("python-docx templates missing from the bundle")

    monkeypatch.setattr(dp_mod, "scan_documents", boom)
    report = tmp_path / "fail.txt"

    assert main.selftest(str(report)) == 1
    text = report.read_text(encoding="utf-8")
    assert "SELFTEST FAILED" in text
    assert "templates missing from the bundle" in text


# ----------------------------------------------------- rules editor

def test_add_rule_creates_a_row_and_clears_the_editor(app):
    app.find_var.set("2022 CBC")
    app.replace_var.set("2025 CBC")
    app._add_rule()
    assert app.rules == [dp.Rule("2022 CBC", "2025 CBC")]
    assert app.rules_tree.values_of(app.rule_iids[0]) == ("2022 CBC", "2025 CBC")
    assert app.find_var.get() == "" and app.replace_var.get() == ""


def test_add_with_an_existing_find_updates_that_rule_in_place(app):
    app._set_rules([dp.Rule("a", "1"), dp.Rule("b", "2")])
    app.find_var.set("a")
    app.replace_var.set("changed")
    app._add_rule()
    assert app.rules == [dp.Rule("a", "changed"), dp.Rule("b", "2")]


def test_add_with_an_empty_find_warns(app):
    app.replace_var.set("only a replacement")
    app._add_rule()
    assert app.rules == []
    assert tkstub.dialogs.warnings


def test_remove_removes_exactly_the_selected_rules(app):
    app._set_rules([dp.Rule("a"), dp.Rule("b"), dp.Rule("c")])
    app.rules_tree.selection_set(app.rule_iids[0], app.rule_iids[2])
    app._remove_rules()
    assert app.rules == [dp.Rule("b")]
    assert len(app.rules_tree.get_children()) == 1


def test_remove_with_nothing_selected_explains(app):
    app._set_rules([dp.Rule("a")])
    app._remove_rules()
    assert app.rules == [dp.Rule("a")]
    assert tkstub.dialogs.info


def test_selecting_a_rule_loads_it_into_the_editor(app):
    app._set_rules([dp.Rule("2022 CBC", "2025 CBC"), dp.Rule("x", "y")])
    app.rules_tree.selection_set(app.rule_iids[1])
    app._on_rule_selected()
    assert (app.find_var.get(), app.replace_var.get()) == ("x", "y")


def test_load_set_replaces_the_rules(app, monkeypatch, tmp_path):
    path = str(tmp_path / "set.csv")
    app_settings.save_rules_csv(path, [dp.Rule("2022 CBC", "2025 CBC"),
                                       dp.Rule("2022 CFC", "2025 CFC")])
    monkeypatch.setattr(main.filedialog, "askopenfilename", lambda **kw: path)
    app._set_rules([dp.Rule("old")])
    app._load_rule_set()
    assert app.rules == [dp.Rule("2022 CBC", "2025 CBC"), dp.Rule("2022 CFC", "2025 CFC")]
    assert "Loaded 2 rules" in app.progress_var.get()


def test_load_set_reports_a_bad_file_and_keeps_the_rules(app, monkeypatch, tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("find,replace\n,oops\n", encoding="utf-8")
    monkeypatch.setattr(main.filedialog, "askopenfilename", lambda **kw: str(path))
    app._set_rules([dp.Rule("old")])
    app._load_rule_set()
    assert app.rules == [dp.Rule("old")]
    assert "Row 2" in tkstub.dialogs.warnings[-1][1]


def test_load_set_cancelled_changes_nothing(app, monkeypatch):
    monkeypatch.setattr(main.filedialog, "askopenfilename", lambda **kw: "")
    app._set_rules([dp.Rule("old")])
    app._load_rule_set()
    assert app.rules == [dp.Rule("old")]


def test_save_set_writes_the_rules_including_a_pending_row(app, monkeypatch, tmp_path):
    path = str(tmp_path / "out.csv")
    monkeypatch.setattr(main.filedialog, "asksaveasfilename", lambda **kw: path)
    app._set_rules([dp.Rule("a", "1")])
    app.find_var.set("b")
    app.replace_var.set("2")
    app._save_rule_set()
    assert app_settings.load_rules_csv(path) == [dp.Rule("a", "1"), dp.Rule("b", "2")]
    assert "Saved 2 rules" in app.progress_var.get()


def test_save_set_with_no_rules_explains(app, monkeypatch):
    monkeypatch.setattr(main.filedialog, "asksaveasfilename", lambda **kw: "x.csv")
    app._save_rule_set()
    assert tkstub.dialogs.info


def test_preview_auto_adds_the_pending_editor_row(app, monkeypatch, tmp_path):
    no_thread(monkeypatch)
    app.folder_var.set(str(tmp_path))
    app.find_var.set("2022 CBC")
    app.replace_var.set("2025 CBC")
    app._start_preview()
    assert app.rules == [dp.Rule("2022 CBC", "2025 CBC")]
    assert app.scan_rules == [dp.Rule("2022 CBC", "2025 CBC")]
    assert app.find_var.get() == ""
    assert app.processing_kind == 'scan'


def test_preview_refuses_zero_rules(app, tmp_path):
    app.folder_var.set(str(tmp_path))
    app._start_preview()
    assert app.processing_kind is None
    assert "text to find" in tkstub.dialogs.warnings[-1][1]


def test_preview_refuses_case_insensitive_duplicates(app, tmp_path):
    app.folder_var.set(str(tmp_path))
    app._set_rules([dp.Rule("CBC", "1"), dp.Rule("cbc", "2")])
    app.case_var.set(False)
    app._start_preview()
    assert app.processing_kind is None
    assert "case is ignored" in tkstub.dialogs.warnings[-1][1]


def test_preview_snapshots_every_scan_setting(app, monkeypatch, tmp_path):
    no_thread(monkeypatch)
    app.folder_var.set(str(tmp_path))
    app._set_rules([dp.Rule("a", "b")])
    app.case_var.set(False)
    app.whole_word_var.set(True)
    app.recursive_var.set(True)
    app._start_preview()
    assert (app.scan_folder, app.scan_case_sensitive, app.scan_whole_word,
            app.scan_recursive) == (str(tmp_path), False, True, True)


def test_a_real_preview_end_to_end(app, monkeypatch, docdir):
    """Thread runs synchronously; the whole scan -> display path is exercised."""
    from docx import Document
    doc = Document()
    doc.add_heading("GENERAL", level=1)
    doc.add_paragraph("Per 2022 CBC.")
    doc.save(str(docdir / "real.docx"))

    class SyncThread:
        def __init__(self, target, args=(), daemon=None):
            self.target, self.args = target, args

        def start(self):
            self.target(*self.args)

    monkeypatch.setattr(main.threading, "Thread", SyncThread)
    app.folder_var.set(str(docdir))
    app.find_var.set("2022 CBC")
    app.replace_var.set("2025 CBC")

    app._start_preview()

    assert app.processing_kind is None
    assert len(app.all_matches) == 1
    assert "Found 1 match in 1 file" in app.progress_var.get()
    match_node = app.tree.get_children(app.tree.get_children()[0])[0]
    assert app.tree.values_of(match_node) == (f"2022 CBC {main.ARROW} 2025 CBC", "GENERAL")


# ----------------------------------------------------- columns and conflicts

def test_result_rows_carry_rule_and_section_columns(app):
    fr = dp.FileResult(file_path='a.docx', file_name='a.docx', matches=[
        make_match(section="GENERAL > REFERENCES")])
    app._display_results([fr], {'files_scanned': 1, 'cancelled': False})
    file_node = app.tree.get_children()[0]
    match_node = app.tree.get_children(file_node)[0]
    assert app.tree.values_of(match_node) == (f"2022 CBC {main.ARROW} 2025 CBC",
                                              "GENERAL > REFERENCES")
    assert app.tree.values_of(file_node) == ("", "")


@pytest.fixture
def conflicted(app):
    fr = dp.FileResult(file_path='a.docx', file_name='a.docx', matches=[
        make_match(conflict=True, find_text="2022"),
        make_match(conflict=True),
        make_match(offset=13),
    ])
    app.scan_rules = [dp.Rule("2022", "2025"), dp.Rule("2022 CBC", "2025 CBC")]
    app._display_results([fr], {'files_scanned': 1, 'cancelled': False})
    return app


def test_conflict_rows_are_marked_and_tagged(conflicted):
    file_node = conflicted.tree.get_children()[0]
    rows = conflicted.tree.get_children(file_node)
    assert conflicted.tree.text_of(rows[0]).startswith(main.WARNING)
    assert conflicted.tree.tags_of(rows[0]) == ("conflict",)
    assert conflicted.tree.text_of(rows[2]).startswith(main.CHECKED)
    assert conflicted.tree.tags_of(rows[2]) == ()
    assert "2 overlapping" in conflicted.progress_var.get()
    assert "1 of 1 selected (2 overlapping, not selectable)" in \
        conflicted.match_count_label.options["text"]
    assert "3 matches, 1 selected, 2 overlapping" in conflicted.tree.text_of(file_node)


def test_conflict_rows_cannot_be_selected(conflicted):
    file_node = conflicted.tree.get_children()[0]
    rows = conflicted.tree.get_children(file_node)
    click(conflicted, rows[0])
    assert conflicted.all_matches[0].is_selected is False

    click(conflicted, file_node)                    # deselect all selectable
    assert [m.is_selected for m in conflicted.all_matches] == [False, False, False]
    click(conflicted, file_node)                    # select all selectable
    assert [m.is_selected for m in conflicted.all_matches] == [False, False, True]

    conflicted._select_all()
    assert [m.is_selected for m in conflicted.all_matches] == [False, False, True]
    assert conflicted.apply_btn.state == "normal"


def test_file_checkbox_reflects_selectable_matches_only(conflicted):
    file_node = conflicted.tree.get_children()[0]
    assert conflicted.tree.text_of(file_node).startswith(main.CHECKED)
    conflicted._deselect_all()
    assert conflicted.tree.text_of(file_node).startswith(main.UNCHECKED)
    assert conflicted.apply_btn.state == "disabled"


def test_apply_with_only_conflicts_selected_is_refused(conflicted, monkeypatch):
    no_thread(monkeypatch)
    conflicted._deselect_all()
    conflicted.all_matches[0].is_selected = True      # hand-forced
    conflicted._start_apply()
    assert conflicted.processing_kind is None
    assert tkstub.dialogs.info[-1][0] == "Nothing to Apply"


# ----------------------------------------------------- double-click

def test_double_click_opens_the_document_behind_a_match(previewed, monkeypatch):
    opened = []
    monkeypatch.setattr(main, "_open_document", opened.append)
    file_node = previewed.tree.get_children()[0]
    match_node = previewed.tree.get_children(file_node)[0]
    previewed.tree.next_identify_row = match_node
    assert previewed._on_tree_double_click(type("E", (), {"x": 5, "y": 5})()) == "break"
    assert opened == ['C:\\specs\\a.docx']

    previewed.tree.next_identify_row = file_node
    previewed._on_tree_double_click(type("E", (), {"x": 5, "y": 5})())
    assert opened == ['C:\\specs\\a.docx', 'C:\\specs\\a.docx']


def test_double_click_is_inert_while_processing_and_off_rows(previewed, monkeypatch):
    opened = []
    monkeypatch.setattr(main, "_open_document", opened.append)
    previewed.tree.next_identify_row = ""
    assert previewed._on_tree_double_click(type("E", (), {"x": 5, "y": 5})()) is None
    previewed.processing_kind = 'apply'
    previewed.tree.next_identify_row = previewed.tree.get_children()[0]
    assert previewed._on_tree_double_click(type("E", (), {"x": 5, "y": 5})()) == "break"
    assert opened == []


def test_double_click_failure_is_shown_not_raised(previewed, monkeypatch):
    def boom(path):
        raise OSError("no association")
    monkeypatch.setattr(main, "_open_document", boom)
    previewed.tree.next_identify_row = previewed.tree.get_children()[0]
    previewed._on_tree_double_click(type("E", (), {"x": 5, "y": 5})())
    assert "no association" in tkstub.dialogs.errors[-1][1]


# ----------------------------------------------------- restore dialog

def _backup(doc, stamp, seq=0):
    return dp.Backup(doc, f"{doc}.{stamp}{'-' + str(seq) if seq else ''}.bak",
                     datetime.strptime(stamp, "%Y%m%d-%H%M%S"), seq)


@pytest.fixture
def backups(tmp_path):
    a = os.path.join(str(tmp_path), "a.docx")
    b = os.path.join(str(tmp_path), "sub", "b.docx")
    return [_backup(a, "20260910-090000"), _backup(a, "20260909-141530", 1),
            _backup(a, "20260909-141530"), _backup(b, "20260101-000000")]


def test_restore_dialog_prechecks_the_newest_backup_per_document(app, backups, tmp_path):
    dialog = main.RestoreDialog(app, str(tmp_path), backups)
    rows = dialog.tree.get_children()
    assert [dialog.checked[r] for r in rows] == [True, False, False, True]
    assert dialog.tree.text_of(rows[0]) == f"{main.CHECKED} a.docx"
    assert dialog.tree.text_of(rows[3]) == f"{main.CHECKED} {os.path.join('sub', 'b.docx')}"
    assert dialog.tree.values_of(rows[1]) == ("2026-09-09 14:15:30 (#1)",)
    assert dialog.selected_backups() == [backups[0], backups[3]]


def test_restore_dialog_allows_one_backup_per_document(app, backups, tmp_path):
    dialog = main.RestoreDialog(app, str(tmp_path), backups)
    rows = dialog.tree.get_children()
    dialog.tree.next_identify_row = rows[2]
    dialog._on_click(type("E", (), {"x": 5, "y": 5})())
    assert [dialog.checked[r] for r in rows] == [False, False, True, True]
    dialog.toggle(rows[2])
    assert [dialog.checked[r] for r in rows] == [False, False, False, True]


def test_restore_dialog_restores_the_ticked_backups_after_confirming(
        app, backups, tmp_path, monkeypatch):
    called = []
    monkeypatch.setattr(main, "restore_backups", lambda chosen: (
        called.append(list(chosen)) or {'restored': len(chosen), 'restored_paths': [], 'errors': []}))
    fr = dp.FileResult(file_path='a.docx', file_name='a.docx', matches=[make_match()])
    app._display_results([fr], {'files_scanned': 1, 'cancelled': False})

    dialog = main.RestoreDialog(app, str(tmp_path), backups)
    tkstub.dialogs.answer = True
    dialog._restore()

    assert called == [[backups[0], backups[3]]]
    assert "lost" in tkstub.dialogs.questions[-1][1]
    assert dialog.top.destroyed is True
    assert app.all_matches == []
    assert "Restored 2 documents" in app.progress_var.get()


def test_restore_dialog_declined_restores_nothing(app, backups, tmp_path, monkeypatch):
    monkeypatch.setattr(main, "restore_backups", lambda chosen: pytest.fail("should not run"))
    dialog = main.RestoreDialog(app, str(tmp_path), backups)
    tkstub.dialogs.answer = False
    dialog._restore()
    assert dialog.top.destroyed is False


def test_restore_dialog_with_nothing_ticked_explains(app, backups, tmp_path, monkeypatch):
    monkeypatch.setattr(main, "restore_backups", lambda chosen: pytest.fail("should not run"))
    dialog = main.RestoreDialog(app, str(tmp_path), backups)
    for iid in list(dialog.checked):
        dialog.checked[iid] = False
    dialog._restore()
    assert tkstub.dialogs.info[-1][0] == "Nothing Selected"


def test_restore_errors_are_shown(app, backups, tmp_path, monkeypatch):
    monkeypatch.setattr(main, "restore_backups", lambda chosen: {
        'restored': 1, 'restored_paths': [], 'errors': ["Could not restore b.docx: locked"]})
    dialog = main.RestoreDialog(app, str(tmp_path), backups)
    tkstub.dialogs.answer = True
    dialog._restore()
    assert "locked" in tkstub.dialogs.errors[-1][1]
    assert "Restored 1 document" in app.progress_var.get()


def test_restore_button_lists_real_backups(app, monkeypatch, docdir):
    from docx import Document
    doc = Document()
    doc.add_paragraph("2022 CBC")
    doc.save(str(docdir / "r.docx"))
    matches = [m for r in dp.scan_documents(str(docdir), "2022 CBC") for m in r.matches]
    dp.apply_changes(matches, "2025 CBC", create_backups=True)

    opened = []
    monkeypatch.setattr(main, "RestoreDialog", lambda a, f, b: opened.append(b))
    app.folder_var.set(str(docdir))
    app._open_restore_dialog()
    assert len(opened) == 1 and len(opened[0]) == 1
    assert opened[0][0].document_path == str(docdir / "r.docx")


def test_restore_button_with_no_backups_explains(app, tmp_path):
    app.folder_var.set(str(tmp_path))
    app._open_restore_dialog()
    assert tkstub.dialogs.info[-1][0] == "No Backups"


def test_restore_button_is_refused_while_processing_or_without_a_folder(app, monkeypatch):
    monkeypatch.setattr(main, "RestoreDialog", lambda *a: pytest.fail("should not open"))
    app._open_restore_dialog()
    assert tkstub.dialogs.warnings
    app.folder_var.set("C:\\nope\\missing")
    app._open_restore_dialog()
    app.processing_kind = 'scan'
    app.folder_var._value = os.getcwd()
    app._open_restore_dialog()


# ----------------------------------------------------- settings

def test_settings_are_restored_on_construction(tmp_path):
    path = str(tmp_path / "settings.json")
    app_settings.save_settings({
        "folder": "C:\\specs", "rules": [dp.Rule("2022 CBC", "2025 CBC")],
        "case_sensitive": False, "whole_word": True, "recursive": True,
        "create_backups": False, "write_change_log": False, "tracked_changes": True,
        "author": "Abe", "geometry": "800x600+1+2"}, path)
    tkstub.dialogs.reset()
    app = main.FindReplaceApp(main.tk.Tk(), settings_path=path)
    assert app.folder_var.get() == "C:\\specs"
    assert app.rules == [dp.Rule("2022 CBC", "2025 CBC")]
    assert app.rules_tree.values_of(app.rule_iids[0]) == ("2022 CBC", "2025 CBC")
    assert (app.case_var.get(), app.whole_word_var.get(), app.recursive_var.get()) == (False, True, True)
    assert (app.backup_var.get(), app.changelog_var.get(), app.tracked_var.get()) == (False, False, True)
    assert app.author_var.get() == "Abe"
    assert app.all_matches == []                 # nothing to invalidate on load


def test_settings_are_saved_on_close(app):
    app.folder_var.set("D:\\jobs")
    app._set_rules([dp.Rule("a", "b")])
    app.tracked_var.set(True)
    app.author_var.set("Abe")
    app._on_close()
    assert app.root.destroyed is True
    saved = app_settings.load_settings(app.settings_path)
    assert saved["folder"] == "D:\\jobs"
    assert saved["rules"] == [dp.Rule("a", "b")]
    assert saved["tracked_changes"] is True
    assert saved["author"] == "Abe"


def test_settings_are_saved_when_closing_after_a_cancelled_scan(app):
    import threading
    app._set_rules([dp.Rule("a", "b")])
    app.processing_kind = 'scan'
    app.scan_cancel = threading.Event()
    tkstub.dialogs.answer = True
    app._on_close()
    app._finish_processing()
    assert app.root.destroyed is True
    assert app_settings.load_settings(app.settings_path)["rules"] == [dp.Rule("a", "b")]


def test_corrupt_settings_are_ignored(tmp_path):
    path = tmp_path / "settings.json"
    path.write_text("{oops", encoding="utf-8")
    tkstub.dialogs.reset()
    app = main.FindReplaceApp(main.tk.Tk(), settings_path=str(path))
    assert app.rules == [] and app.case_var.get() is True


def test_unwritable_settings_do_not_stop_a_close(app, monkeypatch):
    def boom(*a, **k):
        raise OSError("read-only")
    monkeypatch.setattr(main.app_settings, "save_settings", boom)
    app._on_close()
    assert app.root.destroyed is True


def test_selftest_covers_tracked_mode_and_settings(tmp_path):
    report = tmp_path / "selftest.txt"
    assert main.selftest(str(report)) == 0
    text = report.read_text(encoding="utf-8")
    assert "tracked changes OK" in text
    assert "rule set and settings OK" in text
