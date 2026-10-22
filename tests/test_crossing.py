"""
Matches that cross an element that is not text.

A picture, a note or comment reference mark, a symbol and a complex field's
begin/separate/end characters add nothing to the paragraph text, so a find can
match straight across one. Applying such a match used to move the element (the
replacement went into the first covered w:t, ahead of it) or split a field
(the new text landed half inside its cached result), and both came back as
applied. Every test here builds the structure for real, applies in place and
as tracked changes, and asserts on the saved document through `layout`, which
shows each non-text element where it sits: the element has not moved, the
text around it is untouched, and the match is skipped with its reason.
"""

import os

import pytest
from docx import Document

import document_processor as dp
from conftest import (accept_all, add_bookmark_end, add_bookmark_start, add_deletion,
                      add_note_reference, add_notes_part, add_run, add_run_with_element,
                      add_simple_field, body_layouts, count_drawings, layout,
                      reject_all, revisions)

MODES = pytest.mark.parametrize("tracked", [False, True], ids=["in_place", "tracked"])


def _scan(path, find, **kw):
    return [m for r in dp.scan_documents(os.path.dirname(path), find, **kw) for m in r.matches]


def _apply(matches, replace=None, tracked=False):
    return dp.apply_changes(matches, replace, create_backups=False,
                            tracked_changes=tracked, author="Tester")


def _final_layout(path, tracked):
    """The first paragraph as Word shows it once the edit is accepted."""
    return accept_all(path, render=layout)[0] if tracked else body_layouts(path)[0]


def _assert_skipped_untouched(path, find, replace, reason, tracked):
    """The scan lists the match; apply skips it, and the file is not rewritten."""
    before = open(path, "rb").read()
    matches = _scan(path, find)
    assert len(matches) == 1
    result = _apply(matches, replace, tracked)
    assert result['total_replaced'] == 0
    assert result['total_skipped'] == 1
    assert result['skip_reasons'] == {reason: 1}
    assert matches[0].applied is False
    assert matches[0].skip_reason == reason
    assert open(path, "rb").read() == before
    assert revisions(path) == []


# ---------------------------------------------------------------- pictures

@MODES
def test_match_across_a_picture_is_skipped_and_the_picture_stays(figure_reference_doc, tracked):
    """'See Figure[picture] 3': "Figure 3" -> "Detail 7" used to give
    'See Detail 7[picture] for riser detail'."""
    assert body_layouts(figure_reference_doc)[0] == "See Figure[drawing] 3 for riser detail"
    _assert_skipped_untouched(figure_reference_doc, "Figure 3", "Detail 7",
                              dp._REASON_CROSSES_OBJECT, tracked)
    assert body_layouts(figure_reference_doc)[0] == "See Figure[drawing] 3 for riser detail"
    assert count_drawings(figure_reference_doc) == 1


@MODES
def test_match_across_a_picture_inside_a_spanned_run_is_skipped(image_in_run_doc, tracked):
    _assert_skipped_untouched(image_in_run_doc, "2022 CBC", "2025 CBC",
                              dp._REASON_CROSSES_OBJECT, tracked)
    assert body_layouts(image_in_run_doc)[0] == "See Figure 20[drawing]22 CBC ref"


@MODES
def test_other_matches_in_the_paragraph_still_apply_around_the_picture(
        figure_reference_doc, tracked):
    """The skip mutates nothing, even when the file is saved for another edit."""
    rules = [dp.Rule("Figure 3", "Detail 7"), dp.Rule("riser", "standpipe")]
    matches = _scan(figure_reference_doc, rules)
    result = _apply(matches, tracked=tracked)
    assert result['total_replaced'] == 1
    assert result['skip_reasons'] == {dp._REASON_CROSSES_OBJECT: 1}
    crossing = next(m for m in matches if m.match_text == "Figure 3")
    assert crossing.applied is False

    assert _final_layout(figure_reference_doc, tracked) == \
        "See Figure[drawing] 3 for standpipe detail"
    if tracked:
        assert reject_all(figure_reference_doc, render=layout)[0] == \
            "See Figure[drawing] 3 for riser detail"
    assert count_drawings(figure_reference_doc) == 1


@MODES
@pytest.mark.parametrize("find, replace, expected", [
    ("Figure", "Detail", "See Detail[drawing] 3 for riser detail"),
    (" 3 for", " 4 for", "See Figure[drawing] 4 for riser detail"),
])
def test_match_ending_or_starting_at_a_picture_applies_and_the_picture_stays(
        figure_reference_doc, tracked, find, replace, expected):
    """A picture at the edge of a match is not inside it."""
    result = _apply(_scan(figure_reference_doc, find), replace, tracked)
    assert result['total_replaced'] == 1 and result['total_skipped'] == 0
    assert _final_layout(figure_reference_doc, tracked) == expected
    if tracked:
        assert reject_all(figure_reference_doc, render=layout)[0] == \
            "See Figure[drawing] 3 for riser detail"


# ---------------------------------------------------------------- complex fields

_FIELD_LAYOUT = "Refer to Article {begin}{separate}2.3{end} above."


@MODES
def test_match_across_the_end_of_a_field_is_skipped(article_field_doc, tracked):
    """"2.3 above" -> "2.4 below" used to write "2.4 below" into the field
    result and delete " above" outside it; the next field update brought
    back "2.3" and lost "below"."""
    assert body_layouts(article_field_doc)[0] == _FIELD_LAYOUT
    _assert_skipped_untouched(article_field_doc, "2.3 above", "2.4 below",
                              dp._REASON_CROSSES_FIELD, tracked)
    assert body_layouts(article_field_doc)[0] == _FIELD_LAYOUT


@MODES
@pytest.mark.parametrize("find", ["Article 2.3", "Article 2.3 above"])
def test_match_across_the_start_of_a_field_is_skipped(article_field_doc, tracked, find):
    _assert_skipped_untouched(article_field_doc, find, "Section 9",
                              dp._REASON_CROSSES_FIELD, tracked)
    assert body_layouts(article_field_doc)[0] == _FIELD_LAYOUT


@MODES
def test_match_wholly_inside_a_field_result_is_still_replaced(article_field_doc, tracked):
    """Today's behaviour, kept: the cached result is edited and the field
    characters stay around it (README "Field results can revert")."""
    result = _apply(_scan(article_field_doc, "2.3"), "2.4", tracked)
    assert result['total_replaced'] == 1 and result['total_skipped'] == 0
    assert _final_layout(article_field_doc, tracked) == \
        "Refer to Article {begin}{separate}2.4{end} above."
    if tracked:
        assert reject_all(article_field_doc, render=layout)[0] == _FIELD_LAYOUT
        assert accept_all(article_field_doc)[0] == "Refer to Article 2.4 above."


@MODES
@pytest.mark.parametrize("find, replace, expected", [
    ("Article ", "Section ", "Refer to Section {begin}{separate}2.3{end} above."),
    (" above.", " below.", "Refer to Article {begin}{separate}2.3{end} below."),
])
def test_match_beside_a_field_applies_outside_it(article_field_doc, tracked,
                                                 find, replace, expected):
    result = _apply(_scan(article_field_doc, find), replace, tracked)
    assert result['total_replaced'] == 1 and result['total_skipped'] == 0
    assert _final_layout(article_field_doc, tracked) == expected
    if tracked:
        assert reject_all(article_field_doc, render=layout)[0] == _FIELD_LAYOUT


@MODES
def test_match_across_a_simple_field_is_skipped(docdir, tracked):
    """A w:fldSimple is not scanned, but it shows its result between the words."""
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Refer to ")
    add_simple_field(p, r' REF _Ref2 \h ', "2.3")
    add_run(p, " above.")
    path = str(docdir / "simple_field.docx")
    doc.save(path)

    _assert_skipped_untouched(path, "to  above", "to above", dp._REASON_CROSSES_FIELD,
                              tracked)
    assert body_layouts(path)[0] == "Refer to [fldSimple] above."


# ---------------------------------------------------------------- reference marks

@MODES
@pytest.mark.parametrize("kind", ["footnote", "endnote", "comment"])
def test_match_across_a_reference_mark_is_skipped(docdir, tracked, kind):
    doc = Document()
    add_notes_part(doc, kind, ["The note."])
    p = doc.add_paragraph()
    add_run(p, "Per NFPA 13")
    add_note_reference(p, kind, 2)
    add_run(p, " requirements.")
    path = str(docdir / f"{kind}_ref.docx")
    doc.save(path)

    _assert_skipped_untouched(path, "13 requirements", "13 criteria",
                              dp._REASON_CROSSES_MARK, tracked)
    assert body_layouts(path)[0] == f"Per NFPA 13[{kind}] requirements."


@MODES
def test_match_ending_at_a_reference_mark_keeps_the_mark_after_it(docdir, tracked):
    doc = Document()
    add_notes_part(doc, "footnote", ["The note."])
    p = doc.add_paragraph()
    add_run(p, "Per NFPA 13")
    add_note_reference(p, "footnote", 2)
    add_run(p, " requirements.")
    path = str(docdir / "footnote_edge.docx")
    doc.save(path)

    result = _apply(_scan(path, "NFPA 13"), "NFPA 13R", tracked)
    assert result['total_replaced'] == 1
    assert _final_layout(path, tracked) == "Per NFPA 13R[footnote] requirements."


# ---------------------------------------------------------------- symbols and the rest

@MODES
def test_match_across_a_symbol_is_skipped(docdir, tracked):
    """Insert > Symbol stores a w:sym, which is not in the text: "68F" would
    otherwise match across the degree sign and move it."""
    doc = Document()
    p = doc.add_paragraph()
    add_run_with_element(p, "Maintain 68", 'sym w:font="Symbol" w:char="F0B0"', "F minimum.")
    path = str(docdir / "symbol.docx")
    doc.save(path)

    _assert_skipped_untouched(path, "68F", "70F", dp._REASON_CROSSES_SYMBOL, tracked)
    assert body_layouts(path)[0] == "Maintain 68[sym]F minimum."


@MODES
def test_match_across_a_tracked_deletion_is_skipped(docdir, tracked):
    """Editing across a pending deletion would move it; Reject All would then
    put the deleted words back in the wrong place."""
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Per the 2022 ")
    add_deletion(p, "edition of the ")
    add_run(p, "CBC.")
    path = str(docdir / "deletion.docx")
    doc.save(path)

    before = open(path, "rb").read()
    matches = _scan(path, "2022 CBC")
    result = _apply(matches, "2025 CBC", tracked)
    assert result['skip_reasons'] == {dp._REASON_CROSSES_DELETION: 1}
    assert open(path, "rb").read() == before
    assert body_layouts(path)[0] == "Per the 2022 [del]CBC."


@MODES
def test_an_unanticipated_element_is_refused_not_moved(docdir, tracked):
    """Anything not text and not a known neutral marker is opaque."""
    doc = Document()
    p = doc.add_paragraph()
    add_run_with_element(p, "Wet-pipe sprin", "softHyphen", "kler system.")
    path = str(docdir / "soft_hyphen.docx")
    doc.save(path)

    _assert_skipped_untouched(path, "sprinkler", "standpipe", dp._REASON_CROSSES_OTHER,
                              tracked)
    assert body_layouts(path)[0] == "Wet-pipe sprin[softHyphen]kler system."


# ---------------------------------------------------------------- neutral markers

@MODES
def test_page_break_cache_and_empty_bookmark_do_not_block_a_match(docdir, tracked):
    """Word scatters w:lastRenderedPageBreak and the empty _GoBack bookmark
    through ordinary text; neither shows anything, so neither is a boundary."""
    doc = Document()
    p = doc.add_paragraph()
    add_run_with_element(p, "Per 2022 ", "lastRenderedPageBreak", "CBC and ")
    add_run(p, "2022 ")
    add_bookmark_start(p, 0, "_GoBack")
    add_bookmark_end(p, 0)
    add_run(p, "CFC.")
    path = str(docdir / "neutral.docx")
    doc.save(path)

    matches = _scan(path, [dp.Rule("2022 CBC", "2025 CBC"), dp.Rule("2022 CFC", "2025 CFC")])
    result = _apply(matches, tracked=tracked)
    assert result['total_replaced'] == 2 and result['total_skipped'] == 0
    if tracked:
        assert accept_all(path)[0] == "Per 2025 CBC and 2025 CFC."
        assert reject_all(path)[0] == "Per 2022 CBC and 2022 CFC."
    else:
        assert body_layouts(path)[0] == "Per 2025 CBC and 2025 CFC[bm][/bm]."
