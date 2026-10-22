"""
Field codes, nested fields and fields that span paragraphs.

Word shows a complex field's result and never its code: everything from the
field's begin to its separate, which includes the cached result of any field
nested inside the code. That nested result is an ordinary w:t, and the
traversal used to read every w:t it reached. With
{ IF { REF SystemType } = "Wet" "..." "..." } the text read "WetWet-pipe ...",
a find for "Wet" matched the REF's result and edited the IF's instruction, and
a match could run from text that is not on the page into text that is.

Every test builds the structure for real. Where something is applied, it is
applied in place and as tracked changes, and the saved document is asserted
through `layout`, which shows each field character where it sits and every
w:t whether it is code or not, so a test can say the code was left exactly as
it was.
"""

import os

import pytest
from docx import Document
from docx.oxml.ns import qn

import document_processor as dp
from conftest import (accept_all, add_complex_field, add_field_begin, add_field_code,
                      add_field_end, add_field_separate, add_run, add_vml_textbox,
                      all_layouts, all_text, body_layouts, layout, nested_field,
                      reject_all, revisions)

MODES = pytest.mark.parametrize("tracked", [False, True], ids=["in_place", "tracked"])


def _scan(path, find, **kw):
    return [m for r in dp.scan_documents(os.path.dirname(path), find, **kw) for m in r.matches]


def _apply(matches, replace=None, tracked=False):
    return dp.apply_changes(matches, replace, create_backups=False,
                            tracked_changes=tracked, author="Tester")


def _final_layouts(path, tracked):
    """Every paragraph as Word shows it once the edit is accepted."""
    return accept_all(path, render=layout) if tracked else body_layouts(path)


def _open_fields(path):
    return [loc.open_fields for loc in dp._locate_paragraphs(Document(path))]


# ---------------------------------------------------------------- a field nested in a code

_IF_LAYOUT = ("Provide {begin}{begin}{separate}Wet{end}{separate}"
              "Wet-pipe system per 2022 CBC{end} throughout.")


def test_a_result_nested_in_a_field_code_is_not_text(nested_if_field_doc):
    assert body_layouts(nested_if_field_doc) == [_IF_LAYOUT]
    assert all_text(nested_if_field_doc) == [
        "Provide Wet-pipe system per 2022 CBC throughout."]


def test_the_scan_finds_only_the_text_word_shows(nested_if_field_doc):
    """The REF's "Wet" used to be found too, at offset 8, ahead of the IF's
    own, and "WetWet" matched from the code into the result."""
    (match,) = _scan(nested_if_field_doc, "Wet")
    assert (match.char_offset, match.context_before, match.context_after) == (
        8, "Provide ", "-pipe system per 2022 CBC throughout.")
    assert _scan(nested_if_field_doc, "WetWet") == []


@MODES
def test_replacing_in_the_result_leaves_the_code_as_it_was(nested_if_field_doc, tracked):
    result = _apply(_scan(nested_if_field_doc, "Wet"), "Dry", tracked)
    assert result['total_replaced'] == 1 and result['total_skipped'] == 0
    assert _final_layouts(nested_if_field_doc, tracked) == [
        "Provide {begin}{begin}{separate}Wet{end}{separate}"
        "Dry-pipe system per 2022 CBC{end} throughout."]
    if tracked:
        assert reject_all(nested_if_field_doc, render=layout) == [_IF_LAYOUT]


def test_everything_in_the_code_is_an_opaque_position_where_the_field_begins(
        nested_if_field_doc):
    """The REF's w:t is still in the walk, as an opaque node marked in_code,
    so nothing can edit it and the node map owns exactly the text."""
    p_el = dp._collect_paragraphs(Document(nested_if_field_doc))[0][0]
    _text, nodes = dp._paragraph_text_and_nodes(p_el)
    assert [n.element.text for n in nodes if n.editable] == [
        "Provide ", "Wet-pipe system per 2022 CBC", " throughout."]
    (ref_result,) = [n for n in nodes if n.element.tag == qn("w:t") and not n.editable]
    assert (ref_result.element.text, ref_result.start, ref_result.length,
            ref_result.opaque, ref_result.in_code) == ("Wet", 8, 0, True, True)
    assert all(n.opaque and n.start == 8 for n in nodes if n.in_code)


@pytest.mark.parametrize("instruction, result, text, hidden", [
    # The case the bug was reported with.
    ([" IF ", nested_field(" REF x ", "a"), ' = "a" "yes" "no" '], "yes", "yes", ["a"]),
    # A field in the code of a field in the code: two levels down.
    ([" IF ", nested_field([" IF ", nested_field(" REF x ", "deep"),
                            ' = "deep" "mid" "" '], "mid"),
      ' = "mid" "top" "" '], "top", "top", ["deep", "mid"]),
    # A field in the code of a field that is itself in another field's result.
    (" QUOTE ", ["Per ", nested_field([" IF ", nested_field(" REF x ", "Wet"),
                                       ' = "Wet" "NFPA 13" "" '], "NFPA 13"), "."],
     "Per NFPA 13.", ["Wet"]),
    # Fields with no separate (index entries), in a code and in a result.
    ([" IF ", nested_field(' XE "code" '), ' 1 = 1 "yes" "" '],
     ["yes", nested_field(' XE "result" '), " too"], "yes too", []),
], ids=["if_ref", "two_levels", "in_a_result", "no_separate"])
def test_only_text_outside_every_field_code_is_read(docdir, instruction, result, text, hidden):
    """Nesting is counted per field, not with one in-code flag: the inner
    field's separate does not end the outer field's code, and a field with no
    separate closes at its end."""
    doc = Document()
    add_complex_field(doc.add_paragraph(), instruction, result)
    path = str(docdir / "nesting.docx")
    doc.save(path)

    assert all_text(path) == [text]
    for word in hidden:
        assert _scan(path, word) == []


def test_a_heading_label_leaves_out_field_code(docdir):
    doc = Document()
    heading = doc.add_heading("", level=1)
    add_complex_field(heading, [" IF ", nested_field(" REF SystemType ", "Wet"),
                                ' = "Wet" "WET-PIPE SYSTEMS" "" '], "WET-PIPE SYSTEMS")
    doc.add_paragraph("Body cites 2022 CBC.")
    path = str(docdir / "heading_field.docx")
    doc.save(path)

    (match,) = _scan(path, "2022 CBC")
    assert match.section == "WET-PIPE SYSTEMS"


# ---------------------------------------------------------------- fields that span paragraphs

_TOC_LAYOUTS = [
    "{begin}{separate}PART 1 GENERAL\t{begin}{separate}1{end}",
    "PART 2 PRODUCTS\t{begin}{separate}3{end}",
    "PART 3 EXECUTION\t{begin}{separate}5{end}{end}",
    "Body text after the table of contents.",
]


def test_a_table_of_contents_is_read_through_every_entry(toc_doc):
    """Entries 2 and 3 begin inside the TOC's result, which is on the page,
    and each nested PAGEREF result is part of it."""
    assert body_layouts(toc_doc) == _TOC_LAYOUTS
    assert _open_fields(toc_doc) == [(), (False,), (False,), ()]
    assert all_text(toc_doc) == ["PART 1 GENERAL\t1", "PART 2 PRODUCTS\t3",
                                 "PART 3 EXECUTION\t5",
                                 "Body text after the table of contents."]


@MODES
def test_table_of_contents_entries_are_still_replaced(toc_doc, tracked):
    """As before: the entry text and a nested page number are results (README
    "Field results can revert")."""
    rules = [dp.Rule("PRODUCTS", "MATERIALS"), dp.Rule("5", "6")]
    result = _apply(_scan(toc_doc, rules), tracked=tracked)
    assert result['total_replaced'] == 2 and result['total_skipped'] == 0
    assert _final_layouts(toc_doc, tracked) == [
        _TOC_LAYOUTS[0],
        "PART 2 MATERIALS\t{begin}{separate}3{end}",
        "PART 3 EXECUTION\t{begin}{separate}6{end}{end}",
        _TOC_LAYOUTS[3]]
    if tracked:
        assert reject_all(toc_doc, render=layout) == _TOC_LAYOUTS


_ACROSS_LAYOUTS = [
    "Sprinklers: {begin}{begin}{separate}Wet{end}",
    "{begin}{separate}NFPA 13{end}",
    "{begin}{separate}NFPA 13{end}{separate}Provide wet-pipe system.",
    "Coordinate with {begin}{separate}NFPA 13{end}.",
    "Test per {begin}{separate}NFPA 13{end}.{end} Field ends here.",
]


def test_a_paragraph_that_begins_inside_a_field_code_is_read_as_code(
        field_code_across_paragraphs_doc):
    """Paragraph 2 lies wholly in the IF's code and paragraph 3 begins there,
    so the REF results in them are not on the page. Read one paragraph at a
    time, both would show "NFPA 13"."""
    path = field_code_across_paragraphs_doc
    assert body_layouts(path) == _ACROSS_LAYOUTS
    assert _open_fields(path) == [(), (True,), (True,), (False,), (False,)]
    assert all_text(path) == ["Sprinklers: ", "", "Provide wet-pipe system.",
                              "Coordinate with NFPA 13.",
                              "Test per NFPA 13. Field ends here."]


def test_the_scan_skips_results_nested_in_a_code_that_spans_paragraphs(
        field_code_across_paragraphs_doc):
    path = field_code_across_paragraphs_doc
    assert [(m.location_detail, m.char_offset) for m in _scan(path, "NFPA 13")] == [
        ("Paragraph 4", 16), ("Paragraph 5", 9)]
    assert _scan(path, "Wet") == []


@MODES
def test_a_table_inside_a_field_code_is_not_read(docdir, tracked):
    """Word lets an IF's text hold a table, and its result then holds a copy.
    A table is part of its story, so its cells begin in the field state of
    the paragraph before it: the copy in the code is not on the page."""
    doc = Document()
    p1 = doc.add_paragraph()
    add_run(p1, "Before ")
    add_field_begin(p1, ' IF 1 = 1 "')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Code table cites 2022 CBC"
    p2 = doc.add_paragraph()
    add_field_code(p2, '" "" ')
    add_field_separate(p2, "")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Result table cites 2022 CBC"
    p3 = doc.add_paragraph()
    add_field_end(p3)
    add_run(p3, " after.")
    path = str(docdir / "table_in_code.docx")
    doc.save(path)

    assert all_text(path) == ["Before ", "", "", "Result table cites 2022 CBC", " after."]
    assert _open_fields(path) == [(), (True,), (True,), (False,), (False,)]
    matches = _scan(path, "2022 CBC")
    assert [m.location_detail for m in matches] == ["Table 2, Row 1, Cell 1"]
    result = _apply(matches, "2025 CBC", tracked)
    assert result['total_replaced'] == 1 and result['total_skipped'] == 0

    shown = "Result table cites [del]2025 CBC" if tracked else "Result table cites 2025 CBC"
    assert all_layouts(path) == ["Before {begin}", "Code table cites 2022 CBC",
                                 "{separate}", shown, "{end} after."]
    if tracked:
        assert accept_all(path)[3] == "Result table cites 2025 CBC"
        assert reject_all(path)[3] == "Result table cites 2022 CBC"


@MODES
def test_replacing_after_a_field_code_that_spans_paragraphs(
        field_code_across_paragraphs_doc, tracked):
    """Paragraph 3 begins inside the code, so its offsets only line up with
    the scan's when apply also starts it there."""
    path = field_code_across_paragraphs_doc
    rules = [dp.Rule("NFPA 13", "NFPA 13R"), dp.Rule("wet-pipe", "dry-pipe")]
    result = _apply(_scan(path, rules), tracked=tracked)
    assert result['total_replaced'] == 3 and result['total_skipped'] == 0
    assert _final_layouts(path, tracked) == [
        _ACROSS_LAYOUTS[0],
        _ACROSS_LAYOUTS[1],
        "{begin}{separate}NFPA 13{end}{separate}Provide dry-pipe system.",
        "Coordinate with {begin}{separate}NFPA 13R{end}.",
        "Test per {begin}{separate}NFPA 13R{end}.{end} Field ends here."]
    if tracked:
        assert reject_all(path, render=layout) == _ACROSS_LAYOUTS


# ---------------------------------------------------------------- text boxes

@MODES
def test_a_text_box_inside_a_field_code_is_not_collected(docdir, tracked):
    """A drawing in an IF's text is part of the instruction; Word shows the
    copy in the result. Only that copy is scanned and changed."""
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "See ")
    add_field_begin(p, ' IF 1 = 1 "')
    add_vml_textbox(p, "Code copy per 2022 CBC")
    add_field_code(p, '" "" ')
    add_field_separate(p, "")
    add_vml_textbox(p, "Shown copy per 2022 CBC")
    add_field_end(p)
    path = str(docdir / "textbox_in_code.docx")
    doc.save(path)

    assert all_text(path) == ["See ", "Shown copy per 2022 CBC"]
    matches = _scan(path, "2022 CBC")
    assert [m.location_detail for m in matches] == ["Paragraph 1, Text Box"]
    result = _apply(matches, "2025 CBC", tracked)
    assert result['total_replaced'] == 1 and result['total_skipped'] == 0

    shown = "Shown copy per [del]2025 CBC" if tracked else "Shown copy per 2025 CBC"
    assert all_layouts(path) == ["See {begin}[pict]{separate}[pict]{end}",
                                 "Code copy per 2022 CBC", shown]
    if tracked:
        assert [(kind, text) for kind, _id, _a, text in revisions(path)] == [
            ("del", "2022 CBC"), ("ins", "2025 CBC")]
        assert accept_all(path)[1] == "Shown copy per 2025 CBC"
        assert reject_all(path)[1] == "Shown copy per 2022 CBC"


def test_a_text_box_starts_with_no_field_open(docdir):
    """A text box is a story of its own. This one is walked after its
    paragraph, where the IF is still open, but must not inherit it."""
    doc = Document()
    p1 = doc.add_paragraph()
    add_run(p1, "Intro ")
    add_vml_textbox(p1, "Boxed per 2022 CBC")
    add_field_begin(p1, ' IF 1 = 1 "first')
    p2 = doc.add_paragraph()
    add_field_code(p2, 'second" "" ')
    add_field_separate(p2, "Shown per 2022 CBC")
    add_field_end(p2)
    path = str(docdir / "textbox_beside_field.docx")
    doc.save(path)

    assert all_text(path) == ["Intro ", "Boxed per 2022 CBC", "Shown per 2022 CBC"]
    assert _open_fields(path) == [(), (), (True,)]
    assert [m.location_detail for m in _scan(path, "2022 CBC")] == [
        "Paragraph 1, Text Box", "Paragraph 2"]
