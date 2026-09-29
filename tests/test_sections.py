"""
Section context: the heading path recorded on each body paragraph.
"""

import os

from docx import Document

import document_processor as dp
from conftest import (add_custom_paragraph_style, add_notes_part,
                      add_paragraph_with_outline_level, add_styled_paragraph,
                      add_vml_textbox, drop_styles_part, has_styles_part)


def _sections(path):
    """[(text, section)] for every paragraph in the document."""
    doc = Document(path)
    return [(dp.paragraph_text(loc.element), loc.section)
            for loc in dp._locate_paragraphs(doc)]


def _scan(path, find):
    return [m for r in dp.scan_documents(os.path.dirname(path), find) for m in r.matches]


def test_body_paragraph_records_the_headings_above_it(heading_doc):
    matches = _scan(heading_doc, "2022 CBC")
    by_text = {m.context_before.strip(): m.section for m in matches}
    assert by_text["Intro cites"] == "GENERAL"
    assert by_text["Body cites"] == "GENERAL > REFERENCES"


def test_table_cell_inherits_the_current_path(heading_doc):
    matches = _scan(heading_doc, "2022 CBC")
    cell = [m for m in matches if m.location_type == 'table'][0]
    assert cell.section == "GENERAL > REFERENCES"


def test_a_heading_records_the_path_including_itself(heading_doc):
    sections = dict(_sections(heading_doc))
    assert sections["GENERAL"] == "GENERAL"
    assert sections["REFERENCES"] == "GENERAL > REFERENCES"


def test_a_later_top_level_heading_pops_the_lower_level(heading_doc):
    matches = _scan(heading_doc, "2022 CBC")
    last = [m for m in matches if m.context_before.startswith("Later")][0]
    assert last.section == "PRODUCTS"


def test_text_before_any_heading_has_no_section(docdir):
    doc = Document()
    doc.add_paragraph("Preamble 2022 CBC")
    doc.add_heading("GENERAL", level=1)
    path = str(docdir / "pre.docx")
    doc.save(path)
    assert _scan(path, "2022 CBC")[0].section == ""


def test_explicit_outline_level_overrides_the_style(docdir):
    doc = Document()
    add_paragraph_with_outline_level(doc, "Part One", 0)
    add_paragraph_with_outline_level(doc, "Article", 1)
    doc.add_paragraph("Under 2022 CBC")
    path = str(docdir / "lvl.docx")
    doc.save(path)
    assert _scan(path, "2022 CBC")[0].section == "Part One > Article"


def test_paragraph_level_nine_is_body_text(docdir):
    doc = Document()
    doc.add_heading("GENERAL", level=1)
    p = doc.add_heading("Not a heading", level=2)
    p._p.get_or_add_pPr().append(
        Document().add_paragraph()._p.makeelement(dp._w('outlineLvl'), {dp._w('val'): "9"}))
    doc.add_paragraph("Under 2022 CBC")
    path = str(docdir / "nine.docx")
    doc.save(path)
    assert _scan(path, "2022 CBC")[0].section == "GENERAL"


def test_custom_style_inherits_level_through_based_on(docdir):
    doc = Document()
    add_custom_paragraph_style(doc, "ART", "Article", based_on="Heading2")
    doc.add_heading("GENERAL", level=1)
    add_styled_paragraph(doc, "SUMMARY", "ART")
    doc.add_paragraph("Under 2022 CBC")
    path = str(docdir / "based.docx")
    doc.save(path)
    assert _scan(path, "2022 CBC")[0].section == "GENERAL > SUMMARY"


def test_style_with_own_level_nine_is_not_a_heading(docdir):
    # Word's TOC Heading is based on Heading 1 with an explicit level 9.
    doc = Document()
    add_custom_paragraph_style(doc, "TOCHeading", "TOC Heading",
                               based_on="Heading1", outline_level=9)
    add_styled_paragraph(doc, "Contents", "TOCHeading")
    doc.add_paragraph("Under 2022 CBC")
    path = str(docdir / "toc.docx")
    doc.save(path)
    assert _scan(path, "2022 CBC")[0].section == ""


def test_style_named_heading_n_without_a_level_falls_back_to_its_name(docdir):
    doc = Document()
    add_custom_paragraph_style(doc, "MyH3", "heading 3")
    add_styled_paragraph(doc, "Third", "MyH3")
    doc.add_paragraph("Under 2022 CBC")
    path = str(docdir / "named.docx")
    doc.save(path)
    assert _scan(path, "2022 CBC")[0].section == "Third"


def test_custom_style_with_its_own_level(docdir):
    doc = Document()
    add_custom_paragraph_style(doc, "PRT", "PRT", outline_level=0)
    add_custom_paragraph_style(doc, "ART", "ART", outline_level=1)
    add_styled_paragraph(doc, "GENERAL", "PRT")
    add_styled_paragraph(doc, "REFERENCES", "ART")
    doc.add_paragraph("Under 2022 CBC")
    add_styled_paragraph(doc, "PRODUCTS", "PRT")
    doc.add_paragraph("Later 2022 CBC")
    path = str(docdir / "masterspec.docx")
    doc.save(path)
    assert [m.section for m in _scan(path, "2022 CBC")] == ["GENERAL > REFERENCES", "PRODUCTS"]


def test_heading_label_collapses_whitespace_and_truncates(docdir):
    doc = Document()
    doc.add_heading("1.3\tREFERENCES   AND\nSTANDARDS", level=1)
    doc.add_heading("x" * 100, level=2)
    doc.add_paragraph("Under 2022 CBC")
    path = str(docdir / "label.docx")
    doc.save(path)
    section = _scan(path, "2022 CBC")[0].section
    first, second = section.split(" > ")
    assert first == "1.3 REFERENCES AND STANDARDS"
    assert len(second) == 60 and second.endswith("…")


def test_empty_heading_paragraphs_are_ignored(docdir):
    doc = Document()
    doc.add_heading("GENERAL", level=1)
    doc.add_heading("", level=1)
    doc.add_paragraph("Under 2022 CBC")
    path = str(docdir / "emptyh.docx")
    doc.save(path)
    assert _scan(path, "2022 CBC")[0].section == "GENERAL"


def test_headers_and_notes_have_no_section(docdir):
    doc = Document()
    doc.add_heading("GENERAL", level=1)
    doc.add_paragraph("Body 2022 CBC")
    doc.sections[0].header.paragraphs[0].text = "Header 2022 CBC"
    add_notes_part(doc, "footnote", ["Note 2022 CBC"])
    path = str(docdir / "hn.docx")
    doc.save(path)
    by_type = {m.location_type: m.section for m in _scan(path, "2022 CBC")}
    assert by_type == {'body': "GENERAL", 'header': "", 'footnote': ""}


def test_textbox_inherits_but_does_not_push(docdir):
    doc = Document()
    doc.add_heading("GENERAL", level=1)
    p = doc.add_paragraph("Anchor")
    add_vml_textbox(p, "Box 2022 CBC")
    doc.add_paragraph("After 2022 CBC")
    path = str(docdir / "tb.docx")
    doc.save(path)
    assert [m.section for m in _scan(path, "2022 CBC")] == ["GENERAL", "GENERAL"]


def test_scanning_a_document_without_styles_does_not_add_a_styles_part(docdir):
    doc = Document()
    doc.add_paragraph("Bare 2022 CBC")
    drop_styles_part(doc)
    path = str(docdir / "nostyles.docx")
    doc.save(path)
    assert not has_styles_part(path)

    matches = _scan(path, "2022 CBC")
    assert len(matches) == 1 and matches[0].section == ""
    dp.apply_changes(matches, "2025 CBC", create_backups=False)
    assert not has_styles_part(path)


def test_collect_paragraphs_keys_match_locate_paragraphs(heading_doc):
    doc = Document(heading_doc)
    four = dp._collect_paragraphs(doc)
    five = dp._locate_paragraphs(doc)
    assert [(k) for _p, _t, _d, k in four] == [loc.key for loc in five]
    assert len(four) == 7


def test_outline_levels_ignore_cycles_and_unknown_bases(docdir):
    doc = Document()
    add_custom_paragraph_style(doc, "A", "A", based_on="B")
    add_custom_paragraph_style(doc, "B", "B", based_on="A")
    add_custom_paragraph_style(doc, "C", "C", based_on="Missing")
    levels = dp._outline_levels_by_style(doc.styles.element)
    assert levels["A"] is None and levels["B"] is None and levels["C"] is None
    assert levels["Heading1"] == 0 and levels["Heading9"] == 8
