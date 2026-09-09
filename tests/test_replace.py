"""
Replacement: the paragraph-level edit that has to preserve everything it is not
explicitly changing.
"""

import pytest
from docx import Document
from docx.oxml.ns import qn

import document_processor as dp
from conftest import add_run, add_hyperlink, count_drawings


def _para(path, predicate=lambda t: True):
    doc = Document(path)
    for p_el, _t, _d, _k in dp._collect_paragraphs(doc):
        if predicate(dp.paragraph_text(p_el)):
            return doc, p_el
    raise AssertionError("no matching paragraph")


def test_replaces_inside_a_hyperlink_at_the_right_offset(hyperlink_doc):
    """
    The regression that ate unrelated prose: offsets computed over text that
    included the hyperlink, applied to a run map that excluded it.
    """
    doc, p_el = _para(hyperlink_doc)
    applied, skipped = dp._replace_in_paragraph(p_el, "2022 CBC", "2025 CBC", [4])
    assert applied == [4]
    assert skipped == []
    assert dp.paragraph_text(p_el) == "See 2025 CBC link and done."


def test_hyperlink_relationship_survives(hyperlink_doc, tmp_path):
    doc, p_el = _para(hyperlink_doc)
    dp._replace_in_paragraph(p_el, "2022 CBC", "2025 CBC", [4])
    out = tmp_path / "out.docx"
    doc.save(str(out))

    reloaded = Document(str(out))
    p = reloaded.element.body.findall(".//" + qn("w:hyperlink"))
    assert len(p) == 1
    assert p[0].get(qn("r:id"))


def test_cross_run_match_preserves_inline_image(image_in_run_doc, tmp_path):
    """
    Assigning Run.text calls clear_content(), which strips every child of the
    run including the image. Editing w:t text in place does not.
    """
    assert count_drawings(image_in_run_doc) == 1

    doc, p_el = _para(image_in_run_doc)
    applied, _ = dp._replace_in_paragraph(p_el, "2022 CBC", "2025 CBC", [11])
    assert applied == [11]

    out = tmp_path / "out.docx"
    doc.save(str(out))
    assert dp.paragraph_text(Document(str(out)).element.body.findall(".//" + qn("w:p"))[0]) \
        == "See Figure 2025 CBC ref"
    assert count_drawings(str(out)) == 1


def test_cross_run_match_preserves_run_formatting(split_run_doc, tmp_path):
    doc, p_el = _para(split_run_doc)
    dp._replace_in_paragraph(p_el, "2022 CBC", "2025 CBC", [14])
    assert dp.paragraph_text(p_el) == "Reference the 2025 CBC for compliance."
    # The three original runs are still three runs; none was emptied away.
    assert len(p_el.findall(qn("w:r"))) == 3


def test_replacement_inside_a_single_run(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Per the 2022 CBC requirements.")
    path = docdir / "single.docx"
    doc.save(str(path))

    doc2, p_el = _para(str(path))
    assert dp.paragraph_text(p_el).index("2022 CBC") == 8
    dp._replace_in_paragraph(p_el, "2022 CBC", "2025 CBC", [8])
    assert dp.paragraph_text(p_el) == "Per the 2025 CBC requirements."


def test_multiple_occurrences_in_one_paragraph(repeated_doc):
    doc, p_el = _para(repeated_doc)
    text = dp.paragraph_text(p_el)
    offsets = dp._find_all_occurrences(text, "2022", True)
    applied, skipped = dp._replace_in_paragraph(p_el, "2022", "2025", offsets)
    assert sorted(applied) == offsets
    assert skipped == []
    assert dp.paragraph_text(p_el) == "aaaa and 2025 2025 2025"


def test_selective_replacement_leaves_others_alone(repeated_doc):
    """Only the offsets handed in are touched — the whole point of the preview."""
    doc, p_el = _para(repeated_doc)
    offsets = dp._find_all_occurrences(dp.paragraph_text(p_el), "2022", True)
    dp._replace_in_paragraph(p_el, "2022", "2025", [offsets[1]])
    assert dp.paragraph_text(p_el) == "aaaa and 2022 2025 2022"


def test_growing_replacement_is_not_applied_twice(repeated_doc):
    """A replacement containing the search text must not compound."""
    doc, p_el = _para(repeated_doc)
    offsets = dp._find_all_occurrences(dp.paragraph_text(p_el), "2022", True)
    dp._replace_in_paragraph(p_el, "2022", "2022 and 2025", offsets)
    assert dp.paragraph_text(p_el) == (
        "aaaa and 2022 and 2025 2022 and 2025 2022 and 2025"
    )


def test_stale_offset_is_skipped_not_misapplied(repeated_doc):
    """An offset that no longer holds the search text is reported, not guessed at."""
    doc, p_el = _para(repeated_doc)
    applied, skipped = dp._replace_in_paragraph(p_el, "2022", "2025", [0, 9])
    assert applied == [9]
    assert skipped == [0]
    assert dp.paragraph_text(p_el) == "aaaa and 2025 2022 2022"


def test_offset_past_end_of_paragraph_is_skipped(repeated_doc):
    doc, p_el = _para(repeated_doc)
    applied, skipped = dp._replace_in_paragraph(p_el, "2022", "2025", [9999])
    assert applied == []
    assert skipped == [9999]


def test_case_insensitive_replacement(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "the cbc and the CBC")
    path = docdir / "case.docx"
    doc.save(str(path))

    doc2, p_el = _para(str(path))
    offsets = dp._find_all_occurrences(dp.paragraph_text(p_el), "cbc", False)
    applied, skipped = dp._replace_in_paragraph(p_el, "cbc", "IBC", offsets, case_sensitive=False)
    assert len(applied) == 2
    assert skipped == []
    assert dp.paragraph_text(p_el) == "the IBC and the IBC"


def test_leading_and_trailing_spaces_are_preserved(docdir):
    """
    Word strips whitespace from a w:t without xml:space="preserve".
    Any node we write has to carry it.
    """
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "start 2022 end")
    path = docdir / "space.docx"
    doc.save(str(path))

    doc2, p_el = _para(str(path))
    dp._replace_in_paragraph(p_el, "2022", "  spaced  ", [6])
    out = docdir / "space_out.docx"
    doc2.save(str(out))

    reloaded_doc = Document(str(out))
    p2 = dp._collect_paragraphs(reloaded_doc)[0][0]
    assert dp.paragraph_text(p2) == "start   spaced   end"


def test_tracked_insertion_can_be_replaced(tracked_changes_doc):
    doc, p_el = _para(tracked_changes_doc)
    text = dp.paragraph_text(p_el)
    offset = text.index("2022 CBC")
    applied, _ = dp._replace_in_paragraph(p_el, "2022 CBC", "2025 CBC", [offset])
    assert applied == [offset]
    assert "2025 CBC inserted" in dp.paragraph_text(p_el)


def test_deleted_text_is_never_touched(tracked_changes_doc, tmp_path):
    doc, p_el = _para(tracked_changes_doc)
    offset = dp.paragraph_text(p_el).index("2022 CBC")
    dp._replace_in_paragraph(p_el, "2022 CBC", "2025 CBC", [offset])

    out = tmp_path / "out.docx"
    doc.save(str(out))
    del_texts = [e.text for e in
                 Document(str(out)).element.body.findall(".//" + qn("w:delText"))]
    assert del_texts == [" 2022 CBC deleted"]


def test_empty_find_text_replaces_nothing(repeated_doc):
    doc, p_el = _para(repeated_doc)
    before = dp.paragraph_text(p_el)
    applied, skipped = dp._replace_in_paragraph(p_el, "", "x", [0])
    assert applied == []
    assert skipped == [0]
    assert dp.paragraph_text(p_el) == before
