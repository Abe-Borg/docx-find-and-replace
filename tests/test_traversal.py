"""
Traversal: paragraph text extraction and paragraph collection.

These cover the root cause behind the corruption bugs — text, offsets and
paragraph identity all coming from one walk instead of three disagreeing ones.
"""

import os

import pytest
from docx import Document
from docx.oxml.ns import qn

import document_processor as dp
from conftest import (add_bookmark_end, add_bookmark_start, add_hyperlink, add_note_reference,
                      add_notes_part, add_run, add_run_with_element, add_runs)


# ----------------------------------------------------- text extraction

def test_hyperlink_text_is_included(hyperlink_doc):
    """Hyperlink text is part of the paragraph, so it must be in the text."""
    doc = Document(hyperlink_doc)
    p_el = dp._collect_paragraphs(doc)[0][0]
    assert dp.paragraph_text(p_el) == "See 2022 CBC link and done."


@pytest.mark.parametrize("fixture", ["hyperlink_doc", "figure_reference_doc",
                                     "article_field_doc", "nested_if_field_doc"])
def test_text_and_node_map_are_built_from_the_same_walk(request, fixture):
    """
    Every character of the extracted text must be owned by a node in the map.

    This is the invariant that failed before: text came from Paragraph.text
    (which includes hyperlinks) while the offset map came from Paragraph.runs
    (which does not), so replacements landed at the wrong offset. Opaque nodes
    (a picture, field characters, anything in a field's code) own no
    characters and keep it intact.
    """
    doc = Document(request.getfixturevalue(fixture))
    p_el = dp._collect_paragraphs(doc)[0][0]
    text, nodes = dp._paragraph_text_and_nodes(p_el)

    rebuilt = "".join(
        (n.element.text or "") if n.editable else text[n.start:n.start + n.length]
        for n in nodes
    )
    assert rebuilt == text
    assert sum(n.length for n in nodes) == len(text)
    # Nodes are contiguous and in order.
    cursor = 0
    for node in nodes:
        assert node.start == cursor
        cursor += node.length


def test_tracked_insertion_is_visible_deletion_is_not(tracked_changes_doc):
    doc = Document(tracked_changes_doc)
    p_el = dp._collect_paragraphs(doc)[0][0]
    text = dp.paragraph_text(p_el)
    assert "2022 CBC inserted" in text
    assert "deleted" not in text


def test_tabs_and_breaks_become_characters(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "A")
    p._p.append(_tab_element())
    add_run(p, "B")
    path = docdir / "tabs.docx"
    doc.save(str(path))

    reloaded = Document(str(path))
    p_el = dp._collect_paragraphs(reloaded)[0][0]
    assert dp.paragraph_text(p_el) == "A\tB"


def test_a_picture_is_an_opaque_position_that_adds_no_text(figure_reference_doc):
    doc = Document(figure_reference_doc)
    p_el = dp._collect_paragraphs(doc)[0][0]
    text, nodes = dp._paragraph_text_and_nodes(p_el)
    assert text == "See Figure 3 for riser detail"
    assert [(n.element.tag, n.start, n.length, n.editable) for n in nodes if n.opaque] == [
        (qn("w:drawing"), 10, 0, False)]


def test_field_characters_are_opaque_positions_around_the_result(article_field_doc):
    """begin, the field code and separate sit where the result starts; end
    where it stops. The field code itself is still never read as text."""
    doc = Document(article_field_doc)
    p_el = dp._collect_paragraphs(doc)[0][0]
    text, nodes = dp._paragraph_text_and_nodes(p_el)
    assert text == "Refer to Article 2.3 above."
    assert [(n.element.tag, n.start) for n in nodes if n.opaque] == [
        (qn("w:fldChar"), 17), (qn("w:instrText"), 17),
        (qn("w:fldChar"), 17), (qn("w:fldChar"), 20)]


def test_properties_and_markers_are_not_positions(docdir):
    """Run properties, bookmarks and the page-break cache show nothing; only
    the footnote reference mark is opaque."""
    doc = Document()
    add_notes_part(doc, "footnote", ["The note."])
    p = doc.add_paragraph()
    add_bookmark_start(p, 0, "_GoBack")
    add_bookmark_end(p, 0)
    add_run_with_element(p, "Per NFPA", "lastRenderedPageBreak", " 13")
    add_note_reference(p, "footnote", 2)     # its run carries a w:rPr
    path = docdir / "markers.docx"
    doc.save(str(path))

    p_el = dp._collect_paragraphs(Document(str(path)))[0][0]
    text, nodes = dp._paragraph_text_and_nodes(p_el)
    assert text == "Per NFPA 13"
    assert [(n.element.tag, n.start) for n in nodes if n.opaque] == [
        (qn("w:footnoteReference"), 11)]


def test_context_around_a_picture_has_no_stray_character(figure_reference_doc):
    """Opaque positions add nothing to the text, so nothing leaks into the
    context the results tree shows."""
    results = dp.scan_documents(os.path.dirname(figure_reference_doc), "Figure 3")
    (match,) = [m for r in results for m in r.matches]
    assert (match.context_before, match.match_text, match.context_after) == (
        "See ", "Figure 3", " for riser detail")
    assert match.display_context == "See [Figure 3] for riser detail"


def _tab_element():
    from docx.oxml import parse_xml
    return parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:tab/></w:r>'
    )


# ----------------------------------------------------- paragraph identity

def test_merged_cell_is_collected_once(merged_and_linked_doc):
    """
    row.cells yields the same merged cell once per spanned grid column.
    Walking the XML visits the physical <w:tc> exactly once.
    """
    doc = Document(merged_and_linked_doc)
    paragraphs = dp._collect_paragraphs(doc)
    merged = [p for p, _t, _d, _k in paragraphs if "Merged:" in dp.paragraph_text(p)]
    assert len(merged) == 1


def test_linked_header_is_collected_once(merged_and_linked_doc):
    """One header part shared by three linked sections is one paragraph."""
    doc = Document(merged_and_linked_doc)
    paragraphs = dp._collect_paragraphs(doc)
    headers = [p for p, t, _d, _k in paragraphs if t == 'header']
    assert len([p for p in headers if "HEADER" in dp.paragraph_text(p)]) == 1


def test_paragraph_keys_are_unique(merged_and_linked_doc):
    doc = Document(merged_and_linked_doc)
    keys = [k for _p, _t, _d, k in dp._collect_paragraphs(doc)]
    assert len(keys) == len(set(keys))


def test_paragraph_keys_are_stable_across_reloads(merged_and_linked_doc):
    """
    Scanning and applying open the file separately, so identity must survive a
    reload. Python object ids do not; structural keys do.
    """
    first = {k: dp.paragraph_text(p) for p, _t, _d, k in
             dp._collect_paragraphs(Document(merged_and_linked_doc))}
    second = {k: dp.paragraph_text(p) for p, _t, _d, k in
              dp._collect_paragraphs(Document(merged_and_linked_doc))}
    assert first == second
    assert first  # non-empty, so the comparison means something


def test_collecting_headers_does_not_mutate_the_document(docdir):
    """
    Reading section.header creates an empty header part as a side effect.
    Collection must not change the file it is only reading.
    """
    doc = Document()
    doc.add_paragraph("No header here")
    path = docdir / "plain.docx"
    doc.save(str(path))
    before = os.path.getsize(str(path))

    reloaded = Document(str(path))
    dp._collect_paragraphs(reloaded)
    resaved = docdir / "plain_resaved.docx"
    reloaded.save(str(resaved))

    part_names = {str(p.partname) for p in Document(str(resaved)).part.package.iter_parts()}
    assert not any("header" in n for n in part_names)
    assert before > 0


# ----------------------------------------------------- coverage

def test_nested_table_is_found(nested_table_doc):
    doc = Document(nested_table_doc)
    texts = [dp.paragraph_text(p) for p, _t, _d, _k in dp._collect_paragraphs(doc)]
    assert any("Nested 2022 CBC" in t for t in texts)


def test_textbox_is_found(textbox_doc):
    doc = Document(textbox_doc)
    texts = [dp.paragraph_text(p) for p, _t, _d, _k in dp._collect_paragraphs(doc)]
    assert any("Textbox 2022 CBC" in t for t in texts)


def test_alternate_content_textbox_counted_once(alt_content_doc):
    """
    Word stores a text box twice (mc:Choice + mc:Fallback). One visible text box
    must report one match, not two.
    """
    doc = Document(alt_content_doc)
    texts = [dp.paragraph_text(p) for p, _t, _d, _k in dp._collect_paragraphs(doc)]
    assert sum(1 for t in texts if "Choice 2022 CBC" in t) == 1


def test_textbox_anchor_paragraph_excludes_textbox_text(textbox_doc):
    """Text box content is its own paragraph, not part of the anchor's text."""
    doc = Document(textbox_doc)
    paragraphs = dp._collect_paragraphs(doc)
    anchor = next(p for p, _t, _d, _k in paragraphs
                  if "Anchor paragraph." in dp.paragraph_text(p))
    assert dp.paragraph_text(anchor) == "Anchor paragraph."


def test_body_table_and_header_location_types(merged_and_linked_doc):
    doc = Document(merged_and_linked_doc)
    types = {t for _p, t, _d, _k in dp._collect_paragraphs(doc)}
    assert {'body', 'table', 'header'} <= types
