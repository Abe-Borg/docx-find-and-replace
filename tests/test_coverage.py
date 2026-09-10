"""
Content classes that live in their own document parts, and the guards that
protect the preview's offsets from a document edited underneath them.
"""

import os
import threading

import pytest
from docx import Document

import document_processor as dp
from conftest import all_text


def _scan(path, find, case=True):
    results = dp.scan_documents(os.path.dirname(path), find, case)
    return [m for r in results for m in r.matches]


# ----------------------------------------------------- notes coverage

def test_footnotes_endnotes_and_comments_are_scanned(notes_doc):
    matches = _scan(notes_doc, "2022 CBC")
    kinds = sorted(m.location_type for m in matches)
    assert kinds == ['body', 'comment', 'endnote', 'footnote']


def test_footnote_separator_is_not_reported(notes_doc):
    """Word's separator entries are rule lines, not content."""
    doc = Document(notes_doc)
    details = [d for _p, t, d, _k in dp._collect_paragraphs(doc) if t == 'footnote']
    assert details == ['Footnote 2']


def test_notes_are_replaced_and_the_file_still_opens(notes_doc):
    matches = _scan(notes_doc, "2022 CBC")
    result = dp.apply_changes(matches, "2025 CBC", create_backups=False)

    assert result['total_replaced'] == 4
    assert result['errors'] == []

    texts = all_text(notes_doc)
    assert any("Footnote cites 2025 CBC" in t for t in texts)
    assert any("Endnote cites 2025 CBC" in t for t in texts)
    assert any("Reviewer asks about 2025 CBC" in t for t in texts)
    Document(notes_doc)  # raises if the saved package is invalid


def test_note_paragraph_keys_are_stable_across_reloads(notes_doc):
    first = {k: dp.paragraph_text(p) for p, _t, _d, k in
             dp._collect_paragraphs(Document(notes_doc))}
    second = {k: dp.paragraph_text(p) for p, _t, _d, k in
              dp._collect_paragraphs(Document(notes_doc))}
    assert first == second
    assert any('footnotes.xml' in k for k in first)


def test_note_locations_are_labelled_by_id(notes_doc):
    doc = Document(notes_doc)
    details = {d for _p, t, d, _k in dp._collect_paragraphs(doc) if t != 'body'}
    assert 'Footnote 2' in details
    assert 'Endnote 2' in details
    assert 'Comment 2' in details


# ----------------------------------------------------- staleness guard

def test_file_edited_after_preview_is_refused(repeated_doc):
    """
    Offsets recorded during the preview describe one version of a file. If the
    document changes underneath them they can still validate against different
    text, so the file is refused rather than edited blind.
    """
    matches = _scan(repeated_doc, "2022")

    doc = Document(repeated_doc)
    doc.add_paragraph("Someone edited this in Word after the preview.")
    doc.save(repeated_doc)

    result = dp.apply_changes(matches, "2025", create_backups=False)

    assert result['total_replaced'] == 0
    assert result['total_skipped'] == len(matches)
    assert len(result['errors']) == 1
    assert "re-run Preview Changes" in result['errors'][0]
    assert all(m.applied is False for m in matches)
    assert "2022 2022 2022" in all_text(repeated_doc)[0]


def test_unchanged_file_is_not_refused(repeated_doc):
    result = dp.apply_changes(_scan(repeated_doc, "2022"), "2025", create_backups=False)
    assert result['total_replaced'] == 3
    assert result['errors'] == []


def test_no_backup_is_written_for_a_refused_file(repeated_doc):
    import glob
    matches = _scan(repeated_doc, "2022")
    doc = Document(repeated_doc)
    doc.add_paragraph("edited")
    doc.save(repeated_doc)

    dp.apply_changes(matches, "2025", create_backups=True)
    assert glob.glob(repeated_doc + ".*.bak") == []


def test_hand_built_match_without_stat_is_not_checked(repeated_doc):
    """A Match with no recorded stat (e.g. built in a test) skips the check."""
    match = dp.Match(
        file_path=repeated_doc, location_type="body", location_detail="Paragraph 1",
        paragraph_key="/word/document.xml#0", char_offset=9,
        context_before="", match_text="2022", context_after="",
    )
    result = dp.apply_changes([match], "2025", create_backups=False)
    assert result['total_replaced'] == 1


# ----------------------------------------------------- progress and cancel

def test_scan_reports_progress_per_file(docdir):
    for name in ("a.docx", "b.docx", "c.docx"):
        d = Document()
        d.add_paragraph("2022 CBC")
        d.save(str(docdir / name))

    seen = []
    results, summary = dp.scan_documents_detailed(
        str(docdir), "2022", True,
        progress_callback=lambda name, i, total: seen.append((name, i, total)),
    )
    assert [s[0] for s in seen] == ["a.docx", "b.docx", "c.docx"]
    assert [s[1] for s in seen] == [0, 1, 2]
    assert all(s[2] == 3 for s in seen)
    assert summary['files_scanned'] == 3
    assert summary['cancelled'] is False


def test_scan_stops_when_cancelled(docdir):
    for name in ("a.docx", "b.docx", "c.docx"):
        d = Document()
        d.add_paragraph("2022 CBC")
        d.save(str(docdir / name))

    cancel = threading.Event()

    def progress(name, idx, total):
        if idx == 1:
            cancel.set()          # cancel while the second file is in flight

    results, summary = dp.scan_documents_detailed(
        str(docdir), "2022", True, progress_callback=progress, cancel_event=cancel,
    )
    assert summary['cancelled'] is True
    assert summary['files_scanned'] == 2      # third never started
    assert len(results) == 2


def test_files_scanned_counts_examined_files_not_the_listing(docdir):
    d = Document()
    d.add_paragraph("nothing here")
    d.save(str(docdir / "only.docx"))
    results, summary = dp.scan_documents_detailed(str(docdir), "2022", True)
    assert summary['files_scanned'] == 1
    assert results == []
