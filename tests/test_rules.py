"""
Rule sets, whole-word matching, conflict detection and recursive scanning.

Every test that applies a change asserts on the resulting document, not on
the return value alone.
"""

import os

import pytest
from docx import Document

import document_processor as dp
from conftest import add_run, add_runs, all_text


def _scan(folder, rules, **kw):
    return [m for r in dp.scan_documents(folder, rules, **kw) for m in r.matches]


def _doc(docdir, name, *paragraphs):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path = docdir / name
    doc.save(str(path))
    return str(path)


# ---------------------------------------------------------------- whole word

@pytest.mark.parametrize("text,find,expected", [
    ("CBC", "CBC", [0]),
    ("CBCX", "CBC", []),
    ("XCBC", "CBC", []),
    ("the CBC.", "CBC", [4]),
    ("(CBC)", "CBC", [1]),
    ("x(2022)", "(2022)", [1]),        # punctuation edges need no boundary
    ("20220 and 2022", "2022", [10]),
    ("2022_1 2022", "2022", [7]),      # underscore is a word character
    ("aaa aa", "aa", [4]),             # failed candidate does not hide a later one
    ("a-b a", "a", [0, 4]),
])
def test_whole_word_occurrences(text, find, expected):
    assert dp._find_all_occurrences(text, find, True, whole_word=True) == expected


def test_whole_word_is_case_insensitive_too():
    assert dp._find_all_occurrences("cbc CBCX Cbc", "CBC", False, whole_word=True) == [0, 9]


def test_whole_word_scan_and_apply(docdir):
    path = _doc(docdir, "w.docx", "2022CBC stays, 2022 CBC changes, x2022 stays.")
    matches = _scan(str(docdir), "2022", whole_word=True)
    assert [m.char_offset for m in matches] == [15]
    result = dp.apply_changes(matches, "2025", create_backups=False)
    assert result['total_replaced'] == 1
    assert all_text(path)[0] == "2022CBC stays, 2025 CBC changes, x2022 stays."


# ---------------------------------------------------------------- rule sets

def test_matches_carry_their_rule(docdir):
    _doc(docdir, "r.docx", "Per 2022 CBC and NFPA 13, 2022 edition.")
    rules = [dp.Rule("2022 CBC", "2025 CBC"),
             dp.Rule("NFPA 13, 2022 edition", "NFPA 13, 2025 edition")]
    matches = _scan(str(docdir), rules)
    assert [(m.find_text, m.replace_text, m.char_offset) for m in matches] == [
        ("2022 CBC", "2025 CBC", 4),
        ("NFPA 13, 2022 edition", "NFPA 13, 2025 edition", 17),
    ]


def test_each_rule_writes_its_own_replacement(docdir):
    path = _doc(docdir, "r.docx", "Per 2022 CBC and NFPA 13, 2022 edition.",
                "NFPA 13, 2022 edition first, then 2022 CBC.")
    rules = [dp.Rule("2022 CBC", "2025 CBC"),
             dp.Rule("NFPA 13, 2022 edition", "NFPA 13, 2025 edition")]
    matches = _scan(str(docdir), rules)
    result = dp.apply_changes(matches, create_backups=False)
    assert result['total_replaced'] == 4 and result['errors'] == []
    assert all_text(path) == [
        "Per 2025 CBC and NFPA 13, 2025 edition.",
        "NFPA 13, 2025 edition first, then 2025 CBC.",
    ]


def test_matches_within_a_paragraph_are_in_document_order(docdir):
    _doc(docdir, "o.docx", "b a b a")
    matches = _scan(str(docdir), [dp.Rule("a", "A"), dp.Rule("b", "B")])
    assert [m.char_offset for m in matches] == [0, 2, 4, 6]
    assert [m.find_text for m in matches] == ["b", "a", "b", "a"]


def test_cross_run_match_from_a_second_rule(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_runs(p, "See 20", "22 CFC and 2022 ", "CBC.")
    path = docdir / "s.docx"
    doc.save(str(path))
    rules = [dp.Rule("2022 CBC", "2025 CBC"), dp.Rule("2022 CFC", "2025 CFC")]
    matches = _scan(str(docdir), rules)
    dp.apply_changes(matches, create_backups=False)
    assert all_text(str(path))[0] == "See 2025 CFC and 2025 CBC."


def test_override_replacement_wins_over_rule_text(docdir):
    path = _doc(docdir, "r.docx", "2022 CBC")
    matches = _scan(str(docdir), [dp.Rule("2022 CBC", "2025 CBC")])
    dp.apply_changes(matches, "OVERRIDE", create_backups=False)
    assert all_text(path)[0] == "OVERRIDE"


def test_hand_built_match_without_find_text_uses_match_text(docdir):
    path = _doc(docdir, "h.docx", "2022 CBC")
    matches = _scan(str(docdir), "2022 CBC")
    for m in matches:
        m.find_text = ""
        m.replace_text = "2025 CBC"
    dp.apply_changes(matches, create_backups=False)
    assert all_text(path)[0] == "2025 CBC"


def test_bare_string_and_single_rule_are_equivalent(docdir):
    _doc(docdir, "r.docx", "2022 CBC")
    assert ([m.char_offset for m in _scan(str(docdir), "2022 CBC")]
            == [m.char_offset for m in _scan(str(docdir), [dp.Rule("2022 CBC")])])


def test_rules_with_empty_find_text_are_ignored(docdir):
    _doc(docdir, "r.docx", "2022 CBC")
    assert _scan(str(docdir), [dp.Rule("", "x")]) == []
    assert len(_scan(str(docdir), [dp.Rule("", "x"), dp.Rule("2022 CBC")])) == 1


def test_duplicate_rules_are_rejected(docdir):
    with pytest.raises(ValueError):
        dp.scan_documents(str(docdir), [dp.Rule("CBC", "a"), dp.Rule("CBC", "b")])


def test_case_variant_duplicates_only_clash_when_case_is_ignored(docdir):
    _doc(docdir, "r.docx", "CBC cbc")
    rules = [dp.Rule("CBC", "1"), dp.Rule("cbc", "2")]
    assert len(_scan(str(docdir), rules, case_sensitive=True)) == 2
    with pytest.raises(ValueError):
        dp.scan_documents(str(docdir), rules, case_sensitive=False)


# ---------------------------------------------------------------- conflicts

@pytest.fixture
def conflict_matches(docdir):
    path = _doc(docdir, "c.docx", "2022 CBC and 2022 CFC")
    rules = [dp.Rule("2022", "2025"), dp.Rule("2022 CBC", "2025 CBC")]
    return path, _scan(str(docdir), rules)


def test_overlapping_matches_are_flagged_and_deselected(conflict_matches):
    _path, matches = conflict_matches
    flagged = [(m.find_text, m.char_offset) for m in matches if m.conflict]
    assert flagged == [("2022 CBC", 0), ("2022", 0)] or flagged == [("2022", 0), ("2022 CBC", 0)]
    assert all(not m.is_selected for m in matches if m.conflict)
    clear = [m for m in matches if not m.conflict]
    assert [(m.find_text, m.char_offset, m.is_selected) for m in clear] == [("2022", 13, True)]


def test_apply_skips_conflicts_and_applies_the_rest(conflict_matches):
    path, matches = conflict_matches
    result = dp.apply_changes(matches, create_backups=False)
    assert result['total_replaced'] == 1
    assert all_text(path)[0] == "2022 CBC and 2025 CFC"


def test_a_hand_selected_conflict_is_still_refused(conflict_matches):
    path, matches = conflict_matches
    for m in matches:
        m.is_selected = True
    result = dp.apply_changes(matches, create_backups=False)
    assert result['total_replaced'] == 1
    assert result['total_skipped'] == 2
    assert result['skip_reasons'] == {dp._REASON_CONFLICT: 2}
    assert all(m.skip_reason == dp._REASON_CONFLICT for m in matches if m.conflict)
    assert all_text(path)[0] == "2022 CBC and 2025 CFC"


def test_file_result_counts_conflicts(conflict_matches):
    path, _ = conflict_matches
    fr = dp.scan_documents(os.path.dirname(path),
                           [dp.Rule("2022", "2025"), dp.Rule("2022 CBC", "2025 CBC")])[0]
    assert (fr.match_count, fr.conflict_count, fr.selectable_count, fr.selected_count) == (3, 2, 1, 1)


def test_overlapping_hand_built_edits_leave_the_paragraph_untouched(docdir):
    path = _doc(docdir, "e.docx", "2022 CBC")
    doc = Document(path)
    p_el = dp._collect_paragraphs(doc)[0][0]
    outcomes = dp._apply_edits(p_el, [dp._Edit(0, "2022", "X"), dp._Edit(0, "2022 CBC", "Y")])
    assert [(o.applied, o.reason) for o in outcomes] == [
        (False, dp._REASON_OVERLAP), (False, dp._REASON_OVERLAP)]
    assert dp.paragraph_text(p_el) == "2022 CBC"


def test_three_way_overlap_flags_the_non_adjacent_pair_too(docdir):
    """"2022 CBC", "2022" and "CBC": the third overlaps the first, not the second."""
    path = _doc(docdir, "t.docx", "2022 CBC here")
    rules = [dp.Rule("2022 CBC", "a"), dp.Rule("2022", "b"), dp.Rule("CBC", "c")]
    matches = _scan(str(docdir), rules)
    assert [(m.find_text, m.conflict) for m in matches] == [
        ("2022 CBC", True), ("2022", True), ("CBC", True)]
    for m in matches:
        m.is_selected = True
    result = dp.apply_changes(matches, create_backups=False)
    assert result['total_replaced'] == 0 and result['total_skipped'] == 3
    assert all_text(path)[0] == "2022 CBC here"


def test_three_way_overlap_in_hand_built_edits(docdir):
    path = _doc(docdir, "t.docx", "2022 CBC here")
    doc = Document(path)
    p_el = dp._collect_paragraphs(doc)[0][0]
    outcomes = dp._apply_edits(p_el, [dp._Edit(0, "2022 CBC", "a"), dp._Edit(0, "2022", "b"),
                                      dp._Edit(5, "CBC", "c"), dp._Edit(9, "here", "there")])
    assert sorted((o.offset, o.applied) for o in outcomes) == [
        (0, False), (0, False), (5, False), (9, True)]
    assert dp.paragraph_text(p_el) == "2022 CBC there"


def test_partial_overlap_is_also_refused(docdir):
    path = _doc(docdir, "e.docx", "abcdef")
    doc = Document(path)
    p_el = dp._collect_paragraphs(doc)[0][0]
    outcomes = dp._apply_edits(p_el, [dp._Edit(0, "abc", "X"), dp._Edit(2, "cde", "Y"),
                                      dp._Edit(5, "f", "Z")])
    assert sorted((o.offset, o.applied) for o in outcomes) == [(0, False), (2, False), (5, True)]
    assert dp.paragraph_text(p_el) == "abcdeZ"


def test_identical_duplicate_edits_apply_once(docdir):
    path = _doc(docdir, "e.docx", "2022 CBC")
    doc = Document(path)
    p_el = dp._collect_paragraphs(doc)[0][0]
    outcomes = dp._apply_edits(p_el, [dp._Edit(0, "2022", "2025")] * 2)
    assert [(o.offset, o.applied) for o in outcomes] == [(0, True)]
    assert dp.paragraph_text(p_el) == "2025 CBC"


def test_skip_reason_is_recorded_for_a_moved_match(docdir):
    path = _doc(docdir, "m.docx", "2022 CBC")
    matches = _scan(str(docdir), "2022 CBC")
    matches[0].char_offset = 3
    matches[0].file_mtime = 0.0
    matches[0].file_size = 0
    result = dp.apply_changes(matches, "x", create_backups=False)
    assert result['total_skipped'] == 1
    assert matches[0].skip_reason == dp._REASON_MOVED
    assert result['skip_reasons'] == {dp._REASON_MOVED: 1}
    assert all_text(path)[0] == "2022 CBC"


# ---------------------------------------------------------------- recursion

def test_non_recursive_scan_sees_only_the_top_folder(nested_folder_docs, docdir):
    results = dp.scan_documents(str(docdir), "2022 CBC")
    assert [r.file_name for r in results] == ["a.docx"]


def test_recursive_scan_walks_subfolders_in_order(nested_folder_docs, docdir):
    results = dp.scan_documents(str(docdir), "2022 CBC", recursive=True)
    assert [r.file_name for r in results] == [
        "a.docx", os.path.join("sub", "b.docx"), os.path.join("sub", "deeper", "c.docx")]
    assert [r.file_path for r in results] == nested_folder_docs


def test_recursive_scan_skips_lock_files(nested_folder_docs, docdir):
    _results, summary = dp.scan_documents_detailed(str(docdir), "2022 CBC", recursive=True)
    assert summary['files_scanned'] == 3


def test_recursive_progress_reports_relative_names(nested_folder_docs, docdir):
    seen = []
    dp.scan_documents_detailed(str(docdir), "2022 CBC", recursive=True,
                               progress_callback=lambda name, i, n: seen.append(name))
    assert seen == ["a.docx", os.path.join("sub", "b.docx"),
                    os.path.join("sub", "deeper", "c.docx")]


def test_recursive_apply_modifies_deep_files_and_backs_up_beside_them(nested_folder_docs, docdir):
    matches = _scan(str(docdir), [dp.Rule("2022 CBC", "2025 CBC")], recursive=True)
    result = dp.apply_changes(matches, create_backups=True)
    assert result['files_modified'] == 3
    deep = nested_folder_docs[2]
    assert all_text(deep)[0] == "Nested cites 2025 CBC."
    backup = result['backup_map'][deep]
    assert os.path.dirname(backup) == os.path.dirname(deep)
    assert sorted(result['backup_map']) == sorted(nested_folder_docs)
    assert result['backups'] == [result['backup_map'][p] for p in nested_folder_docs]


def test_unreadable_top_folder_is_reported_in_recursive_mode(tmp_path):
    missing = str(tmp_path / "nope")
    results, _summary = dp.scan_documents_detailed(missing, "x", recursive=True)
    assert len(results) == 1 and results[0].error
