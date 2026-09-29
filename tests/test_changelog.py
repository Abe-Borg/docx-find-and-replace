"""
Change log CSV, backup discovery and restore.
"""

import csv
import os
import shutil
import time
from datetime import datetime

import pytest
from docx import Document

import document_processor as dp
from conftest import all_text


def _scan(path, find):
    return [m for r in dp.scan_documents(os.path.dirname(path), find) for m in r.matches]


def _read_csv(path):
    with open(path, "rb") as fh:
        raw = fh.read()
    assert raw.startswith(b"\xef\xbb\xbf")      # BOM, so Excel reads UTF-8
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.reader(fh))


# ---------------------------------------------------------------- change log

def test_change_log_records_every_match(docdir, tmp_path):
    doc = Document()
    doc.add_heading("GENERAL", level=1)
    doc.add_paragraph("A 2022 CBC, then\tanother 2022 CBC, then a third 2022 CBC.")
    path = str(docdir / "log.docx")
    doc.save(path)

    matches = _scan(path, "2022 CBC")
    matches[1].is_selected = False
    matches[2].char_offset = 12                 # will be skipped as moved
    for m in matches:
        m.replace_text = "2025 CBC"
    result = dp.apply_changes(matches, create_backups=True)

    log = str(tmp_path / "changes.csv")
    assert dp.write_change_log(log, matches, result['backup_map']) == log
    rows = _read_csv(log)
    assert rows[0] == list(dp.CHANGE_LOG_COLUMNS)
    assert len(rows) == 4

    by_status = {r[8]: r for r in rows[1:]}
    applied = by_status["applied"]
    assert applied[0] == path
    assert applied[1] == "body: Paragraph 2"
    assert applied[2] == "GENERAL"
    assert applied[3:5] == ["2022 CBC", "2025 CBC"]
    assert applied[6] == "2022 CBC"
    assert applied[9] == result['backup_map'][path]
    assert "not selected" in by_status
    assert f"skipped: {dp._REASON_MOVED}" in by_status
    # Tabs and commas in context survive the round trip.
    assert any("\t" in r[5] or "\t" in r[7] for r in rows[1:])
    assert any("," in r[5] or "," in r[7] for r in rows[1:])


def test_change_log_round_trips_newlines_and_unicode(tmp_path):
    m = dp.Match(file_path="C:\\specs\\ünïcode.docx", location_type='body',
                 location_detail='Paragraph 1', paragraph_key='/word/document.xml#0',
                 char_offset=3, context_before="line one\nline, two ", match_text="2022 CBC",
                 context_after=" – dash", find_text="2022 CBC", replace_text="2025 CBC",
                 section="GENERAL > “REFERENCES”", applied=True)
    log = str(tmp_path / "u.csv")
    dp.write_change_log(log, [m])
    rows = _read_csv(log)
    assert rows[1] == ["C:\\specs\\ünïcode.docx", "body: Paragraph 1", "GENERAL > “REFERENCES”",
                       "2022 CBC", "2025 CBC", "line one\nline, two ", "2022 CBC", " – dash",
                       "applied", ""]


def test_change_log_with_no_matches_still_has_a_header(tmp_path):
    log = str(tmp_path / "empty.csv")
    dp.write_change_log(log, [])
    assert _read_csv(log) == [list(dp.CHANGE_LOG_COLUMNS)]


# ---------------------------------------------------------------- find_backups

def _touch(path, content=b"x"):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as fh:
        fh.write(content)
    return path


def test_find_backups_parses_stamp_and_sequence(tmp_path):
    d = str(tmp_path)
    _touch(os.path.join(d, "spec.docx"))
    _touch(os.path.join(d, "spec.docx.20260909-141530.bak"))
    _touch(os.path.join(d, "spec.docx.20260909-141530-1.bak"))
    _touch(os.path.join(d, "spec.docx.20260910-090000.bak"))
    _touch(os.path.join(d, "Other.Spec.v2.DOCX.20260101-000000.BAK"))
    _touch(os.path.join(d, "unrelated.bak"))
    _touch(os.path.join(d, "spec.docx.bak"))
    _touch(os.path.join(d, "spec.docx.20261399-141530.bak"))     # impossible date

    backups = dp.find_backups(d)
    assert [(os.path.basename(b.document_path), b.timestamp, b.sequence) for b in backups] == [
        ("Other.Spec.v2.DOCX", datetime(2026, 1, 1, 0, 0, 0), 0),
        ("spec.docx", datetime(2026, 9, 10, 9, 0, 0), 0),
        ("spec.docx", datetime(2026, 9, 9, 14, 15, 30), 1),
        ("spec.docx", datetime(2026, 9, 9, 14, 15, 30), 0),
    ]
    assert backups[1].backup_path == os.path.join(d, "spec.docx.20260910-090000.bak")


def test_find_backups_recursive(tmp_path):
    d = str(tmp_path)
    _touch(os.path.join(d, "a.docx.20260101-000000.bak"))
    _touch(os.path.join(d, "sub", "b.docx.20260101-000000.bak"))
    assert [os.path.basename(b.backup_path) for b in dp.find_backups(d)] == [
        "a.docx.20260101-000000.bak"]
    assert [os.path.basename(b.backup_path) for b in dp.find_backups(d, recursive=True)] == [
        "a.docx.20260101-000000.bak", "b.docx.20260101-000000.bak"]
    assert dp.find_backups(d, recursive=True)[1].document_path == os.path.join(d, "sub", "b.docx")


def test_find_backups_sees_what_apply_changes_wrote(docdir):
    doc = Document()
    doc.add_paragraph("2022 CBC")
    path = str(docdir / "real.docx")
    doc.save(path)
    result = dp.apply_changes(_scan(path, "2022 CBC"), "2025 CBC", create_backups=True)
    backups = dp.find_backups(str(docdir))
    assert [b.backup_path for b in backups] == result['backups']
    assert backups[0].document_path == path


# ---------------------------------------------------------------- restore

def test_restore_puts_the_original_bytes_back_and_keeps_the_backup(docdir):
    doc = Document()
    doc.add_paragraph("2022 CBC")
    path = str(docdir / "r.docx")
    doc.save(path)
    original = open(path, "rb").read()

    dp.apply_changes(_scan(path, "2022 CBC"), "2025 CBC", create_backups=True)
    assert all_text(path)[0] == "2025 CBC"

    backups = dp.find_backups(str(docdir))
    result = dp.restore_backups(backups)
    assert result == {'restored': 1, 'restored_paths': [path], 'errors': []}
    assert open(path, "rb").read() == original
    assert os.path.exists(backups[0].backup_path)
    assert all_text(path)[0] == "2022 CBC"
    assert not [f for f in os.listdir(str(docdir)) if f.endswith(".tmp")]


def test_restore_refuses_two_backups_for_one_document_before_touching_anything(tmp_path):
    d = str(tmp_path)
    doc = _touch(os.path.join(d, "s.docx"), b"current")
    b1 = _touch(os.path.join(d, "s.docx.20260101-000000.bak"), b"one")
    b2 = _touch(os.path.join(d, "s.docx.20260102-000000.bak"), b"two")
    backups = dp.find_backups(d)
    assert len(backups) == 2
    with pytest.raises(ValueError):
        dp.restore_backups(backups)
    assert open(doc, "rb").read() == b"current"


def test_restore_reports_a_missing_backup_and_continues(tmp_path):
    d = str(tmp_path)
    good_doc = _touch(os.path.join(d, "good.docx"), b"changed")
    good = dp._parse_backup(_touch(os.path.join(d, "good.docx.20260101-000000.bak"), b"orig"))
    missing = dp.Backup(os.path.join(d, "gone.docx"),
                        os.path.join(d, "gone.docx.20260101-000000.bak"),
                        datetime(2026, 1, 1), 0)
    result = dp.restore_backups([missing, good])
    assert result['restored'] == 1
    assert result['restored_paths'] == [good_doc]
    assert len(result['errors']) == 1 and "gone.docx" in result['errors'][0]
    assert open(good_doc, "rb").read() == b"orig"
    assert not [f for f in os.listdir(d) if f.endswith(".tmp")]


def test_restore_progress_callback(tmp_path):
    d = str(tmp_path)
    _touch(os.path.join(d, "a.docx"), b"x")
    _touch(os.path.join(d, "a.docx.20260101-000000.bak"), b"y")
    seen = []
    dp.restore_backups(dp.find_backups(d), progress_callback=lambda n, i, t: seen.append((n, i, t)))
    assert seen == [("a.docx", 0, 1)]


def test_restore_makes_the_preview_stale(docdir):
    doc = Document()
    doc.add_paragraph("2022 CBC")
    path = str(docdir / "stale.docx")
    doc.save(path)
    dp.apply_changes(_scan(path, "2022 CBC"), "2025 CBC", create_backups=True)
    matches = _scan(path, "2025 CBC")
    time.sleep(0.01)
    dp.restore_backups(dp.find_backups(str(docdir)))
    result = dp.apply_changes(matches, "2028 CBC", create_backups=False)
    assert result['total_replaced'] == 0 and result['total_skipped'] == 1
    assert all_text(path)[0] == "2022 CBC"
