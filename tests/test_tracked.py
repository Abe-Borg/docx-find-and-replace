"""
Tracked-changes output mode.

Every test saves the document, reloads it and asserts on the XML that Word
will read: the deletions, the insertions, their ids, and what the text looks
like after Word's Accept All and Reject All.
"""

import re

import pytest
from docx import Document
from docx.oxml.ns import qn

import document_processor as dp
from conftest import (accept_all, add_deletion, add_formatted_run, add_hyperlink,
                      add_notes_part, add_run, add_run_with_rpr_change, add_runs,
                      add_tab_run, add_tracked_insertion, all_revision_ids, all_text,
                      count_drawings, reject_all, revisions)


def _scan(path, find, **kw):
    import os
    return [m for r in dp.scan_documents(os.path.dirname(path), find, **kw) for m in r.matches]


def _apply(matches, replace=None, **kw):
    kw.setdefault("create_backups", False)
    kw.setdefault("author", "Tester")
    return dp.apply_changes(matches, replace, tracked_changes=True, **kw)


def _body(path):
    return Document(path).element.body


# ---------------------------------------------------------------- basics

def test_single_run_becomes_one_deletion_and_one_insertion(docdir):
    doc = Document()
    doc.add_paragraph("Reference the 2022 CBC for compliance.")
    path = str(docdir / "one.docx")
    doc.save(path)

    result = _apply(_scan(path, "2022 CBC"), "2025 CBC")
    assert result['total_replaced'] == 1 and result['errors'] == []

    assert revisions(path) == [("del", "1", "Tester", "2022 CBC"),
                               ("ins", "2", "Tester", "2025 CBC")]
    assert all_text(path)[0] == "Reference the 2025 CBC for compliance."
    assert accept_all(path)[0] == "Reference the 2025 CBC for compliance."
    assert reject_all(path)[0] == "Reference the 2022 CBC for compliance."

    del_text = _body(path).find(".//" + qn("w:delText"))
    assert del_text.get(qn("xml:space")) == "preserve"
    Document(path)   # still opens


def test_untouched_text_stays_in_plain_runs(docdir):
    doc = Document()
    doc.add_paragraph("Reference the 2022 CBC for compliance.")
    path = str(docdir / "one.docx")
    doc.save(path)
    _apply(_scan(path, "2022 CBC"), "2025 CBC")

    p = _body(path).find(qn("w:p"))
    plain = [r for r in p.findall(qn("w:r"))]     # direct children only
    assert ["".join(t.text for t in r.iter(qn("w:t"))) for r in plain] == [
        "Reference the ", " for compliance."]


def test_split_runs_each_get_their_own_deletion(split_run_doc):
    _apply(_scan(split_run_doc, "2022 CBC"), "2025 CBC")
    revs = revisions(split_run_doc)
    assert [(t, txt) for t, _i, _a, txt in revs] == [
        ("del", "20"), ("del", "22 "), ("del", "CBC"), ("ins", "2025 CBC")]
    assert accept_all(split_run_doc)[0] == "Reference the 2025 CBC for compliance."
    assert reject_all(split_run_doc)[0] == "Reference the 2022 CBC for compliance."


def test_insertion_copies_formatting_of_the_first_deleted_run(formatted_split_doc):
    _apply(_scan(formatted_split_doc, "2022 CBC"), "2025 CBC")
    body = _body(formatted_split_doc)
    ins_run = body.find(".//" + qn("w:ins") + "/" + qn("w:r"))
    rpr = ins_run.find(qn("w:rPr"))
    assert rpr.find(qn("w:b")) is not None
    assert rpr.find(qn("w:color")) is None       # the first run was bold, not red

    # The split-off plain suffix keeps its own (absent) formatting.
    assert accept_all(formatted_split_doc)[0] == "2025 CBC ref"


def test_prefix_and_suffix_of_a_partially_covered_run_keep_formatting(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_formatted_run(p, "See 2022 CBC now", bold=True, color="FF0000")
    path = str(docdir / "partial.docx")
    doc.save(path)

    _apply(_scan(path, "2022 CBC"), "2025 CBC")
    p_el = _body(path).find(qn("w:p"))
    runs = p_el.findall(qn("w:r"))
    assert ["".join(t.text for t in r.iter(qn("w:t"))) for r in runs] == ["See ", " now"]
    for r in runs:
        assert r.find(qn("w:rPr") + "/" + qn("w:b")) is not None
        assert r.find(qn("w:rPr") + "/" + qn("w:color")).get(qn("w:val")) == "FF0000"
    assert accept_all(path)[0] == "See 2025 CBC now"
    assert reject_all(path)[0] == "See 2022 CBC now"


# ---------------------------------------------------------------- run contents

def test_inline_image_is_not_marked_deleted(image_in_run_doc):
    _apply(_scan(image_in_run_doc, "2022 CBC"), "2025 CBC")
    body = _body(image_in_run_doc)
    assert count_drawings(image_in_run_doc) == 1
    drawing = body.find(".//" + qn("w:drawing"))
    ancestors = {a.tag for a in drawing.iterancestors()}
    assert qn("w:del") not in ancestors and qn("w:ins") not in ancestors
    assert accept_all(image_in_run_doc)[0] == "See Figure 2025 CBC ref"
    assert reject_all(image_in_run_doc)[0] == "See Figure 2022 CBC ref"


def test_tracked_formatting_change_is_kept_on_the_original_run_only(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_run_with_rpr_change(p, "Old 2022 CBC text")
    path = str(docdir / "rprchange.docx")
    doc.save(path)

    _apply(_scan(path, "2022 CBC"), "2025 CBC")
    body = _body(path)
    changes = body.findall(".//" + qn("w:rPrChange"))
    assert len(changes) == 1
    assert changes[0].getparent().getparent().getparent().tag == qn("w:del")
    ids = all_revision_ids(path)
    assert len(ids) == len(set(ids))
    assert int(min(ids, key=int)) == 902
    assert accept_all(path)[0] == "Old 2025 CBC text"


def test_tab_inside_the_match_is_deleted_as_a_tab(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "2022")
    add_tab_run(p)
    add_run(p, "CBC")
    path = str(docdir / "tab.docx")
    doc.save(path)

    _apply(_scan(path, "2022\tCBC"), "2025 CBC")
    body = _body(path)
    dels = body.findall(".//" + qn("w:del"))
    assert len(dels) == 3
    assert dels[1].find(qn("w:r") + "/" + qn("w:tab")) is not None
    assert all_text(path)[0] == "2025 CBC"
    assert reject_all(path)[0] == "2022\tCBC"


def test_empty_replacement_writes_a_deletion_only(docdir):
    doc = Document()
    doc.add_paragraph("Delete 2022 CBC please")
    path = str(docdir / "empty.docx")
    doc.save(path)
    _apply(_scan(path, "2022 CBC"), "")
    assert [t for t, _i, _a, _x in revisions(path)] == ["del"]
    assert accept_all(path)[0] == "Delete  please"


# ---------------------------------------------------------------- ids and attributes

def test_ids_continue_above_existing_revisions(tracked_changes_doc):
    # The fixture holds ids 900 (ins) and 901 (del); the match sits in the
    # plain "Kept " run? No - it sits inside the insertion, which is skipped,
    # so use a plain paragraph with the same existing ids.
    doc = Document(tracked_changes_doc)
    doc.add_paragraph("Plain 2022 CBC")
    doc.save(tracked_changes_doc)

    matches = [m for m in _scan(tracked_changes_doc, "2022 CBC") if m.location_detail == "Paragraph 2"]
    _apply(matches, "2025 CBC")
    new_ids = [i for i in all_revision_ids(tracked_changes_doc) if i not in ("900", "901")]
    assert new_ids == ["902", "903"]


def test_ids_are_unique_across_body_and_footnotes(notes_doc):
    matches = [m for m in _scan(notes_doc, "2022 CBC") if m.location_type != 'comment']
    result = _apply(matches, "2025 CBC")
    assert result['total_replaced'] == 3
    ids = all_revision_ids(notes_doc)
    assert len(ids) == 6 and len(set(ids)) == 6


def test_author_and_date_attributes(docdir):
    doc = Document()
    doc.add_paragraph("2022 CBC")
    path = str(docdir / "attrs.docx")
    doc.save(path)
    _apply(_scan(path, "2022 CBC"), "2025 CBC", author="Jane Reviewer")
    for el in _body(path).iter(qn("w:ins"), qn("w:del")):
        assert el.get(qn("w:author")) == "Jane Reviewer"
        assert re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z", el.get(qn("w:date")))


def test_default_author_is_never_empty(docdir):
    doc = Document()
    doc.add_paragraph("2022 CBC")
    path = str(docdir / "author.docx")
    doc.save(path)
    dp.apply_changes(_scan(path, "2022 CBC"), "2025 CBC", create_backups=False,
                     tracked_changes=True)
    assert all(a for _t, _i, a, _x in revisions(path))


# ---------------------------------------------------------------- hyperlinks

def test_match_inside_a_hyperlink_keeps_the_link(hyperlink_doc):
    _apply(_scan(hyperlink_doc, "2022 CBC"), "2025 CBC")
    body = _body(hyperlink_doc)
    link = body.find(".//" + qn("w:hyperlink"))
    assert link.get(qn("r:id"))
    assert link.find(qn("w:del")) is not None and link.find(qn("w:ins")) is not None
    assert accept_all(hyperlink_doc)[0] == "See 2025 CBC link and done."
    assert reject_all(hyperlink_doc)[0] == "See 2022 CBC link and done."


def test_match_starting_before_a_hyperlink_lands_its_insertion_inside(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_run(p, "Per 2022 ")
    add_hyperlink(p, "CBC link")
    path = str(docdir / "span.docx")
    doc.save(path)

    _apply(_scan(path, "2022 CBC"), "2025 CBC")
    body = _body(path)
    ins = body.find(".//" + qn("w:ins"))
    assert ins.getparent().tag == qn("w:hyperlink")
    assert accept_all(path)[0] == "Per 2025 CBC link"
    assert reject_all(path)[0] == "Per 2022 CBC link"


def test_match_starting_inside_a_hyperlink_lands_its_insertion_outside(docdir):
    doc = Document()
    p = doc.add_paragraph()
    add_hyperlink(p, "See 2022")
    add_run(p, " CBC now")
    path = str(docdir / "span2.docx")
    doc.save(path)

    _apply(_scan(path, "2022 CBC"), "2025 CBC")
    body = _body(path)
    ins = body.find(".//" + qn("w:ins"))
    assert ins.getparent().tag == qn("w:p")
    assert accept_all(path)[0] == "See 2025 CBC now"


# ---------------------------------------------------------------- refusals

def test_match_inside_an_existing_insertion_is_skipped_untouched(tracked_changes_doc):
    before = open(tracked_changes_doc, "rb").read()
    matches = _scan(tracked_changes_doc, "2022 CBC")
    assert len(matches) == 1
    result = _apply(matches, "2025 CBC")
    assert result['total_replaced'] == 0
    assert result['total_skipped'] == 1
    assert matches[0].skip_reason == dp._REASON_NESTED_INS
    assert result['skip_reasons'] == {dp._REASON_NESTED_INS: 1}
    assert open(tracked_changes_doc, "rb").read() == before


def test_comment_matches_are_skipped_while_the_body_is_applied(notes_doc):
    matches = _scan(notes_doc, "2022 CBC")
    result = _apply(matches, "2025 CBC")
    assert result['total_replaced'] == 3
    assert result['total_skipped'] == 1
    comment = [m for m in matches if m.location_type == 'comment'][0]
    assert comment.applied is False
    assert comment.skip_reason == dp._REASON_COMMENT_TRACKED
    texts = all_text(notes_doc)
    assert "Body cites 2025 CBC here." in texts
    assert "Reviewer asks about 2022 CBC." in texts


def test_second_tracked_run_skips_its_own_earlier_insertions(docdir):
    doc = Document()
    doc.add_paragraph("2022 CBC and 2022 CFC")
    path = str(docdir / "twice.docx")
    doc.save(path)
    _apply(_scan(path, "2022 CBC"), "2025 CBC")
    # "2025 CBC" now lives in a w:ins; a rule touching it must be refused.
    result = _apply(_scan(path, "2025"), "2028")
    assert result['total_replaced'] == 0
    assert result['skip_reasons'] == {dp._REASON_NESTED_INS: 1}
    # ...while text outside the insertion is still fair game.
    result = _apply(_scan(path, "2022 CFC"), "2025 CFC")
    assert result['total_replaced'] == 1
    assert accept_all(path)[0] == "2025 CBC and 2025 CFC"
    assert reject_all(path)[0] == "2022 CBC and 2022 CFC"


def test_deleted_text_is_never_scanned(tracked_changes_doc):
    matches = _scan(tracked_changes_doc, "deleted")
    assert matches == []


# ---------------------------------------------------------------- several edits

def test_multiple_matches_in_one_run(repeated_doc):
    _apply(_scan(repeated_doc, "2022"), "2025")
    assert all_text(repeated_doc)[0] == "aaaa and 2025 2025 2025"
    assert accept_all(repeated_doc)[0] == "aaaa and 2025 2025 2025"
    assert reject_all(repeated_doc)[0] == "aaaa and 2022 2022 2022"
    assert [t for t, _i, _a, _x in revisions(repeated_doc)] == ["del", "ins"] * 3


def test_two_rules_in_one_paragraph(docdir):
    doc = Document()
    doc.add_paragraph("Per 2022 CBC and NFPA 13, 2022 edition.")
    path = str(docdir / "rules.docx")
    doc.save(path)
    rules = [dp.Rule("2022 CBC", "2025 CBC"),
             dp.Rule("NFPA 13, 2022 edition", "NFPA 13, 2025 edition")]
    _apply(_scan(path, rules))
    assert accept_all(path)[0] == "Per 2025 CBC and NFPA 13, 2025 edition."
    assert reject_all(path)[0] == "Per 2022 CBC and NFPA 13, 2022 edition."
    ids = all_revision_ids(path)
    assert len(ids) == 4 and len(set(ids)) == 4


def test_header_and_table_cells_are_tracked_too(merged_and_linked_doc):
    result = _apply(_scan(merged_and_linked_doc, "2022 CBC"), "2025 CBC")
    assert result['total_replaced'] == 4 and result['errors'] == []
    texts = accept_all(merged_and_linked_doc)
    assert "HEADER 2025 CBC" in texts
    assert "Merged: 2025 CBC note" in texts
    ids = all_revision_ids(merged_and_linked_doc)
    assert len(ids) == len(set(ids))


def test_textbox_is_tracked(textbox_doc):
    _apply(_scan(textbox_doc, "2022 CBC"), "2025 CBC")
    assert "Textbox 2025 CBC" in accept_all(textbox_doc)
    assert "Textbox 2022 CBC" in reject_all(textbox_doc)


def test_in_place_mode_is_unchanged_by_the_tracked_machinery(split_run_doc):
    dp.apply_changes(_scan(split_run_doc, "2022 CBC"), "2025 CBC", create_backups=False)
    assert revisions(split_run_doc) == []
    assert all_text(split_run_doc)[0] == "Reference the 2025 CBC for compliance."
